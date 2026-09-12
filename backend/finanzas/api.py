import csv
import hashlib
import io
import logging
import os
import re
import secrets
import unicodedata
from calendar import monthrange
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

from django.conf import settings
from django.core.files.storage import default_storage
from django.db import IntegrityError, transaction
from django.db.models import Count, Exists, OuterRef, Prefetch, Q, Subquery, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from ninja import File, Router, Schema
from ninja.errors import HttpError
from ninja.files import UploadedFile
from django.db.models.functions import Coalesce
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from accounts.audit import audit
from accounts.security import (
    get_allowed_entity_ids,
    get_current_capa,
    plan_has_module,
    require_plan_feature,
    require_plan_module,
    require_write_access,
)
from billing.background_jobs import enqueue_background_job, store_background_upload
from crm.models import Cliente
from empresas.models import CapaNegocio, EntidadNegocio, PERIODICIDAD_CHOICES

from .bank_statement_parser import (
    BankStatementParseError,
    ParsedBankStatement,
    parse_bank_statement_pdf,
)
from .models import (
    CargaConciliacion,
    CuentaBancaria,
    CuentaPorCobrar,
    CuentaPorPagar,
    EventoFinanciero,
    EventoFinancieroPartida,
    PagoCuentaPorCobrar,
    PagoCuentaPorPagar,
    ProgramacionCuentaPorPagar,
    SaldoCliente,
    Transaccion,
)
from .eventos import (
    apply_financial_event,
    build_events_dashboard,
    reconcile_event_with_transaction,
    serialize_event_detail,
    update_event_case_status,
    upsert_financial_event,
)
from .services import (
    auto_match_bank_transaction,
    apply_bank_transaction_to_payable,
    apply_bank_transaction_to_receivables,
    auto_reconcile_pending_bank_transactions,
    build_bank_reconciliation_suggestions,
    build_bank_transaction_detail,
    build_cxc_customer_detail,
    build_global_cxc_dashboard,
    build_global_cxp_dashboard,
    build_finance_summary,
    import_bank_transactions,
    serialize_bank_account,
    serialize_reconciliation_load,
    reject_payable_payment,
    reject_receivable_payment,
    resolve_due_date,
    serialize_payment_cxc,
    serialize_payment_cxp,
    serialize_transaction,
    set_bank_transaction_status,
    decimal_to_float,
    generate_payables_for_entity,
    register_payment_for_entity,
    register_payment_for_payable,
    validate_payable_payment,
    validate_receivable_payment,
)

router = Router()
logger = logging.getLogger(__name__)


def build_suggestion_audit_context(payload) -> dict:
    raw_reasons = getattr(payload, "sugerencia_razones", None) or []
    reasons = [
        str(reason).strip()[:240]
        for reason in raw_reasons
        if str(reason).strip()
    ][:5]
    suggestion_id = (getattr(payload, "sugerencia_id", "") or "").strip()
    confidence = (getattr(payload, "sugerencia_confianza", "") or "").strip()
    readiness = (getattr(payload, "sugerencia_estado", "") or "").strip()
    score = getattr(payload, "sugerencia_score", None)
    if not any([suggestion_id, confidence, readiness, score is not None, reasons]):
        return {}
    return {
        "id": suggestion_id or None,
        "score": score,
        "confidence_band": confidence or None,
        "readiness_state": readiness or None,
        "requires_manual_review": bool(
            getattr(payload, "sugerencia_requiere_revision", False)
        ),
        "reasons": reasons,
    }


def build_transaction_audit_context(
    transaction_item: Transaccion,
    *,
    previous_status: str | None = None,
    notes: str = "",
) -> dict:
    serialized = serialize_transaction(transaction_item)
    fecha_pago = serialized.get("fecha_pago")
    if hasattr(fecha_pago, "isoformat"):
        fecha_pago = fecha_pago.isoformat()
    auto_match = serialized.get("auto_match") or {}
    auto_match_summary = {}
    if auto_match:
        auto_match_summary = {
            "scope": auto_match.get("scope"),
            "match_status": auto_match.get("match_status"),
            "matched": auto_match.get("matched"),
            "confidence_score": auto_match.get("confidence_score"),
            "confidence_band": auto_match.get("confidence_band"),
            "matched_type": auto_match.get("matched_type"),
            "matched_id": auto_match.get("matched_id"),
            "reasons": (auto_match.get("reasons") or [])[:5],
        }
    return {
        "id": serialized.get("id"),
        "estatus_anterior": previous_status,
        "estatus_actual": serialized.get("estatus_conciliacion"),
        "tipo_movimiento": serialized.get("tipo_movimiento"),
        "monto": serialized.get("monto"),
        "monto_aplicado": serialized.get("monto_aplicado"),
        "monto_disponible": serialized.get("monto_disponible"),
        "fecha_pago": fecha_pago,
        "cuenta_bancaria_id": serialized.get("cuenta_bancaria_id"),
        "cuenta_bancaria_nombre": serialized.get("cuenta_bancaria_nombre"),
        "cuenta_por_cobrar_id": serialized.get("cuenta_por_cobrar_id"),
        "cuenta_por_pagar_id": serialized.get("cuenta_por_pagar_id"),
        "evento_id": serialized.get("evento_id"),
        "referencia": serialized.get("numero_referencia") or serialized.get("folio_bancario"),
        "concepto_bancario": serialized.get("concepto_bancario"),
        "nota_operador": notes.strip()[:500] or None,
        "auto_match": auto_match_summary,
    }


def require_finance_module(request, module_key: str) -> None:
    require_plan_module(request, module_key)


def plan_allows_finance_module(request, module_key: str) -> bool:
    return plan_has_module(request, module_key)


def require_finance_summary_access(request) -> None:
    if any(
        plan_allows_finance_module(request, module_key)
        for module_key in ("cxc", "cxp", "conciliacion")
    ):
        return
    raise HttpError(403, "Tu plan actual no incluye este modulo financiero.")


def require_cxc_access(request) -> None:
    require_finance_module(request, "cxc")


def resolve_receivable_capa(cuenta: CuentaPorCobrar):
    if cuenta.entidad_relacionada_id:
        return cuenta.entidad_relacionada.capa_negocio
    return cuenta.cliente_relacionado.entidad_relacionada.capa_negocio


def resolve_event_capa(request, event: EventoFinanciero):
    if event.entidad_relacionada_id:
        return event.entidad_relacionada.capa_negocio
    if event.cliente_relacionado_id:
        return event.cliente_relacionado.entidad_relacionada.capa_negocio
    return get_current_capa(request)


def require_batch_import_access(request) -> None:
    require_plan_feature(request, "batch_import")


def token_fingerprint(value: str) -> str:
    if not value:
        return "empty"
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:10]


def validate_cxc_sync_token(request) -> None:
    expected_token = (getattr(settings, "FINANZAS_CXC_SYNC_TOKEN", "") or "").strip()
    authorization = (request.headers.get("Authorization") or "").strip()
    bearer_token = ""
    if authorization.lower().startswith("bearer "):
        bearer_token = authorization[7:].strip()
    provided_token = (
        request.headers.get("X-BetterP-CxC-Sync-Token")
        or bearer_token
        or request.GET.get("token")
        or ""
    ).strip()
    if not expected_token:
        raise HttpError(503, "Generacion programada de CxC sin token configurado.")
    if not secrets.compare_digest(provided_token, expected_token):
        logger.warning(
            "Token invalido en generacion programada de CxC. expected_len=%s "
            "provided_len=%s expected_fp=%s provided_fp=%s",
            len(expected_token),
            len(provided_token),
            token_fingerprint(expected_token),
            token_fingerprint(provided_token),
        )
        raise HttpError(401, "Token invalido para generacion programada de CxC.")


def receivable_scope_filter(allowed_entity_ids: list[int]) -> Q:
    return Q(entidad_relacionada_id__in=allowed_entity_ids) | Q(
        entidad_relacionada__isnull=True,
        cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
    )


def payable_scope_filter(allowed_entity_ids: list[int]) -> Q:
    return Q(entidad_relacionada_id__in=allowed_entity_ids)


def event_scope_filter(allowed_entity_ids: list[int]) -> Q:
    return (
        Q(entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            entidad_relacionada__isnull=True,
            cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
        | Q(partidas__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(partidas__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            partidas__cuenta_por_cobrar_relacionada__entidad_relacionada_id__in=allowed_entity_ids
        )
        | Q(
            partidas__cuenta_por_cobrar_relacionada__entidad_relacionada__isnull=True,
            partidas__cuenta_por_cobrar_relacionada__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
        | Q(
            partidas__cuenta_por_pagar_relacionada__entidad_relacionada_id__in=allowed_entity_ids
        )
    )


def transaction_scope_filter(allowed_entity_ids: list[int], capa_id: int | None = None) -> Q:
    scope = (
        Q(cuenta_por_cobrar_relacionada__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            cuenta_por_cobrar_relacionada__entidad_relacionada__isnull=True,
            cuenta_por_cobrar_relacionada__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
        | Q(cuenta_por_pagar_relacionada__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(cuenta_bancaria_relacionada__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(evento_relacionado__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            evento_relacionado__entidad_relacionada__isnull=True,
            evento_relacionado__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
        | Q(evento_relacionado__partidas__entidad_relacionada_id__in=allowed_entity_ids)
    )
    if capa_id is not None:
        scope |= Q(cuenta_bancaria_relacionada__capa_negocio_id=capa_id)
    return scope


def scoped_receivable_queryset(request):
    return CuentaPorCobrar.objects.filter(receivable_scope_filter(get_allowed_entity_ids(request)))


def scoped_payable_queryset(request):
    return CuentaPorPagar.objects.filter(payable_scope_filter(get_allowed_entity_ids(request)))


def scoped_payment_cxc_queryset(request):
    allowed_entity_ids = get_allowed_entity_ids(request)
    return PagoCuentaPorCobrar.objects.filter(
        Q(cuenta_por_cobrar__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            cuenta_por_cobrar__entidad_relacionada__isnull=True,
            cuenta_por_cobrar__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
    )


def scoped_payment_cxp_queryset(request):
    return PagoCuentaPorPagar.objects.filter(
        cuenta_por_pagar__entidad_relacionada_id__in=get_allowed_entity_ids(request)
    )


def scoped_event_queryset(request):
    return EventoFinanciero.objects.filter(
        event_scope_filter(get_allowed_entity_ids(request))
    ).distinct()


def event_mutation_queryset(request):
    return scoped_event_queryset(request).select_related(
        "entidad_relacionada",
        "cliente_relacionado",
        "caso_relacionado",
        "evidencia_pago_relacionada",
        "transaccion_relacionada",
    ).prefetch_related(
        Prefetch(
            "partidas",
            queryset=EventoFinancieroPartida.objects.select_related(
                "entidad_relacionada",
                "cliente_relacionado",
                "cuenta_por_cobrar_relacionada",
                "cuenta_por_cobrar_relacionada__entidad_relacionada",
                "cuenta_por_cobrar_relacionada__cliente_relacionado",
                "cuenta_por_cobrar_relacionada__cliente_relacionado__entidad_relacionada",
                "cuenta_por_pagar_relacionada",
                "cuenta_por_pagar_relacionada__entidad_relacionada",
            ).order_by("orden", "id"),
        )
    )


def event_detail_queryset(request):
    return event_mutation_queryset(request).prefetch_related(
        Prefetch(
            "pagos_cxc",
            queryset=PagoCuentaPorCobrar.objects.select_related(
                "cuenta_por_cobrar",
                "cuenta_por_cobrar__entidad_relacionada",
                "cuenta_por_cobrar__cliente_relacionado",
                "cuenta_por_cobrar__cliente_relacionado__entidad_relacionada",
            ).order_by("fecha_pago", "id"),
        ),
        Prefetch(
            "pagos_cxp",
            queryset=PagoCuentaPorPagar.objects.select_related(
                "cuenta_por_pagar",
                "cuenta_por_pagar__entidad_relacionada",
            ).order_by("fecha_pago", "id"),
        ),
    )


def event_reconciliation_queryset(request):
    return event_mutation_queryset(request).prefetch_related(
        Prefetch(
            "pagos_cxc",
            queryset=PagoCuentaPorCobrar.objects.select_related(
                "cuenta_por_cobrar",
            ).order_by("fecha_pago", "id"),
            to_attr="_prefetched_reconciliation_cxc_payments",
        ),
        Prefetch(
            "pagos_cxp",
            queryset=PagoCuentaPorPagar.objects.select_related(
                "cuenta_por_pagar",
            ).order_by("fecha_pago", "id"),
            to_attr="_prefetched_reconciliation_cxp_payments",
        ),
    )


def serialize_event_mutation_result(request, event: EventoFinanciero, include_detail: bool):
    if include_detail:
        detailed_event = get_object_or_404(event_detail_queryset(request), id=event.id)
        return serialize_event_detail(detailed_event)
    return {
        "id": event.id,
        "estatus": event.estatus,
        "transaccion_id": event.transaccion_relacionada_id,
    }


def scoped_transaction_queryset(request):
    current_capa = get_current_capa(request)
    return Transaccion.objects.filter(
        transaction_scope_filter(
            get_allowed_entity_ids(request),
            current_capa.id if current_capa else None,
        )
    ).distinct()


def get_scoped_event_or_404(request, event_id: int):
    return get_object_or_404(scoped_event_queryset(request), id=event_id)


def get_scoped_transaction_or_404(request, transaccion_id: int):
    return get_object_or_404(scoped_transaction_queryset(request), id=transaccion_id)


def transaction_auto_match_queryset(request):
    return scoped_transaction_queryset(request).select_related(
        "cuenta_bancaria_relacionada",
        "cuenta_bancaria_relacionada__entidad_relacionada",
        "cuenta_por_cobrar_relacionada",
        "cuenta_por_cobrar_relacionada__entidad_relacionada",
        "cuenta_por_cobrar_relacionada__cliente_relacionado",
        "cuenta_por_cobrar_relacionada__cliente_relacionado__entidad_relacionada",
        "cuenta_por_pagar_relacionada",
        "cuenta_por_pagar_relacionada__entidad_relacionada",
        "evento_relacionado",
        "evento_relacionado__entidad_relacionada",
        "evento_relacionado__cliente_relacionado",
        "evento_relacionado__cliente_relacionado__entidad_relacionada",
    )


def ensure_client_scope(request, cliente_id: int | None) -> None:
    if not cliente_id:
        return
    get_object_or_404(
        Cliente,
        id=cliente_id,
        entidad_relacionada_id__in=get_allowed_entity_ids(request),
    )


def ensure_scoped_ids(queryset, ids: set[int], error_message: str) -> None:
    if not ids:
        return
    found_ids = set(queryset.filter(id__in=ids).values_list("id", flat=True))
    if found_ids != ids:
        raise HttpError(404, error_message)


def ensure_financial_payload_scope(request, payload: "EventoFinancieroIn") -> None:
    allowed_entity_ids = get_allowed_entity_ids(request)
    entity_ids = {payload.entidad_id} if payload.entidad_id else set()
    client_ids = {payload.cliente_id} if payload.cliente_id else set()
    cxc_ids: set[int] = set()
    cxp_ids: set[int] = set()
    for item in payload.partidas:
        if item.entidad_id:
            entity_ids.add(item.entidad_id)
        if item.cliente_id:
            client_ids.add(item.cliente_id)
        if item.cuenta_cxc_id:
            cxc_ids.add(item.cuenta_cxc_id)
        if item.cuenta_cxp_id:
            cxp_ids.add(item.cuenta_cxp_id)

    ensure_scoped_ids(
        EntidadNegocio.objects.filter(id__in=allowed_entity_ids),
        entity_ids,
        "La entidad seleccionada no esta disponible.",
    )
    ensure_scoped_ids(
        Cliente.objects.filter(entidad_relacionada_id__in=allowed_entity_ids),
        client_ids,
        "El cliente seleccionado no esta disponible.",
    )
    ensure_scoped_ids(
        scoped_receivable_queryset(request),
        cxc_ids,
        "La cuenta por cobrar seleccionada no esta disponible.",
    )
    ensure_scoped_ids(
        scoped_payable_queryset(request),
        cxp_ids,
        "La cuenta por pagar seleccionada no esta disponible.",
    )
    if payload.transaccion_id:
        get_scoped_transaction_or_404(request, payload.transaccion_id)






def resolve_receivable_materialization_date(
    fecha_hasta: date | None = None,
) -> date | None:
    today = timezone.localdate()
    current_month_start = date(today.year, today.month, 1)
    if fecha_hasta and fecha_hasta < current_month_start:
        return None
    return date(
        today.year,
        today.month,
        monthrange(today.year, today.month)[1],
    )


def current_receivable_period_start() -> date:
    today = timezone.localdate()
    return date(today.year, today.month, 1)


class GenerarCargosIn(Schema):
    fecha_corte: Optional[date] = None
    entidad_id: Optional[int] = None


class PagoCxCIn(Schema):
    monto: Decimal
    fecha_pago: date
    metodo: str = "TRANSFERENCIA"
    canal_origen: str = "MANUAL"
    referencia: str = ""
    notas: str = ""
    cliente_id: Optional[int] = None
    cuenta_id: Optional[int] = None
    evidencia_id: Optional[int] = None


class EstadoCxCIn(Schema):
    estatus: str


class ProgramacionCxPIn(Schema):
    nombre: str
    categoria: str = "OTROS"
    naturaleza: str = "OPERATIVO"
    proveedor_nombre: str
    banco_pago: str = ""
    cuenta_pago: str = ""
    clabe_pago: str = ""
    prioridad: str = "MEDIA"
    periodicidad: str = "MENSUAL"
    fecha_inicio: date
    fecha_fin: Optional[date] = None
    dia_vencimiento: Optional[int] = None
    monto_base: Decimal
    dias_gracia: int = 0
    genera_recargo: bool = False
    prorrateable: bool = False
    activo: bool = True
    observaciones: Optional[str] = None


class GastoManualIn(Schema):
    entidad_id: Optional[int] = None
    proveedor_nombre: str
    categoria: str = "OTROS"
    naturaleza: str = "OPERATIVO"
    concepto: str
    banco_pago: str = ""
    cuenta_pago: str = ""
    clabe_pago: str = ""
    prioridad: str = "MEDIA"
    periodicidad: str = "UNICO"
    fecha_emision: date
    fecha_vencimiento: date
    fecha_periodo_inicio: Optional[date] = None
    fecha_periodo_fin: Optional[date] = None
    monto_proyectado: Decimal
    monto_real: Optional[Decimal] = None
    dias_gracia: int = 0
    genera_recargo: bool = False
    observaciones: Optional[str] = None


class PagoCxPIn(Schema):
    monto: Decimal
    fecha_pago: date
    metodo: str = "TRANSFERENCIA"
    canal_origen: str = "MANUAL"
    referencia: str = ""
    notas: str = ""
    evidencia_id: Optional[int] = None


class ValidarPagoIn(Schema):
    transaccion_id: Optional[int] = None
    notas_validacion: str = ""


class RechazarPagoIn(Schema):
    motivo: str = ""


class EventoPartidaIn(Schema):
    orden: int = 1
    tipo_destino: str = "CXC"
    entidad_id: Optional[int] = None
    cliente_id: Optional[int] = None
    cuenta_cxc_id: Optional[int] = None
    cuenta_cxp_id: Optional[int] = None
    concepto: str
    beneficiario: str = ""
    unidad_referencia: str = ""
    periodo_referencia: str = ""
    monto_partida: Decimal
    confianza: Optional[Decimal] = None
    requiere_revision_manual: bool = False
    estatus: str = "PROPUESTA"
    metadata: Optional[dict] = None
    observaciones: str = ""


class EventoFinancieroIn(Schema):
    caso_id: Optional[int] = None
    evidencia_id: Optional[int] = None
    entidad_id: Optional[int] = None
    cliente_id: Optional[int] = None
    transaccion_id: Optional[int] = None
    origen_evento: str = "MANUAL"
    tipo_movimiento: str = "INGRESO"
    unidad_detectada: str = ""
    beneficiario_principal: str = ""
    referencia_principal: str = ""
    fecha_evento: Optional[date] = None
    monto_total_reportado: Optional[Decimal] = None
    confianza_global: Optional[Decimal] = None
    requiere_revision_manual: bool = False
    texto_consolidado: str = ""
    raw_data: Optional[dict] = None
    propuesta_ia: Optional[dict] = None
    observaciones: str = ""
    estatus: str = "PROPUESTO"
    partidas: list[EventoPartidaIn] = []


class EventoEstatusIn(Schema):
    estatus: str
    observaciones: str = ""


class ConciliarEventoIn(Schema):
    transaccion_id: int
    notas: str = ""


class VincularEventoIn(Schema):
    evento_id: int
    notas: str = ""


class CuentaBancariaIn(Schema):
    nombre: str
    alias: str = ""
    banco: str = ""
    numero_cuenta: str = ""
    clabe: str = ""
    ultima_4: str = ""
    moneda: str = "MXN"
    activa: bool = True


class AplicarTransaccionCxCIn(Schema):
    entidad_id: Optional[int] = None
    cliente_id: Optional[int] = None
    cuenta_id: Optional[int] = None
    monto: Optional[Decimal] = None
    notas: str = ""
    sugerencia_id: str = ""
    sugerencia_score: Optional[int] = None
    sugerencia_confianza: str = ""
    sugerencia_estado: str = ""
    sugerencia_requiere_revision: bool = False
    sugerencia_razones: list[str] = []


class AplicarTransaccionCxPIn(Schema):
    cuenta_id: int
    monto: Optional[Decimal] = None
    notas: str = ""
    sugerencia_id: str = ""
    sugerencia_score: Optional[int] = None
    sugerencia_confianza: str = ""
    sugerencia_estado: str = ""
    sugerencia_requiere_revision: bool = False
    sugerencia_razones: list[str] = []


class EstadoTransaccionIn(Schema):
    estatus: str
    notas: str = ""


class AutoConciliarPendientesIn(Schema):
    cuenta_bancaria_id: Optional[int] = None
    carga_id: Optional[int] = None
    limite: int = 300


def validate_day(day: Optional[int]) -> Optional[int]:
    if day is None:
        return None
    if day < 1 or day > 31:
        raise HttpError(400, "El dia de vencimiento debe estar entre 1 y 31.")
    return day


def parse_date_or_raise(value: str, label: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{label} no tiene formato de fecha valido.") from exc


def load_rows_from_file(upload: UploadedFile) -> list[list[object]]:
    return load_rows_from_file_object(upload.file, upload.name or "")


def load_rows_from_file_object(file_object, file_name: str) -> list[list[object]]:
    extension = os.path.splitext(file_name or "")[1].lower()
    if extension in {".xlsx", ".xlsm"}:
        file_object.seek(0)
        workbook = load_workbook(file_object, data_only=True)
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    if extension == ".csv":
        file_object.seek(0)
        raw_content = file_object.read()
        content = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                content = raw_content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if content is None:
            raise HttpError(
                400,
                "No se pudo leer el CSV. Guarda el archivo como UTF-8 o Windows-1252.",
            )
        reader = csv.reader(io.StringIO(content))
        return [list(row) for row in reader]

    raise HttpError(400, "Solo se permiten archivos .xlsx, .xlsm, .csv o .pdf.")


CXP_HEADER_ALIASES = {
    "unidad": {"unidad", "unidad negocio", "unidad de negocio", "entidad"},
    "concepto": {"concepto", "gasto", "tipo gasto", "tipo de gasto"},
    "beneficiario": {"beneficiario", "proveedor", "a quien se paga", "pagar a"},
    "banco": {"banco", "banco pago", "banco de pago"},
    "cuenta": {"cuenta", "cuenta pago", "numero cuenta", "numero de cuenta"},
    "clabe": {"clabe", "clabe pago"},
    "fecha_vencimiento": {"fecha vencimiento", "vencimiento", "fecha de vencimiento"},
    "prioridad": {"prioridad", "semaforo", "semáforo"},
    "monto": {"monto", "importe", "total", "monto fijo"},
    "periodicidad": {"periodicidad", "frecuencia"},
    "categoria": {"categoria", "categoría", "tipo de gasto"},
    "naturaleza": {"naturaleza", "tipo operacion", "tipo operación"},
    "observaciones": {"observaciones", "notas", "comentarios"},
}


CXC_HEADER_ALIASES = {
    "unidad": {"unidad", "unidad negocio", "unidad de negocio", "entidad"},
    "cliente": {"cliente", "residente", "huesped", "inquilino", "propietario"},
    "concepto": {"concepto", "cargo", "descripcion", "detalle"},
    "monto": {"monto", "importe", "total", "monto total"},
    "fecha_vencimiento": {"fecha vencimiento", "vencimiento", "fecha de vencimiento"},
    "fecha_emision": {"fecha emision", "emision", "fecha de emision"},
    "periodo_inicio": {"periodo inicio", "fecha periodo inicio", "desde"},
    "periodo_fin": {"periodo fin", "fecha periodo fin", "hasta"},
    "origen": {"origen", "tipo", "tipo cargo", "tipo de cargo"},
    "periodicidad": {"periodicidad", "frecuencia"},
    "referencia": {"referencia", "referencia unica", "folio", "id externo"},
    "estatus": {"estatus", "estado"},
    "observaciones": {"observaciones", "notas", "comentarios"},
}


CXC_TEMPLATE_HEADERS = [
    "UNIDAD",
    "CLIENTE",
    "CONCEPTO",
    "MONTO",
    "FECHA VENCIMIENTO",
    "FECHA EMISION",
    "PERIODO INICIO",
    "PERIODO FIN",
    "ORIGEN",
    "PERIODICIDAD",
    "REFERENCIA",
    "ESTATUS",
    "OBSERVACIONES",
]

CXC_TEMPLATE_ROWS = [
    ["Edificio Reforma", "Cliente Ejemplo 01", "Renta Hab 01", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-001", "PENDIENTE", "Renta mensual habitacion ocupada"],
    ["Edificio Reforma", "Cliente Ejemplo 02", "Renta Hab 03", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-002", "PENDIENTE", "Renta mensual habitacion ocupada"],
    ["Edificio Reforma", "Cliente Ejemplo 03", "Renta Hab 04", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-003", "PENDIENTE", "Renta mensual habitacion ocupada"],
    ["Edificio Reforma", "Cliente Ejemplo 04", "Deposito Hab 02", "9200", "2026-05-03", "2026-05-01", "2026-05-01", "2026-05-31", "DEPOSITO", "UNICO", "CXC-PLANTILLA-004", "PENDIENTE", "Deposito por nueva asignacion"],
    ["Edificio Reforma", "Cliente Ejemplo 01", "Limpieza extraordinaria", "650", "2026-05-08", "2026-05-01", "2026-05-01", "2026-05-31", "AJUSTE", "UNICO", "CXC-PLANTILLA-005", "PENDIENTE", "Servicio adicional solicitado"],
    ["Torre Centro", "Cliente Ejemplo 05", "Renta Hab 01", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-006", "PENDIENTE", "Renta mensual suite"],
    ["Torre Centro", "Cliente Ejemplo 06", "Renta Hab 03", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-007", "PENDIENTE", "Renta mensual suite"],
    ["Torre Centro", "Cliente Ejemplo 07", "Renta Hab 04", "9200", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-008", "PENDIENTE", "Renta mensual suite"],
    ["Torre Centro", "Cliente Ejemplo 08", "Deposito Depto 01", "15500", "2026-05-04", "2026-05-01", "2026-05-01", "2026-05-31", "DEPOSITO", "UNICO", "CXC-PLANTILLA-009", "PENDIENTE", "Deposito por departamento"],
    ["Torre Centro", "Cliente Ejemplo 05", "Ajuste de servicios", "850", "2026-05-10", "2026-05-01", "2026-05-01", "2026-05-31", "AJUSTE", "UNICO", "CXC-PLANTILLA-010", "PENDIENTE", "Diferencia de servicios del periodo"],
    ["Condominio Norte", "Residente Ejemplo 01", "Cuota mantenimiento Mayo", "4500", "2026-05-10", "2026-05-01", "2026-05-01", "2026-05-31", "REGLA", "MENSUAL", "CXC-PLANTILLA-011", "PENDIENTE", "Cuota ordinaria mensual"],
    ["Condominio Norte", "Residente Ejemplo 02", "Cuota mantenimiento Mayo", "4500", "2026-05-10", "2026-05-01", "2026-05-01", "2026-05-31", "REGLA", "MENSUAL", "CXC-PLANTILLA-012", "PENDIENTE", "Cuota ordinaria mensual"],
    ["Condominio Norte", "Residente Ejemplo 03", "Cuota mantenimiento Mayo", "4500", "2026-05-10", "2026-05-01", "2026-05-01", "2026-05-31", "REGLA", "MENSUAL", "CXC-PLANTILLA-013", "PENDIENTE", "Cuota ordinaria mensual"],
    ["Condominio Norte", "Residente Ejemplo 04", "Reserva amenidades", "1200", "2026-05-12", "2026-05-01", "2026-05-01", "2026-05-31", "AJUSTE", "UNICO", "CXC-PLANTILLA-014", "PENDIENTE", "Uso de amenidad comun"],
    ["Condominio Norte", "Residente Ejemplo 01", "Recargo administrativo", "350", "2026-05-15", "2026-05-01", "2026-05-01", "2026-05-31", "AJUSTE", "UNICO", "CXC-PLANTILLA-015", "PENDIENTE", "Cargo por gestion administrativa"],
    ["Edificio Reforma", "Cliente Ejemplo 02", "Penalizacion salida tardia", "500", "2026-05-18", "2026-05-01", "2026-05-01", "2026-05-31", "AJUSTE", "UNICO", "CXC-PLANTILLA-016", "PENDIENTE", "Penalizacion operativa"],
    ["Torre Centro", "Cliente Ejemplo 07", "Estacionamiento mensual", "1100", "2026-05-08", "2026-05-01", "2026-05-01", "2026-05-31", "REGLA", "MENSUAL", "CXC-PLANTILLA-017", "PENDIENTE", "Servicio adicional mensual"],
    ["Condominio Norte", "Residente Ejemplo 02", "Agua areas comunes", "780", "2026-05-14", "2026-05-01", "2026-05-01", "2026-05-31", "REGLA", "MENSUAL", "CXC-PLANTILLA-018", "PENDIENTE", "Prorrateo de servicios comunes"],
    ["Edificio Reforma", "Cliente Ejemplo 04", "Renta Depto 01", "15500", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-019", "PENDIENTE", "Renta mensual departamento"],
    ["Torre Centro", "Cliente Ejemplo 08", "Renta Depto 01", "15500", "2026-05-05", "2026-05-01", "2026-05-01", "2026-05-31", "RENTA", "MENSUAL", "CXC-PLANTILLA-020", "PENDIENTE", "Renta mensual departamento"],
]

CXP_TEMPLATE_HEADERS = [
    "UNIDAD",
    "CONCEPTO",
    "BENEFICIARIO",
    "BANCO",
    "CUENTA",
    "CLABE",
    "FECHA VENCIMIENTO",
    "PRIORIDAD",
    "MONTO",
    "PERIODICIDAD",
    "CATEGORIA",
    "NATURALEZA",
    "OBSERVACIONES",
]

CXP_TEMPLATE_ROWS = [
    ["Edificio Reforma", "Luz areas comunes", "CFE", "BBVA", "0000000001", "012180000000000001", "2026-05-10", "ALTA", "18500", "MENSUAL", "LUZ", "OPERATIVO", "Variable: sustituir por monto real del mes antes de importar"],
    ["Edificio Reforma", "Internet administrativo", "Telmex", "BBVA", "0000000002", "012180000000000002", "2026-05-08", "MEDIA", "1250", "MENSUAL", "INTERNET", "OPERATIVO", "Servicio de internet del edificio"],
    ["Edificio Reforma", "Limpieza mensual", "Proveedor de limpieza", "SANTANDER", "0000000003", "014180000000000003", "2026-05-15", "ALTA", "9200", "MENSUAL", "LIMPIEZA", "OPERATIVO", "Servicio de limpieza mensual"],
    ["Edificio Reforma", "Recoleccion de basura", "Recolector local", "BANORTE", "0000000004", "072180000000000004", "2026-05-12", "MEDIA", "1650", "MENSUAL", "BASURA", "OPERATIVO", "Retiro de residuos"],
    ["Edificio Reforma", "Renta inmueble", "Arrendador", "BBVA", "0000000005", "012180000000000005", "2026-05-05", "CRITICA", "78000", "MENSUAL", "RENTA", "OPERATIVO", "Renta base del inmueble"],
    ["Torre Centro", "Agua areas comunes", "Organismo de agua", "BBVA", "0000000006", "012180000000000006", "2026-05-13", "ALTA", "7200", "MENSUAL", "AGUA", "OPERATIVO", "Variable: capturar monto real del recibo del periodo"],
    ["Torre Centro", "Gas", "Gas Natural Fenosa", "BANORTE", "0000000007", "072180000000000007", "2026-05-17", "MEDIA", "4300", "MENSUAL", "GAS", "OPERATIVO", "Suministro de gas"],
    ["Torre Centro", "Mantenimiento elevador", "Proveedor de elevadores", "SANTANDER", "0000000008", "014180000000000008", "2026-05-18", "MEDIA", "6800", "MENSUAL", "MANTENIMIENTO", "OPERATIVO", "Mantenimiento preventivo"],
    ["Torre Centro", "Contabilidad", "Despacho contable", "BBVA", "0000000009", "012180000000000009", "2026-05-20", "BAJA", "3500", "MENSUAL", "CONTABILIDAD", "ADMINISTRATIVO", "Honorarios contables"],
    ["Condominio Norte", "Seguridad", "Proveedor de seguridad", "BANORTE", "0000000010", "072180000000000010", "2026-05-06", "ALTA", "24500", "MENSUAL", "SEGURIDAD", "OPERATIVO", "Turnos de vigilancia"],
    ["Condominio Norte", "Jardineria", "Proveedor de jardineria", "BBVA", "0000000011", "012180000000000011", "2026-05-11", "MEDIA", "4800", "MENSUAL", "MANTENIMIENTO", "OPERATIVO", "Jardineria areas comunes"],
    ["Condominio Norte", "Reparacion bomba", "Proveedor hidraulico", "SANTANDER", "0000000012", "014180000000000012", "2026-05-16", "ALTA", "6900", "UNICO", "REPARACIONES", "OPERATIVO", "Reparacion correctiva"],
    ["Condominio Norte", "Impuesto predial", "Municipio", "BBVA", "0000000013", "012180000000000013", "2026-05-25", "ALTA", "12800", "ANUAL", "IMPUESTOS", "ADMINISTRATIVO", "Pago de impuesto programado"],
    ["Condominio Norte", "Nomina conserjeria", "Equipo operativo", "BBVA", "0000000014", "012180000000000014", "2026-05-30", "ALTA", "16200", "QUINCENAL", "NOMINA", "ADMINISTRATIVO", "Pago neto operativo"],
    ["Condominio Norte", "Otros insumos", "Proveedor general", "SANTANDER", "0000000015", "014180000000000015", "2026-05-19", "BAJA", "2100", "UNICO", "OTROS", "OPERATIVO", "Insumos operativos"],
]


def resolve_cxp_header_mapping(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows, start=1):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            normalized = normalize_cxp_key(value)
            if not normalized:
                continue
            for logical_name, aliases in CXP_HEADER_ALIASES.items():
                if normalized in aliases and logical_name not in mapping:
                    mapping[logical_name] = column_index
                    break
        if {"unidad", "concepto", "monto"}.issubset(mapping):
            return row_index, mapping
    raise ValueError("No se encontro un formato normalizado de cuentas por pagar.")


def resolve_cxc_header_mapping(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows, start=1):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            normalized = normalize_cxp_key(value)
            if not normalized:
                continue
            for logical_name, aliases in CXC_HEADER_ALIASES.items():
                if normalized in aliases and logical_name not in mapping:
                    mapping[logical_name] = column_index
                    break
        if {"unidad", "cliente", "concepto", "monto"}.issubset(mapping):
            return row_index, mapping
    raise ValueError("No se encontro un formato normalizado de cuentas por cobrar.")


def find_cxp_matrix(rows: list[list[object]]) -> tuple[int, int, int, list[tuple[int, str]]]:
    for row_index, row in enumerate(rows):
        normalized = [normalize_cxp_key(value) for value in row]
        if "concepto" not in normalized:
            continue
        concept_col = normalized.index("concepto")
        type_col = normalized.index("tipo de gasto") if "tipo de gasto" in normalized else max(concept_col - 1, 0)
        entity_columns: list[tuple[int, str]] = []
        for column_index in range(concept_col + 1, len(row)):
            header = normalize_cxp_text(row[column_index])
            key = normalize_cxp_key(header)
            if not header:
                continue
            if key in {"total general", "total", "gran total"}:
                break
            entity_columns.append((column_index, header))
        if len(entity_columns) >= 2:
            return row_index, type_col, concept_col, entity_columns
    raise ValueError("No se encontro una matriz de gastos con entidades en columnas.")


def build_entity_lookup(allowed_entity_ids: list[int]) -> dict[str, EntidadNegocio]:
    entidades = EntidadNegocio.objects.filter(id__in=allowed_entity_ids)
    lookup: dict[str, EntidadNegocio] = {}
    for entidad in entidades:
        lookup[normalize_cxp_key(entidad.nombre_comercial)] = entidad
        if entidad.razon_social:
            lookup[normalize_cxp_key(entidad.razon_social)] = entidad
    return lookup


def build_client_lookup(allowed_entity_ids: list[int]) -> dict[tuple[int, str], Cliente]:
    clientes = Cliente.objects.filter(entidad_relacionada_id__in=allowed_entity_ids)
    lookup: dict[tuple[int, str], Cliente] = {}
    for cliente in clientes:
        values = [
            cliente.razon_social,
            cliente.nombre_comercial,
            cliente.identificador,
            cliente.correo_principal,
            cliente.telefono,
        ]
        for value in values:
            key = normalize_cxp_key(value)
            if key:
                lookup.setdefault((cliente.entidad_relacionada_id, key), cliente)
    return lookup


def normalize_cxc_origin(value: object) -> str:
    text = normalize_cxp_key(value)
    aliases = {
        "renta": "RENTA",
        "deposito": "DEPOSITO",
        "deposito garantia": "DEPOSITO",
        "regla": "REGLA",
        "regla de negocio": "REGLA",
        "ajuste": "AJUSTE",
        "ajuste manual": "AJUSTE",
        "saldo a favor": "SALDO_A_FAVOR",
    }
    normalized = aliases.get(text, normalize_cxp_text(value).upper())
    return normalized if normalized in dict(CuentaPorCobrar.ORIGEN_CHOICES) else "AJUSTE"


def normalize_cxc_status(value: object) -> str:
    text = normalize_cxp_key(value)
    aliases = {
        "pendiente": "PENDIENTE",
        "parcial": "PARCIAL",
        "pago parcial": "PARCIAL",
        "vencido": "VENCIDO",
        "por conciliar": "POR_CONCILIAR",
        "conciliado": "CONCILIADO",
        "pagado": "CONCILIADO",
        "pagada": "CONCILIADO",
        "cancelado": "CANCELADO",
        "cancelada": "CANCELADO",
        "incobrable": "INCOBRABLE",
    }
    normalized = aliases.get(text, normalize_cxp_text(value).upper())
    return (
        normalized
        if normalized in dict(CuentaPorCobrar.ESTATUS_ADEUDO_CHOICES)
        else "PENDIENTE"
    )


def normalize_cxc_periodicity(value: object) -> str:
    normalized = normalize_cxp_text(value).upper()
    valid_periods = dict(PERIODICIDAD_CHOICES)
    return normalized if normalized in valid_periods else "MENSUAL"


def get_mapped_row_value(
    row: list[object],
    mapping: dict[str, int],
    key: str,
    default: object = "",
) -> object:
    if key not in mapping:
        return default
    column_index = mapping[key]
    if column_index >= len(row):
        return default
    return row[column_index]


def create_cxp_batch_item(
    *,
    entidad: EntidadNegocio,
    concepto: str,
    beneficiario: str,
    monto: Decimal,
    categoria: str,
    naturaleza: str,
    fecha_vencimiento: date,
    fecha_periodo_inicio: date,
    fecha_periodo_fin: date,
    referencia_unica: str,
    prioridad: str = "MEDIA",
    periodicidad: str = "MENSUAL",
    banco_pago: str = "",
    cuenta_pago: str = "",
    clabe_pago: str = "",
    observaciones: str = "",
) -> bool:
    _, created = CuentaPorPagar.objects.get_or_create(
        referencia_unica=referencia_unica,
        defaults={
            "entidad_relacionada": entidad,
            "proveedor_nombre": beneficiario or concepto,
            "banco_pago": banco_pago or None,
            "cuenta_pago": cuenta_pago or None,
            "clabe_pago": clabe_pago or None,
            "prioridad": prioridad if prioridad in {"BAJA", "MEDIA", "ALTA", "CRITICA"} else "MEDIA",
            "categoria": categoria,
            "naturaleza": naturaleza,
            "concepto": concepto,
            "periodicidad": normalize_cxc_periodicity(periodicidad),
            "tipo_registro": "BATCH",
            "fecha_periodo_inicio": fecha_periodo_inicio,
            "fecha_periodo_fin": fecha_periodo_fin,
            "fecha_programada": fecha_periodo_inicio,
            "fecha_emision": timezone.localdate(),
            "fecha_vencimiento": fecha_vencimiento,
            "monto_proyectado": monto,
            "monto_real": monto,
            "estatus": "PENDIENTE",
            "genera_recargo": categoria in {"RENTA", "LUZ", "AGUA", "CREDITO"},
            "observaciones": observaciones or None,
        },
    )
    return created


def import_cxp_normalized_rows(
    rows: list[list[object]],
    *,
    allowed_entity_ids: list[int],
    period_start: date,
    period_end: date,
    default_due_date: date,
) -> dict:
    header_row, mapping = resolve_cxp_header_mapping(rows)
    entity_lookup = build_entity_lookup(allowed_entity_ids)
    result = {"creados": 0, "existentes": 0, "omitidos": 0, "errores": []}

    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        if all(not normalize_cxp_text(value) for value in row):
            continue
        unidad = normalize_cxp_text(row[mapping["unidad"]] if mapping["unidad"] < len(row) else "")
        concepto = normalize_cxp_text(row[mapping["concepto"]] if mapping["concepto"] < len(row) else "")
        monto = parse_cxp_amount(row[mapping["monto"]] if mapping["monto"] < len(row) else None)
        entidad = entity_lookup.get(normalize_cxp_key(unidad))
        if not unidad or not concepto or not monto or monto <= 0 or not entidad:
            result["omitidos"] += 1
            if unidad and concepto and (not monto or monto <= 0):
                result["errores"].append(
                    f"Fila {row_number}: monto debe ser mayor a 0. "
                    "Si es un cargo variable, captura el monto real del mes antes de importar."
                )
            if unidad and concepto and monto and not entidad:
                result["errores"].append(f"Fila {row_number}: no se encontro la entidad '{unidad}'.")
            continue
        categoria = normalize_cxp_text(row[mapping["categoria"]]) if "categoria" in mapping and mapping["categoria"] < len(row) else ""
        categoria = categoria.upper() if categoria else infer_cxp_category(concepto)
        if categoria not in dict(ProgramacionCuentaPorPagar.CATEGORIA_CHOICES):
            categoria = infer_cxp_category(concepto, categoria)
        naturaleza = normalize_cxp_text(row[mapping["naturaleza"]]) if "naturaleza" in mapping and mapping["naturaleza"] < len(row) else ""
        naturaleza = naturaleza.upper() if naturaleza else infer_cxp_nature("", categoria)
        if naturaleza not in dict(ProgramacionCuentaPorPagar.NATURALEZA_CHOICES):
            naturaleza = infer_cxp_nature(naturaleza, categoria)
        beneficiario = normalize_cxp_text(row[mapping["beneficiario"]]) if "beneficiario" in mapping and mapping["beneficiario"] < len(row) else concepto
        due_date = parse_cxp_optional_date(row[mapping["fecha_vencimiento"]]) if "fecha_vencimiento" in mapping and mapping["fecha_vencimiento"] < len(row) else None
        prioridad_raw = row[mapping["prioridad"]] if "prioridad" in mapping and mapping["prioridad"] < len(row) else ""
        prioridad = normalize_cxp_priority(prioridad_raw, categoria, concepto)
        periodicidad = normalize_cxc_periodicity(
            get_mapped_row_value(row, mapping, "periodicidad")
        )
        banco = normalize_cxp_text(row[mapping["banco"]]) if "banco" in mapping and mapping["banco"] < len(row) else ""
        cuenta = normalize_cxp_text(row[mapping["cuenta"]]) if "cuenta" in mapping and mapping["cuenta"] < len(row) else ""
        clabe = normalize_cxp_text(row[mapping["clabe"]]) if "clabe" in mapping and mapping["clabe"] < len(row) else ""
        observaciones = normalize_cxp_text(row[mapping["observaciones"]]) if "observaciones" in mapping and mapping["observaciones"] < len(row) else ""
        ref = (
            f"CXP-BATCH:{period_start:%Y%m}:{entidad.id}:"
            f"{build_cxp_reference_token(concepto)}:{row_number}"
        )
        created = create_cxp_batch_item(
            entidad=entidad,
            concepto=concepto,
            beneficiario=beneficiario,
            monto=monto,
            categoria=categoria,
            naturaleza=naturaleza,
            fecha_vencimiento=due_date or default_due_date,
            fecha_periodo_inicio=period_start,
            fecha_periodo_fin=period_end,
            referencia_unica=ref,
            prioridad=prioridad,
            periodicidad=periodicidad,
            banco_pago=banco,
            cuenta_pago=cuenta,
            clabe_pago=clabe,
            observaciones=observaciones,
        )
        result["creados" if created else "existentes"] += 1
    return result


def create_cxc_batch_item(
    *,
    entidad: EntidadNegocio,
    cliente: Cliente,
    concepto: str,
    monto: Decimal,
    fecha_vencimiento: date,
    fecha_emision: date,
    fecha_periodo_inicio: date,
    fecha_periodo_fin: date,
    referencia_unica: str,
    origen: str = "AJUSTE",
    periodicidad: str = "MENSUAL",
    estatus: str = "PENDIENTE",
    observaciones: str = "",
) -> bool:
    _, created = CuentaPorCobrar.objects.get_or_create(
        entidad_relacionada=entidad,
        referencia_unica=referencia_unica[:180],
        defaults={
            "cliente_relacionado": cliente,
            "concepto": concepto,
            "origen": origen,
            "periodicidad": periodicidad,
            "fecha_periodo_inicio": fecha_periodo_inicio,
            "fecha_periodo_fin": fecha_periodo_fin,
            "fecha_programada": fecha_periodo_inicio,
            "fecha_emision": fecha_emision,
            "fecha_vencimiento": fecha_vencimiento,
            "monto_total": monto,
            "monto_pagado": Decimal("0"),
            "estatus_adeudo": estatus,
            "observaciones": observaciones or None,
        },
    )
    return created


def resolve_cxc_batch_due_date(
    *,
    cliente: Cliente,
    entidad: EntidadNegocio,
    period_start: date,
    default_due_date: date,
) -> date:
    capa = entidad.capa_negocio
    due_day = None
    if entidad.tipo_fecha_corte == "INDIVIDUAL":
        due_day = cliente.dia_corte_individual
    elif entidad.tipo_fecha_corte == "GENERAL":
        due_day = entidad.dia_corte_general

    if due_day is None and capa:
        if capa.fecha_vencimiento_modo == "INDIVIDUAL":
            due_day = cliente.dia_corte_individual
        elif capa.fecha_vencimiento_modo == "GLOBAL":
            due_day = capa.dia_vencimiento_default

    if due_day is None and capa:
        due_day = capa.dia_vencimiento_default

    if due_day:
        return resolve_due_date(period_start, due_day)
    return default_due_date


def import_cxc_normalized_rows(
    rows: list[list[object]],
    *,
    allowed_entity_ids: list[int],
    period_start: date,
    period_end: date,
    default_due_date: date,
) -> dict:
    header_row, mapping = resolve_cxc_header_mapping(rows)
    entity_lookup = build_entity_lookup(allowed_entity_ids)
    client_lookup = build_client_lookup(allowed_entity_ids)
    result = {"creados": 0, "existentes": 0, "omitidos": 0, "errores": []}

    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        if all(not normalize_cxp_text(value) for value in row):
            continue

        unidad = normalize_cxp_text(get_mapped_row_value(row, mapping, "unidad"))
        cliente_raw = normalize_cxp_text(get_mapped_row_value(row, mapping, "cliente"))
        concepto = normalize_cxp_text(get_mapped_row_value(row, mapping, "concepto"))
        monto = parse_cxp_amount(get_mapped_row_value(row, mapping, "monto"))
        entidad = entity_lookup.get(normalize_cxp_key(unidad))

        if not unidad or not cliente_raw or not concepto or not monto or monto <= 0:
            result["omitidos"] += 1
            result["errores"].append(
                f"Fila {row_number}: unidad, cliente, concepto y monto son obligatorios."
            )
            continue

        if not entidad:
            result["omitidos"] += 1
            result["errores"].append(f"Fila {row_number}: no se encontro la entidad '{unidad}'.")
            continue

        cliente = client_lookup.get((entidad.id, normalize_cxp_key(cliente_raw)))
        if not cliente:
            result["omitidos"] += 1
            result["errores"].append(
                f"Fila {row_number}: no se encontro el cliente '{cliente_raw}' en '{unidad}'."
            )
            continue

        row_period_start = (
            parse_cxp_optional_date(get_mapped_row_value(row, mapping, "periodo_inicio"))
            or period_start
        )
        row_period_end = (
            parse_cxp_optional_date(get_mapped_row_value(row, mapping, "periodo_fin"))
            or period_end
        )
        due_date = (
            parse_cxp_optional_date(get_mapped_row_value(row, mapping, "fecha_vencimiento"))
            or resolve_cxc_batch_due_date(
                cliente=cliente,
                entidad=entidad,
                period_start=row_period_start,
                default_due_date=default_due_date,
            )
        )
        fecha_emision = (
            parse_cxp_optional_date(get_mapped_row_value(row, mapping, "fecha_emision"))
            or timezone.localdate()
        )
        reference_raw = normalize_cxp_text(get_mapped_row_value(row, mapping, "referencia"))
        reference = reference_raw or (
            f"CXC-BATCH:{period_start:%Y%m}:{entidad.id}:{cliente.id}:"
            f"{build_cxp_reference_token(concepto)}:{row_number}"
        )
        observaciones = normalize_cxp_text(get_mapped_row_value(row, mapping, "observaciones"))

        created = create_cxc_batch_item(
            entidad=entidad,
            cliente=cliente,
            concepto=concepto,
            monto=monto,
            fecha_vencimiento=due_date,
            fecha_emision=fecha_emision,
            fecha_periodo_inicio=row_period_start,
            fecha_periodo_fin=row_period_end,
            referencia_unica=reference,
            origen=normalize_cxc_origin(get_mapped_row_value(row, mapping, "origen")),
            periodicidad=normalize_cxc_periodicity(
                get_mapped_row_value(row, mapping, "periodicidad")
            ),
            estatus=normalize_cxc_status(get_mapped_row_value(row, mapping, "estatus")),
            observaciones=observaciones,
        )
        result["creados" if created else "existentes"] += 1
    return result


def build_cxp_batch_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CxP Batch"
    descriptions = [
        "Unidad de negocio o edificio donde se registrara el compromiso.",
        "Concepto del gasto que vera el equipo operativo.",
        "Proveedor o beneficiario del pago. Usa datos ficticios si solo estas probando.",
        "Banco sugerido para el pago. Opcional.",
        "Numero de cuenta o referencia. Opcional.",
        "CLABE de 18 digitos. Opcional; BetterP valida formato cuando aplique.",
        "Fecha limite de pago. Si queda vacia se usa el ultimo dia del mes de carga.",
        "Selecciona BAJA, MEDIA, ALTA o CRITICA.",
        "Monto mayor a 0. BetterP omite o rechaza importes en cero.",
        "Selecciona UNICO, SEMANAL, QUINCENAL, MENSUAL, BIMESTRAL o ANUAL.",
        "Categoria del compromiso: servicios, impuestos, administracion u otros.",
        "Selecciona OPERATIVO o ADMINISTRATIVO.",
        "Notas internas para contacto, referencia o criterio de pago.",
    ]

    sheet.append(descriptions)
    sheet.append(CXP_TEMPLATE_HEADERS)
    for row in CXP_TEMPLATE_ROWS:
        sheet.append(row)

    description_fill = PatternFill("solid", fgColor="EAF7FF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    for cell in sheet[1]:
        cell.font = Font(italic=True, color="334155")
        cell.fill = description_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A3"
    sheet.row_dimensions[1].height = 72
    sheet.row_dimensions[2].height = 24

    catalog = workbook.create_sheet("Catalogos")
    catalog.append(["PRIORIDAD", "CATEGORIA", "NATURALEZA", "PERIODICIDAD"])
    for index, value in enumerate(["BAJA", "MEDIA", "ALTA", "CRITICA"], start=2):
        catalog.cell(row=index, column=1, value=value)
    for index, (value, _label) in enumerate(
        ProgramacionCuentaPorPagar.CATEGORIA_CHOICES,
        start=2,
    ):
        catalog.cell(row=index, column=2, value=value)
    for index, (value, _label) in enumerate(
        ProgramacionCuentaPorPagar.NATURALEZA_CHOICES,
        start=2,
    ):
        catalog.cell(row=index, column=3, value=value)
    for index, (value, _label) in enumerate(PERIODICIDAD_CHOICES, start=2):
        catalog.cell(row=index, column=4, value=value)
    catalog.sheet_state = "hidden"

    priority_end = 1 + len(["BAJA", "MEDIA", "ALTA", "CRITICA"])
    category_end = 1 + len(ProgramacionCuentaPorPagar.CATEGORIA_CHOICES)
    nature_end = 1 + len(ProgramacionCuentaPorPagar.NATURALEZA_CHOICES)
    periodicity_end = 1 + len(PERIODICIDAD_CHOICES)
    validations = [
        (
            DataValidation(type="list", formula1=f"=Catalogos!$A$2:$A${priority_end}", allow_blank=False),
            "H3:H1000",
        ),
        (
            DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=False),
            "I3:I1000",
        ),
        (
            DataValidation(type="list", formula1=f"=Catalogos!$D$2:$D${periodicity_end}", allow_blank=False),
            "J3:J1000",
        ),
        (
            DataValidation(type="list", formula1=f"=Catalogos!$B$2:$B${category_end}", allow_blank=False),
            "K3:K1000",
        ),
        (
            DataValidation(type="list", formula1=f"=Catalogos!$C$2:$C${nature_end}", allow_blank=False),
            "L3:L1000",
        ),
    ]
    for validation, cell_range in validations:
        sheet.add_data_validation(validation)
        validation.add(cell_range)

    widths = {
        "A": 24,
        "B": 28,
        "C": 26,
        "D": 18,
        "E": 18,
        "F": 22,
        "G": 20,
        "H": 16,
        "I": 14,
        "J": 18,
        "K": 20,
        "L": 20,
        "M": 44,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    instructions = workbook.create_sheet("Instrucciones")
    instructions.append(["Paso", "Detalle"])
    instructions.append(["1", "Usa una fila por compromiso de pago. Los ejemplos son ficticios."])
    instructions.append(["2", "UNIDAD debe coincidir con una entidad accesible de tu capa activa."])
    instructions.append(["3", "MONTO debe ser mayor a 0; en cargos variables captura el monto real del mes antes de importar."])
    instructions.append(["4", "PERIODICIDAD vive en el Excel para indicar si el compromiso es unico, semanal, quincenal, mensual, bimestral o anual."])
    instructions.append(["5", "Si FECHA VENCIMIENTO queda vacia, BetterP usa el ultimo dia del mes de carga."])
    instructions.append(["6", "PRIORIDAD, PERIODICIDAD, CATEGORIA y NATURALEZA tienen selectores para reducir errores de captura."])
    for column in range(1, 3):
        instructions.cell(row=1, column=column).font = Font(bold=True)
        instructions.column_dimensions[instructions.cell(row=1, column=column).column_letter].width = 34

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def import_cxp_matrix_rows(
    rows: list[list[object]],
    *,
    allowed_entity_ids: list[int],
    period_start: date,
    period_end: date,
    default_due_date: date,
) -> dict:
    header_index, type_col, concept_col, entity_columns = find_cxp_matrix(rows)
    entity_lookup = build_entity_lookup(allowed_entity_ids)
    result = {"creados": 0, "existentes": 0, "omitidos": 0, "errores": []}

    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        concepto = normalize_cxp_text(row[concept_col] if concept_col < len(row) else "")
        tipo_gasto = normalize_cxp_text(row[type_col] if type_col < len(row) else "")
        if not concepto or normalize_cxp_key(concepto) in {"total", "total general"}:
            continue
        categoria = infer_cxp_category(concepto, tipo_gasto)
        naturaleza = infer_cxp_nature(tipo_gasto, categoria)
        prioridad = infer_cxp_priority(categoria, concepto)
        for column_index, entity_name in entity_columns:
            monto = parse_cxp_amount(row[column_index] if column_index < len(row) else None)
            if not monto or monto <= 0:
                continue
            entidad = entity_lookup.get(normalize_cxp_key(entity_name))
            if not entidad:
                result["omitidos"] += 1
                error = f"Columna {column_index + 1}: no se encontro la entidad '{entity_name}'."
                if error not in result["errores"]:
                    result["errores"].append(error)
                continue
            ref = (
                f"CXP-BATCH:{period_start:%Y%m}:{entidad.id}:"
                f"{build_cxp_reference_token(concepto)}:{row_number}:{column_index + 1}"
            )
            created = create_cxp_batch_item(
                entidad=entidad,
                concepto=concepto,
                beneficiario=concepto,
                monto=monto,
                categoria=categoria,
                naturaleza=naturaleza,
                fecha_vencimiento=default_due_date,
                fecha_periodo_inicio=period_start,
                fecha_periodo_fin=period_end,
                referencia_unica=ref,
                prioridad=prioridad,
                observaciones=f"Importado desde matriz. Tipo: {tipo_gasto}".strip(),
            )
            result["creados" if created else "existentes"] += 1
    return result


def parse_uploaded_bank_file(upload: UploadedFile) -> ParsedBankStatement:
    return parse_bank_file_object(upload.file, upload.name or "")


def parse_bank_file_object(file_object, file_name: str) -> ParsedBankStatement:
    extension = os.path.splitext(file_name or "")[1].lower()
    if extension == ".pdf":
        try:
            return parse_bank_statement_pdf(fileobj=file_object, filename=file_name or "")
        except BankStatementParseError as exc:
            raise HttpError(400, str(exc)) from exc
        finally:
            file_object.seek(0)

    rows = load_rows_from_file_object(file_object, file_name)
    return ParsedBankStatement(movimientos=build_bank_movements(rows))


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def resolve_bank_headers(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    aliases = {
        "fecha_pago": {
            "fecha",
            "fecha operacion",
            "fecha movimiento",
            "fecha pago",
        },
        "monto": {"monto", "importe", "abono", "cargo"},
        "numero_referencia": {
            "referencia",
            "folio",
            "clave rastreo",
            "numero referencia",
            "referencia numerica",
        },
        "folio_bancario": {"folio bancario", "folio movimiento", "folio operacion"},
        "concepto_bancario": {"concepto", "descripcion", "detalle", "concepto bancario"},
        "tipo_movimiento": {"tipo", "movimiento", "naturaleza", "tipo movimiento"},
        "saldo_resultante": {"saldo", "saldo final", "saldo movimiento", "saldo resultante"},
    }

    for index, row in enumerate(rows, start=1):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            normalized = normalize_header(value)
            if not normalized:
                continue
            for logical_name, candidates in aliases.items():
                if normalized in candidates and logical_name not in mapping:
                    mapping[logical_name] = column_index
                    break
        if {"fecha_pago", "monto"}.issubset(mapping.keys()):
            return index, mapping

    raise HttpError(
        400,
        "No se encontraron encabezados validos. El archivo debe incluir Fecha y Monto.",
    )


def parse_bank_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    if not text:
        raise ValueError("La fecha del movimiento es obligatoria.")

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"No se pudo interpretar la fecha '{text}'.")


def parse_bank_amount(value: object) -> Decimal:
    text = str(value or "").strip()
    if not text:
        raise ValueError("El monto del movimiento es obligatorio.")

    cleaned = text.replace("$", "").replace(",", "").replace("MXN", "").strip()
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"No se pudo interpretar el monto '{text}'.") from exc


def parse_optional_request_date(raw_value: object) -> Optional[date]:
    text = str(raw_value or "").strip()
    if not text:
        return None
    return parse_bank_date(text)


def parse_optional_request_decimal(raw_value: object) -> Optional[Decimal]:
    text = str(raw_value or "").strip()
    if not text:
        return None
    return parse_bank_amount(text)


def normalize_cxp_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal) and value == value.to_integral():
        return str(value.to_integral())
    return str(value).strip()


def normalize_cxp_key(value: object) -> str:
    text = normalize_cxp_text(value).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def build_cxp_reference_token(value: object) -> str:
    normalized = normalize_cxp_key(value).replace(" ", "-")
    return (normalized or "gasto")[:60]


def parse_cxp_amount(value: object) -> Optional[Decimal]:
    text = normalize_cxp_text(value)
    if not text:
        return None
    cleaned = (
        text.replace("$", "")
        .replace(",", "")
        .replace("MXN", "")
        .replace("mxn", "")
        .strip()
    )
    if cleaned in {"-", "0", "0.0", "0.00"}:
        return Decimal("0")
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def parse_cxp_optional_date(value: object) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_cxp_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def resolve_cxp_period(periodo_raw: object) -> tuple[date, date]:
    text = normalize_cxp_text(periodo_raw)
    if text:
        for fmt in ("%Y-%m", "%m/%Y", "%m-%Y"):
            try:
                parsed = datetime.strptime(text, fmt)
                last_day = monthrange(parsed.year, parsed.month)[1]
                return date(parsed.year, parsed.month, 1), date(parsed.year, parsed.month, last_day)
            except ValueError:
                continue
    today = timezone.localdate()
    last_day = monthrange(today.year, today.month)[1]
    return date(today.year, today.month, 1), date(today.year, today.month, last_day)


def infer_cxp_category(concepto: object, tipo_gasto: object = "") -> str:
    text = f"{normalize_cxp_key(tipo_gasto)} {normalize_cxp_key(concepto)}"
    if "nomina" in text:
        return "NOMINA"
    if "renta" in text:
        return "RENTA"
    if "impuesto" in text:
        return "IMPUESTOS"
    if "agua" in text:
        return "AGUA"
    if "luz" in text or "cfe" in text:
        return "LUZ"
    if "gas" in text:
        return "GAS"
    if "internet" in text or "telmex" in text or "total play" in text or "izzi" in text:
        return "INTERNET"
    if "basura" in text or "recoleccion" in text:
        return "BASURA"
    if "limpieza" in text:
        return "LIMPIEZA"
    if "reparacion" in text:
        return "REPARACIONES"
    if "mantenimiento" in text:
        return "MANTENIMIENTO"
    if "credito" in text or "tarjeta" in text or "xepelin" in text:
        return "CREDITO"
    if "publicidad" in text:
        return "PUBLICIDAD"
    if "contador" in text or "contabilidad" in text:
        return "CONTABILIDAD"
    return "OTROS"


def infer_cxp_priority(categoria: str, concepto: object = "") -> str:
    text = normalize_cxp_key(concepto)
    if categoria in {"RENTA", "LUZ", "AGUA", "NOMINA", "CREDITO"}:
        return "ALTA"
    if "urgente" in text or "critico" in text:
        return "CRITICA"
    if categoria in {"INTERNET", "GAS", "BASURA", "REPARACIONES", "MANTENIMIENTO"}:
        return "MEDIA"
    return "BAJA"


def normalize_cxp_priority(value: object, categoria: str, concepto: object = "") -> str:
    text = normalize_cxp_key(value)
    aliases = {
        "baja": "BAJA",
        "verde": "BAJA",
        "media": "MEDIA",
        "normal": "MEDIA",
        "amarillo": "MEDIA",
        "alta": "ALTA",
        "naranja": "ALTA",
        "critica": "CRITICA",
        "critico": "CRITICA",
        "rojo": "CRITICA",
    }
    return aliases.get(text, infer_cxp_priority(categoria, concepto))


def infer_cxp_nature(tipo_gasto: object, categoria: str) -> str:
    text = normalize_cxp_key(tipo_gasto)
    if "operacion" in text or categoria in {"RENTA", "AGUA", "LUZ", "GAS", "INTERNET"}:
        return "OPERATIVO"
    return "ADMINISTRATIVO"


def parse_bank_movement_type(raw_type: object, amount: Decimal) -> str:
    normalized = normalize_header(raw_type)
    if normalized in {"ingreso", "abono", "credito"}:
        return "INGRESO"
    if normalized in {"egreso", "cargo", "debito"}:
        return "EGRESO"
    return "EGRESO" if amount < 0 else "INGRESO"


def build_bank_movements(rows: list[list[object]]) -> list[dict[str, object]]:
    header_row, mapping = resolve_bank_headers(rows)
    movimientos: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        if not any(value not in (None, "") for value in row):
            continue
        try:
            raw_amount = parse_bank_amount(row[mapping["monto"]])
            amount = abs(raw_amount)
            movimientos.append(
                {
                    "fecha_pago": parse_bank_date(row[mapping["fecha_pago"]]),
                    "monto": amount,
                    "tipo_movimiento": parse_bank_movement_type(
                        row[mapping.get("tipo_movimiento", -1)]
                        if "tipo_movimiento" in mapping
                        else "",
                        raw_amount,
                    ),
                    "numero_referencia": (
                        str(row[mapping["numero_referencia"]]).strip()
                        if "numero_referencia" in mapping and row[mapping["numero_referencia"]] is not None
                        else None
                    ),
                    "folio_bancario": (
                        str(row[mapping["folio_bancario"]]).strip()
                        if "folio_bancario" in mapping and row[mapping["folio_bancario"]] is not None
                        else None
                    ),
                    "concepto_bancario": (
                        str(row[mapping["concepto_bancario"]]).strip()
                        if "concepto_bancario" in mapping and row[mapping["concepto_bancario"]] is not None
                        else None
                    ),
                    "saldo_resultante": (
                        abs(parse_bank_amount(row[mapping["saldo_resultante"]]))
                        if "saldo_resultante" in mapping and row[mapping["saldo_resultante"]] not in (None, "")
                        else None
                    ),
                    "metadata": {"row_number": row_number},
                }
            )
        except (IndexError, ValueError) as exc:
            raise HttpError(400, f"Fila {row_number}: {exc}")
    return movimientos


def build_default_bank_account_name(parsed_statement: ParsedBankStatement) -> str:
    suffix = parsed_statement.ultima_4 or ""
    if parsed_statement.banco and suffix:
        return f"{parsed_statement.banco} ****{suffix}"
    if parsed_statement.banco:
        return f"{parsed_statement.banco} principal"
    if suffix:
        return f"Cuenta bancaria ****{suffix}"
    return "Cuenta bancaria"


def build_account_creation_suggestion(
    *,
    parsed_statement: ParsedBankStatement,
    capa=None,
    account_name: str | None = None,
    alias: str | None = None,
) -> dict[str, object]:
    return {
        "capa_id": capa.id if capa is not None else None,
        "capa_nombre": getattr(capa, "nombre", None),
        "titular_detectado": parsed_statement.titular,
        "nombre_cuenta": account_name or build_default_bank_account_name(parsed_statement),
        "alias_cuenta": alias,
        "banco": parsed_statement.banco,
        "numero_cuenta": parsed_statement.numero_cuenta,
        "clabe": parsed_statement.clabe,
        "ultima_4": parsed_statement.ultima_4,
    }


def normalize_legal_name(value: str | None) -> str:
    if not value:
        return ""
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = re.sub(r"[^A-Z0-9 ]+", " ", normalized.upper())
    normalized = re.sub(
        r"\b(SA|SAB|SAPI|CV|DE|RL|SC|AC|SPR|RLMI|DECV|S DE RL DE CV|SA DE CV)\b",
        " ",
        normalized,
    )
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def suggest_statement_entity(
    parsed_statement: ParsedBankStatement,
    allowed_entity_ids: list[int],
) -> Optional[EntidadNegocio]:
    holder_name = normalize_legal_name(parsed_statement.titular)
    if not holder_name:
        return None

    candidates = list(
        EntidadNegocio.objects.filter(id__in=allowed_entity_ids).only(
            "id",
            "nombre_comercial",
            "razon_social",
        )
    )

    matches: list[EntidadNegocio] = []
    for entity in candidates:
        legal_name = normalize_legal_name(entity.razon_social)
        trade_name = normalize_legal_name(entity.nombre_comercial)
        if legal_name and (holder_name == legal_name or holder_name in legal_name or legal_name in holder_name):
            matches.append(entity)
            continue
        if trade_name and (holder_name == trade_name or holder_name in trade_name or trade_name in holder_name):
            matches.append(entity)

    if len(matches) == 1:
        return matches[0]
    return None


def get_active_business_layer(request):
    capa = get_current_capa(request)
    if capa is None:
        raise HttpError(400, "No pudimos identificar la capa de negocio activa para esta conciliación.")
    return capa


def normalize_bank_identifier(value: str | None, *, digits_only: bool = True) -> str:
    text = str(value or "").strip()
    if digits_only:
        return re.sub(r"\D+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def find_duplicate_bank_account(
    *,
    capa_id: int,
    numero_cuenta: str = "",
    clabe: str = "",
    ultima_4: str = "",
    banco: str = "",
    exclude_id: int | None = None,
) -> CuentaBancaria | None:
    normalized_clabe = normalize_bank_identifier(clabe)
    normalized_account = normalize_bank_identifier(numero_cuenta)
    normalized_last4 = normalize_bank_identifier(ultima_4)
    normalized_bank = normalize_bank_identifier(banco, digits_only=False).upper()

    queryset = CuentaBancaria.objects.filter(capa_negocio_id=capa_id)
    if exclude_id is not None:
        queryset = queryset.exclude(id=exclude_id)

    if normalized_clabe:
        existing = queryset.filter(clabe=normalized_clabe).first()
        if existing is not None:
            return existing

    if normalized_account:
        existing = queryset.filter(numero_cuenta=normalized_account).first()
        if existing is not None:
            return existing

    if normalized_bank and normalized_last4:
        existing = queryset.filter(
            banco__iexact=normalized_bank,
            ultima_4=normalized_last4,
        ).first()
        if existing is not None:
            return existing

    return None


def find_matching_bank_account(
    parsed_statement: ParsedBankStatement,
    capa_id: int,
) -> Optional[CuentaBancaria]:
    account_filter = Q()
    if parsed_statement.clabe:
        account_filter |= Q(clabe=parsed_statement.clabe)
    if parsed_statement.numero_cuenta:
        account_filter |= Q(numero_cuenta=parsed_statement.numero_cuenta)
    if parsed_statement.ultima_4:
        account_filter |= Q(ultima_4=parsed_statement.ultima_4)

    if not account_filter:
        return None

    matches = list(
        CuentaBancaria.objects.filter(capa_negocio_id=capa_id, activa=True)
        .filter(account_filter)[:2]
    )
    if len(matches) == 1:
        return matches[0]
    return None








@router.get("/cxc/batch/plantilla/")
def descargar_plantilla_cxc(request):
    require_cxc_access(request)
    require_batch_import_access(request)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="plantilla-cuentas-por-cobrar.csv"'
    writer = csv.writer(response)
    writer.writerow(CXC_TEMPLATE_HEADERS)
    writer.writerows(CXC_TEMPLATE_ROWS)
    return response


@router.post("/cxc/batch/importar/")
def importar_cxc_batch(request, file: UploadedFile = File(...), async_job: bool = False):
    require_cxc_access(request)
    require_batch_import_access(request)
    require_write_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    period_start, period_end = resolve_cxp_period(request.POST.get("periodo"))
    default_due_date = parse_cxp_optional_date(request.POST.get("fecha_vencimiento")) or period_end
    if async_job:
        upload_info = store_background_upload(file, prefix="finanzas-cxc")
        job = enqueue_background_job(
            kind="finanzas.cxc_batch_import",
            payload={
                "file_path": upload_info["path"],
                "file_name": upload_info["original_name"],
                "allowed_entity_ids": allowed_entity_ids,
                "period_start": period_start.isoformat(),
                "period_end": period_end.isoformat(),
                "default_due_date": default_due_date.isoformat(),
            },
            created_by=getattr(getattr(request, "auth", None), "user", None),
            max_attempts=1,
        )
        return {
            "accepted": True,
            "job_id": job.id,
            "status": job.status,
            "mensaje": "Carga batch de CxC enviada a procesamiento en background.",
        }

    return import_cxc_batch_file(
        file.file,
        file.name or "",
        allowed_entity_ids=allowed_entity_ids,
        period_start=period_start,
        period_end=period_end,
        default_due_date=default_due_date,
    )


def import_cxc_batch_from_storage(payload: dict[str, object]) -> dict[str, object]:
    file_path = str(payload.get("file_path") or "")
    file_name = str(payload.get("file_name") or file_path)
    allowed_entity_ids = [
        int(entity_id)
        for entity_id in payload.get("allowed_entity_ids") or []
        if str(entity_id).strip()
    ]
    if not file_path:
        raise ValueError("No se encontro el archivo de importacion.")
    try:
        with default_storage.open(file_path, "rb") as stored_file:
            return import_cxc_batch_file(
                stored_file,
                file_name,
                allowed_entity_ids=allowed_entity_ids,
                period_start=parse_date_or_raise(str(payload["period_start"]), "periodo_inicio"),
                period_end=parse_date_or_raise(str(payload["period_end"]), "periodo_fin"),
                default_due_date=parse_date_or_raise(
                    str(payload["default_due_date"]),
                    "fecha_vencimiento",
                ),
            )
    finally:
        if default_storage.exists(file_path):
            default_storage.delete(file_path)


def import_cxc_batch_file(
    file_object,
    file_name: str,
    *,
    allowed_entity_ids: list[int],
    period_start: date,
    period_end: date,
    default_due_date: date,
) -> dict[str, object]:
    extension = os.path.splitext(file_name or "")[1].lower()

    try:
        with transaction.atomic():
            if extension in {".xlsx", ".xlsm"}:
                file_object.seek(0)
                workbook = load_workbook(file_object, data_only=True)
                sheets = [
                    (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                    for sheet in workbook.worksheets
                ]
            elif extension == ".csv":
                sheets = [(file_name or "CSV", load_rows_from_file_object(file_object, file_name))]
            else:
                raise HttpError(400, "Solo se permiten archivos .xlsx, .xlsm o .csv.")

            last_error = ""
            for sheet_name, rows in sheets:
                try:
                    result = import_cxc_normalized_rows(
                        rows,
                        allowed_entity_ids=allowed_entity_ids,
                        period_start=period_start,
                        period_end=period_end,
                        default_due_date=default_due_date,
                    )
                    return {
                        "mensaje": "Carga batch de cuentas por cobrar procesada.",
                        "formato": "normalizado",
                        "hoja": sheet_name,
                        "periodo_inicio": period_start,
                        "periodo_fin": period_end,
                        **result,
                    }
                except ValueError as exc:
                    last_error = str(exc)
                    continue
            raise HttpError(400, last_error or "No se pudo interpretar el archivo de CxC.")
    except HttpError:
        raise
    except Exception as exc:
        raise HttpError(500, f"No se pudo importar la carga batch de CxC: {exc}")




@router.get("/cxp/")
def listar_cartera_cxp(
    request,
    entidad_id: Optional[int] = None,
    status: Optional[str] = None,
    categoria: Optional[str] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    busqueda: str = "",
    page: int = 1,
    page_size: int = 20,
    include_items: bool = True,
):
    require_finance_module(request, "cxp")
    allowed_entity_ids = get_allowed_entity_ids(request)
    if entidad_id and entidad_id not in allowed_entity_ids:
        raise HttpError(404, "La entidad solicitada no pertenece a tu capa activa.")
    return build_global_cxp_dashboard(
        entidad_id=entidad_id,
        status=status,
        categoria=categoria,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        busqueda=busqueda or "",
        page=page,
        page_size=page_size,
        allowed_entity_ids=allowed_entity_ids,
        include_items=include_items,
    )


@router.get("/cxp/batch/plantilla/")
def descargar_plantilla_cxp(request):
    require_finance_module(request, "cxp")
    require_batch_import_access(request)
    content = build_cxp_batch_template()
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="plantilla-cuentas-por-pagar.xlsx"'
    return response


@router.post("/cxp/batch/importar/")
def importar_cxp_batch(request, file: UploadedFile = File(...), async_job: bool = False):
    require_finance_module(request, "cxp")
    require_batch_import_access(request)
    require_write_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    period_start, period_end = resolve_cxp_period(request.POST.get("periodo"))
    default_due_date = parse_cxp_optional_date(request.POST.get("fecha_vencimiento")) or period_end
    requested_format = normalize_cxp_key(request.POST.get("formato") or "auto")
    if async_job:
        upload_info = store_background_upload(file, prefix="finanzas-cxp")
        job = enqueue_background_job(
            kind="finanzas.cxp_batch_import",
            payload={
                "file_path": upload_info["path"],
                "file_name": upload_info["original_name"],
                "allowed_entity_ids": allowed_entity_ids,
                "period_start": period_start.isoformat(),
                "period_end": period_end.isoformat(),
                "default_due_date": default_due_date.isoformat(),
                "requested_format": requested_format,
            },
            created_by=getattr(getattr(request, "auth", None), "user", None),
            max_attempts=1,
        )
        return {
            "accepted": True,
            "job_id": job.id,
            "status": job.status,
            "mensaje": "Carga batch de CxP enviada a procesamiento en background.",
        }

    return import_cxp_batch_file(
        file.file,
        file.name or "",
        allowed_entity_ids=allowed_entity_ids,
        period_start=period_start,
        period_end=period_end,
        default_due_date=default_due_date,
        requested_format=requested_format,
    )


def import_cxp_batch_from_storage(payload: dict[str, object]) -> dict[str, object]:
    file_path = str(payload.get("file_path") or "")
    file_name = str(payload.get("file_name") or file_path)
    allowed_entity_ids = [
        int(entity_id)
        for entity_id in payload.get("allowed_entity_ids") or []
        if str(entity_id).strip()
    ]
    if not file_path:
        raise ValueError("No se encontro el archivo de importacion.")
    try:
        with default_storage.open(file_path, "rb") as stored_file:
            return import_cxp_batch_file(
                stored_file,
                file_name,
                allowed_entity_ids=allowed_entity_ids,
                period_start=parse_date_or_raise(str(payload["period_start"]), "periodo_inicio"),
                period_end=parse_date_or_raise(str(payload["period_end"]), "periodo_fin"),
                default_due_date=parse_date_or_raise(
                    str(payload["default_due_date"]),
                    "fecha_vencimiento",
                ),
                requested_format=normalize_cxp_key(
                    str(payload.get("requested_format") or "auto")
                ),
            )
    finally:
        if default_storage.exists(file_path):
            default_storage.delete(file_path)


def import_cxp_batch_file(
    file_object,
    file_name: str,
    *,
    allowed_entity_ids: list[int],
    period_start: date,
    period_end: date,
    default_due_date: date,
    requested_format: str,
) -> dict[str, object]:
    extension = os.path.splitext(file_name or "")[1].lower()

    try:
        with transaction.atomic():
            if extension in {".xlsx", ".xlsm"}:
                file_object.seek(0)
                workbook = load_workbook(file_object, data_only=True)
                sheets = [
                    (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                    for sheet in workbook.worksheets
                ]
            elif extension == ".csv":
                sheets = [(file_name or "CSV", load_rows_from_file_object(file_object, file_name))]
            else:
                raise HttpError(400, "Solo se permiten archivos .xlsx, .xlsm o .csv.")

            last_error = ""
            strategies = (
                ["normalizado", "matriz"]
                if requested_format == "auto"
                else [requested_format]
            )
            for strategy in strategies:
                for sheet_name, rows in sheets:
                    try:
                        if strategy == "normalizado":
                            result = import_cxp_normalized_rows(
                                rows,
                                allowed_entity_ids=allowed_entity_ids,
                                period_start=period_start,
                                period_end=period_end,
                                default_due_date=default_due_date,
                            )
                        elif strategy == "matriz":
                            result = import_cxp_matrix_rows(
                                rows,
                                allowed_entity_ids=allowed_entity_ids,
                                period_start=period_start,
                                period_end=period_end,
                                default_due_date=default_due_date,
                            )
                        else:
                            raise HttpError(400, "Formato no soportado para CxP batch.")
                        return {
                            "mensaje": "Carga batch de cuentas por pagar procesada.",
                            "formato": strategy,
                            "hoja": sheet_name,
                            "periodo_inicio": period_start,
                            "periodo_fin": period_end,
                            **result,
                        }
                    except ValueError as exc:
                        last_error = str(exc)
                        continue
            raise HttpError(400, last_error or "No se pudo interpretar el archivo de CxP.")
    except HttpError:
        raise
    except Exception as exc:
        raise HttpError(500, f"No se pudo importar la carga batch de CxP: {exc}")


@router.post("/cxp/manual/")
def crear_gasto_manual_global(request, payload: GastoManualIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    if not payload.entidad_id:
        raise HttpError(400, "Selecciona la entidad del gasto.")
    entidad = get_object_or_404(
        EntidadNegocio,
        id=payload.entidad_id,
        id__in=get_allowed_entity_ids(request),
    )
    if payload.fecha_vencimiento < payload.fecha_emision:
        raise HttpError(400, "La fecha de vencimiento no puede ser anterior a la emision.")
    cuenta = CuentaPorPagar.objects.create(
        entidad_relacionada=entidad,
        proveedor_nombre=payload.proveedor_nombre.strip(),
        banco_pago=payload.banco_pago.strip() or None,
        cuenta_pago=payload.cuenta_pago.strip() or None,
        clabe_pago=payload.clabe_pago.strip() or None,
        prioridad=payload.prioridad if payload.prioridad in {"BAJA", "MEDIA", "ALTA", "CRITICA"} else "MEDIA",
        categoria=payload.categoria,
        naturaleza=payload.naturaleza,
        concepto=payload.concepto.strip(),
        periodicidad=payload.periodicidad,
        tipo_registro="MANUAL",
        fecha_periodo_inicio=payload.fecha_periodo_inicio,
        fecha_periodo_fin=payload.fecha_periodo_fin,
        fecha_programada=payload.fecha_emision,
        fecha_emision=payload.fecha_emision,
        fecha_vencimiento=payload.fecha_vencimiento,
        monto_proyectado=payload.monto_proyectado,
        monto_real=payload.monto_real,
        dias_gracia=max(payload.dias_gracia or 0, 0),
        genera_recargo=payload.genera_recargo,
        estatus="PENDIENTE",
        observaciones=(payload.observaciones or "").strip() or None,
    )
    return {"id": cuenta.id, "mensaje": "Gasto manual creado correctamente."}


@router.get("/cxc/pagos/")
def listar_pagos_cxc(
    request,
    entidad_id: Optional[int] = None,
    cliente_id: Optional[int] = None,
    estatus_validacion: Optional[str] = None,
    page: int = 1,
    page_size: int = 100,
):
    require_cxc_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    queryset = PagoCuentaPorCobrar.objects.select_related(
        "cuenta_por_cobrar",
        "cuenta_por_cobrar__cliente_relacionado",
        "cuenta_por_cobrar__entidad_relacionada",
        "cuenta_por_cobrar__cliente_relacionado__entidad_relacionada",
    ).only(
        "id",
        "cuenta_por_cobrar_id",
        "monto",
        "fecha_pago",
        "metodo",
        "canal_origen",
        "estatus_validacion",
        "fecha_validacion",
        "referencia",
        "transaccion_relacionada_id",
        "evidencia_pago_relacionada_id",
        "evento_financiero_relacionado_id",
        "cuenta_por_cobrar__id",
        "cuenta_por_cobrar__concepto",
        "cuenta_por_cobrar__entidad_relacionada_id",
        "cuenta_por_cobrar__cliente_relacionado_id",
        "cuenta_por_cobrar__entidad_relacionada__id",
        "cuenta_por_cobrar__entidad_relacionada__nombre_comercial",
        "cuenta_por_cobrar__cliente_relacionado__id",
        "cuenta_por_cobrar__cliente_relacionado__razon_social",
        "cuenta_por_cobrar__cliente_relacionado__nombre_comercial",
        "cuenta_por_cobrar__cliente_relacionado__entidad_relacionada_id",
        "cuenta_por_cobrar__cliente_relacionado__entidad_relacionada__id",
        "cuenta_por_cobrar__cliente_relacionado__entidad_relacionada__nombre_comercial",
    ).order_by("-fecha_pago", "-id")
    queryset = queryset.filter(
        Q(cuenta_por_cobrar__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(
            cuenta_por_cobrar__entidad_relacionada__isnull=True,
            cuenta_por_cobrar__cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
        )
    )
    if entidad_id:
        queryset = queryset.filter(
            Q(cuenta_por_cobrar__entidad_relacionada_id=entidad_id)
            | Q(
                cuenta_por_cobrar__entidad_relacionada__isnull=True,
                cuenta_por_cobrar__cliente_relacionado__entidad_relacionada_id=entidad_id,
            )
        )
    if cliente_id:
        queryset = queryset.filter(cuenta_por_cobrar__cliente_relacionado_id=cliente_id)
    if estatus_validacion:
        queryset = queryset.filter(estatus_validacion=estatus_validacion)
    safe_page_size = min(max(page_size, 1), 500)
    safe_page = max(page, 1)
    start = (safe_page - 1) * safe_page_size
    end = start + safe_page_size
    return [serialize_payment_cxc(payment) for payment in queryset[start:end]]


@router.get("/cxp/pagos/")
def listar_pagos_cxp(
    request,
    entidad_id: Optional[int] = None,
    estatus_validacion: Optional[str] = None,
):
    require_finance_module(request, "cxp")
    allowed_entity_ids = get_allowed_entity_ids(request)
    queryset = PagoCuentaPorPagar.objects.select_related(
        "cuenta_por_pagar",
        "cuenta_por_pagar__entidad_relacionada",
        "transaccion_relacionada",
    ).order_by("-fecha_pago", "-id")
    queryset = queryset.filter(cuenta_por_pagar__entidad_relacionada_id__in=allowed_entity_ids)
    if entidad_id:
        queryset = queryset.filter(cuenta_por_pagar__entidad_relacionada_id=entidad_id)
    if estatus_validacion:
        queryset = queryset.filter(estatus_validacion=estatus_validacion)
    return [serialize_payment_cxp(payment) for payment in queryset]


@router.get("/conciliacion/cuentas-bancarias/")
def listar_cuentas_bancarias_conciliacion(request):
    require_finance_module(request, "conciliacion")
    current_capa = get_active_business_layer(request)
    latest_loads = CargaConciliacion.objects.filter(
        cuenta_bancaria_relacionada_id=OuterRef("pk")
    ).order_by("-fecha_hasta", "-fecha_carga", "-id")
    queryset = (
        CuentaBancaria.objects.select_related("capa_negocio", "entidad_relacionada")
        .filter(capa_negocio_id=current_capa.id, activa=True)
        .annotate(
            lotes_cargados_count=Count("cargas_conciliacion", distinct=True),
            movimientos_registrados_count=Count("transacciones", distinct=True),
            ultimo_periodo_desde=Subquery(latest_loads.values("fecha_desde")[:1]),
            ultimo_periodo_hasta=Subquery(latest_loads.values("fecha_hasta")[:1]),
            ultimo_saldo_inicial=Subquery(latest_loads.values("saldo_inicial")[:1]),
            ultimo_saldo_final=Subquery(latest_loads.values("saldo_final")[:1]),
            ultimo_total_ingresos=Subquery(latest_loads.values("total_ingresos")[:1]),
            ultimo_total_egresos=Subquery(latest_loads.values("total_egresos")[:1]),
        )
        .order_by("nombre", "id")
    )
    return [serialize_bank_account(account) for account in queryset]


@router.post("/conciliacion/cuentas-bancarias/")
def crear_cuenta_bancaria_conciliacion(request, payload: CuentaBancariaIn):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    current_capa = get_active_business_layer(request)

    nombre = payload.nombre.strip()
    if not nombre:
        raise HttpError(400, "El nombre de la cuenta bancaria es obligatorio.")

    banco = normalize_bank_identifier(payload.banco, digits_only=False)
    numero_cuenta = normalize_bank_identifier(payload.numero_cuenta)
    clabe = normalize_bank_identifier(payload.clabe)
    ultima_4 = normalize_bank_identifier(payload.ultima_4) or (
        numero_cuenta[-4:] if numero_cuenta else (clabe[-4:] if clabe else "")
    )

    existing_account = find_duplicate_bank_account(
        capa_id=current_capa.id,
        numero_cuenta=numero_cuenta,
        clabe=clabe,
        ultima_4=ultima_4,
        banco=banco,
    )
    if existing_account is not None:
        if not existing_account.activa:
            existing_account.activa = True
            existing_account.save(update_fields=["activa", "fecha_actualizacion"])
        return {
            "id": existing_account.id,
            "mensaje": "La cuenta bancaria ya existía y quedó seleccionada.",
            "cuenta": serialize_bank_account(existing_account),
            "ya_existia": True,
        }

    account = CuentaBancaria.objects.create(
        capa_negocio=current_capa,
        nombre=nombre,
        alias=payload.alias.strip() or None,
        banco=banco or None,
        numero_cuenta=numero_cuenta or None,
        clabe=clabe or None,
        ultima_4=ultima_4 or None,
        moneda=payload.moneda.strip() or "MXN",
        activa=payload.activa,
    )
    return {
        "id": account.id,
        "mensaje": "Cuenta bancaria creada correctamente.",
        "cuenta": serialize_bank_account(account),
    }


@router.put("/conciliacion/cuentas-bancarias/{cuenta_id}/")
def actualizar_cuenta_bancaria_conciliacion(request, cuenta_id: int, payload: CuentaBancariaIn):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    current_capa = get_active_business_layer(request)
    account = get_object_or_404(
        CuentaBancaria,
        id=cuenta_id,
        capa_negocio_id=current_capa.id,
    )

    nombre = payload.nombre.strip()
    if not nombre:
        raise HttpError(400, "El nombre de la cuenta bancaria es obligatorio.")

    banco = normalize_bank_identifier(payload.banco, digits_only=False)
    numero_cuenta = normalize_bank_identifier(payload.numero_cuenta)
    clabe = normalize_bank_identifier(payload.clabe)
    ultima_4 = normalize_bank_identifier(payload.ultima_4) or (
        numero_cuenta[-4:] if numero_cuenta else (clabe[-4:] if clabe else "")
    )

    duplicate = find_duplicate_bank_account(
        capa_id=current_capa.id,
        numero_cuenta=numero_cuenta,
        clabe=clabe,
        ultima_4=ultima_4,
        banco=banco,
        exclude_id=account.id,
    )
    if duplicate is not None:
        raise HttpError(
            409,
            f"Ya existe una cuenta bancaria registrada con esos datos: {duplicate}.",
        )

    account.nombre = nombre
    account.alias = payload.alias.strip() or None
    account.banco = banco or None
    account.numero_cuenta = numero_cuenta or None
    account.clabe = clabe or None
    account.ultima_4 = ultima_4 or None
    account.moneda = payload.moneda.strip() or "MXN"
    account.activa = payload.activa
    account.save(
        update_fields=[
            "nombre",
            "alias",
            "banco",
            "numero_cuenta",
            "clabe",
            "ultima_4",
            "moneda",
            "activa",
            "fecha_actualizacion",
        ]
    )
    return {
        "id": account.id,
        "mensaje": "Cuenta bancaria actualizada correctamente.",
        "cuenta": serialize_bank_account(account),
    }


@router.delete("/conciliacion/cuentas-bancarias/{cuenta_id}/")
def eliminar_cuenta_bancaria_conciliacion(request, cuenta_id: int):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    current_capa = get_active_business_layer(request)
    account = get_object_or_404(
        CuentaBancaria,
        id=cuenta_id,
        capa_negocio_id=current_capa.id,
    )

    has_history = (
        account.cargas_conciliacion.exists()
        or account.transacciones.exists()
    )
    if has_history:
        account.activa = False
        account.save(update_fields=["activa", "fecha_actualizacion"])
        return {
            "id": account.id,
            "mensaje": "La cuenta tiene historial, por eso quedó desactivada y ya no aparecerá para nuevas cargas.",
            "desactivada": True,
        }

    account.delete()
    return {
        "id": cuenta_id,
        "mensaje": "Cuenta bancaria eliminada correctamente.",
        "eliminada": True,
    }


@router.get("/conciliacion/cargas/")
def listar_cargas_conciliacion(
    request,
    cuenta_bancaria_id: Optional[int] = None,
    page: int = 1,
    page_size: int = 12,
):
    require_finance_module(request, "conciliacion")
    current_capa = get_active_business_layer(request)
    page = max(page, 1)
    page_size = min(max(page_size, 1), 50)
    queryset = CargaConciliacion.objects.select_related("cuenta_bancaria_relacionada").order_by(
        "-fecha_carga",
        "-id",
    )
    queryset = queryset.filter(cuenta_bancaria_relacionada__capa_negocio_id=current_capa.id)
    if cuenta_bancaria_id:
        queryset = queryset.filter(
            cuenta_bancaria_relacionada_id=cuenta_bancaria_id,
            cuenta_bancaria_relacionada__capa_negocio_id=current_capa.id,
        )
    total = queryset.count()
    start = (page - 1) * page_size
    items = queryset[start : start + page_size]
    return {
        "items": [serialize_reconciliation_load(item) for item in items],
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_items": total,
            "total_pages": max(1, (total + page_size - 1) // page_size),
        },
    }


def _decimal_from_metadata(value) -> Decimal:
    try:
        return max(Decimal(str(value or "0")), Decimal("0"))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


def _revertir_saldo_a_favor_de_transacciones(transaction_ids: list[int]) -> int:
    transacciones = Transaccion.objects.filter(id__in=transaction_ids).exclude(metadata__isnull=True)
    saldos_revertidos = 0
    for transaccion in transacciones:
        metadata = transaccion.metadata or {}
        saldo_aplicado = _decimal_from_metadata(metadata.get("saldo_a_favor_aplicado"))
        cliente_id = metadata.get("saldo_a_favor_cliente_id")
        if saldo_aplicado <= 0 or not cliente_id:
            continue

        saldo_cliente = SaldoCliente.objects.filter(cliente_relacionado_id=cliente_id).first()
        if not saldo_cliente:
            continue

        saldo_cliente.saldo_a_favor = max(
            (saldo_cliente.saldo_a_favor or Decimal("0")) - saldo_aplicado,
            Decimal("0"),
        )
        saldo_cliente.save(update_fields=["saldo_a_favor", "fecha_actualizacion"])
        saldos_revertidos += 1
    return saldos_revertidos


def _revertir_pagos_de_transacciones(transaction_ids: list[int]) -> dict[str, int]:
    cxc_payments = PagoCuentaPorCobrar.objects.filter(transaccion_relacionada_id__in=transaction_ids)
    cxp_payments = PagoCuentaPorPagar.objects.filter(transaccion_relacionada_id__in=transaction_ids)

    cxc_totals = list(
        cxc_payments.exclude(estatus_validacion="RECHAZADO")
        .values("cuenta_por_cobrar_id")
        .annotate(total=Sum("monto"))
    )
    cxp_totals = list(
        cxp_payments.exclude(estatus_validacion="RECHAZADO")
        .values("cuenta_por_pagar_id")
        .annotate(total=Sum("monto"))
    )
    cxc_count = cxc_payments.count()
    cxp_count = cxp_payments.count()

    cxc_payments.delete()
    cxp_payments.delete()

    for item in cxc_totals:
        cuenta = CuentaPorCobrar.objects.filter(id=item["cuenta_por_cobrar_id"]).first()
        if not cuenta:
            continue
        cuenta.monto_pagado = max(
            (cuenta.monto_pagado or Decimal("0")) - (item["total"] or Decimal("0")),
            Decimal("0"),
        )
        cuenta.estatus_adeudo = cuenta.recalcular_estatus()
        cuenta.save(update_fields=["monto_pagado", "estatus_adeudo"])

    for item in cxp_totals:
        cuenta = CuentaPorPagar.objects.filter(id=item["cuenta_por_pagar_id"]).first()
        if not cuenta:
            continue
        cuenta.monto_pagado = max(
            (cuenta.monto_pagado or Decimal("0")) - (item["total"] or Decimal("0")),
            Decimal("0"),
        )
        cuenta.estatus = cuenta.recalcular_estatus()
        cuenta.save(update_fields=["monto_pagado", "estatus"])

    return {
        "pagos_cxc_revertidos": cxc_count,
        "pagos_cxp_revertidos": cxp_count,
    }


def _reabrir_eventos_de_transacciones(transaction_ids: list[int]) -> int:
    event_ids = set(
        Transaccion.objects.filter(id__in=transaction_ids)
        .exclude(evento_relacionado__isnull=True)
        .values_list("evento_relacionado_id", flat=True)
    )
    event_ids.update(
        EventoFinanciero.objects.filter(transaccion_relacionada_id__in=transaction_ids)
        .values_list("id", flat=True)
    )
    event_ids.discard(None)
    if not event_ids:
        return 0

    EventoFinancieroPartida.objects.filter(evento_relacionado_id__in=event_ids).update(
        cuenta_por_cobrar_relacionada=None,
        cuenta_por_pagar_relacionada=None,
        tipo_destino="NO_IDENTIFICADO",
        estatus="REQUIERE_REVISION",
        requiere_revision_manual=True,
    )
    return EventoFinanciero.objects.filter(id__in=event_ids).update(
        transaccion_relacionada=None,
        estatus="NO_IDENTIFICADO",
        requiere_revision_manual=True,
    )


@router.delete("/conciliacion/cargas/{carga_id}/", response={200: dict, 409: dict})
def eliminar_carga_conciliacion(
    request,
    carga_id: int,
    forzar: bool = False,
    confirmacion: str = "",
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    current_capa = get_active_business_layer(request)
    carga = get_object_or_404(
        CargaConciliacion.objects.select_related("cuenta_bancaria_relacionada"),
        id=carga_id,
        cuenta_bancaria_relacionada__capa_negocio_id=current_capa.id,
    )
    transacciones = Transaccion.objects.filter(carga_relacionada=carga)
    vinculadas = (
        transacciones.filter(
            Q(cuenta_por_cobrar_relacionada__isnull=False)
            | Q(cuenta_por_pagar_relacionada__isnull=False)
            | Q(evento_relacionado__isnull=False)
            | Q(pagos_cxc__isnull=False)
            | Q(pagos_cxp__isnull=False)
        )
        .distinct()
        .count()
    )
    confirmar_borrado_forzado = forzar and confirmacion.strip().upper() == "ELIMINAR"
    if vinculadas and not confirmar_borrado_forzado:
        return 409, {
            "detail": (
                "Esta carga ya tiene movimientos vinculados o conciliados. "
                "Para borrarla debes confirmar la reversa de pagos CxC/CxP y escribir ELIMINAR."
            ),
            "movimientos_bloqueados": vinculadas,
            "requiere_confirmacion_forzada": True,
        }

    with transaction.atomic():
        transaction_ids = list(
            Transaccion.objects.select_for_update()
            .filter(carga_relacionada=carga)
            .values_list("id", flat=True)
        )
        total_transacciones = len(transaction_ids)
        resumen_reversion = {
            "pagos_cxc_revertidos": 0,
            "pagos_cxp_revertidos": 0,
            "eventos_reabiertos": 0,
            "saldos_a_favor_revertidos": 0,
        }
        if transaction_ids:
            resumen_reversion.update(_revertir_pagos_de_transacciones(transaction_ids))
            resumen_reversion["saldos_a_favor_revertidos"] = (
                _revertir_saldo_a_favor_de_transacciones(transaction_ids)
            )
            resumen_reversion["eventos_reabiertos"] = _reabrir_eventos_de_transacciones(
                transaction_ids
            )
            Transaccion.objects.filter(id__in=transaction_ids).update(
                cuenta_por_cobrar_relacionada=None,
                cuenta_por_pagar_relacionada=None,
                evento_relacionado=None,
                estatus_conciliacion="NO_IDENTIFICADO",
                fecha_validacion=None,
            )
            Transaccion.objects.filter(id__in=transaction_ids).delete()
        carga.delete()

    return {
        "success": True,
        "mensaje": (
            "La carga bancaria se elimino correctamente. "
            "Los pagos vinculados se revirtieron y los eventos quedaron sin identificar."
            if vinculadas
            else "La carga bancaria y sus movimientos se eliminaron correctamente."
        ),
        "transacciones_eliminadas": total_transacciones,
        "borrado_forzado": bool(vinculadas),
        **resumen_reversion,
    }


@router.get("/conciliacion/transacciones/")
def listar_transacciones_conciliacion(
    request,
    estatus: Optional[str] = None,
    tipo_movimiento: Optional[str] = None,
    cuenta_bancaria_id: Optional[int] = None,
    carga_id: Optional[int] = None,
    auto_match: str = "",
    busqueda: str = "",
    modo_busqueda: str = "contiene",
    mes: str = "",
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    orden: str = "fecha_desc",
    page: int = 1,
    page_size: int = 30,
):
    require_finance_module(request, "conciliacion")
    allowed_entity_ids = get_allowed_entity_ids(request)
    current_capa = get_active_business_layer(request)
    page = max(page, 1)
    page_size = min(max(page_size, 1), 100)
    queryset = Transaccion.objects.select_related(
        "carga_relacionada",
        "cuenta_bancaria_relacionada",
        "cuenta_por_cobrar_relacionada",
        "cuenta_por_pagar_relacionada",
        "evento_relacionado",
    ).prefetch_related(
        Prefetch(
            "pagos_cxc",
            queryset=PagoCuentaPorCobrar.objects.exclude(
                estatus_validacion="RECHAZADO"
            ).only("id", "monto", "estatus_validacion", "transaccion_relacionada_id"),
            to_attr="_pagos_cxc_vivos",
        ),
        Prefetch(
            "pagos_cxp",
            queryset=PagoCuentaPorPagar.objects.exclude(
                estatus_validacion="RECHAZADO"
            ).only("id", "monto", "estatus_validacion", "transaccion_relacionada_id"),
            to_attr="_pagos_cxp_vivos",
        ),
    )
    queryset = queryset.filter(
        Q(cuenta_por_cobrar_relacionada__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(cuenta_por_pagar_relacionada__entidad_relacionada_id__in=allowed_entity_ids)
        | Q(cuenta_bancaria_relacionada__capa_negocio_id=current_capa.id)
        | Q(evento_relacionado__entidad_relacionada_id__in=allowed_entity_ids)
    )
    if estatus:
        queryset = queryset.filter(estatus_conciliacion=estatus)
    if tipo_movimiento:
        queryset = queryset.filter(tipo_movimiento=tipo_movimiento)
    if cuenta_bancaria_id:
        queryset = queryset.filter(cuenta_bancaria_relacionada_id=cuenta_bancaria_id)
    if carga_id:
        queryset = queryset.filter(carga_relacionada_id=carga_id)
    if mes.strip():
        month_match = re.match(r"^(\d{4})-(\d{2})(?:-\d{2})?$", mes.strip())
        if month_match:
            year_value, month_value = month_match.groups()
            queryset = queryset.filter(
                fecha_pago__year=int(year_value),
                fecha_pago__month=int(month_value),
            )
    if fecha_desde:
        queryset = queryset.filter(fecha_pago__gte=fecha_desde)
    if fecha_hasta:
        queryset = queryset.filter(fecha_pago__lte=fecha_hasta)
    if busqueda.strip():
        term = busqueda.strip()
        if modo_busqueda == "exacto":
            queryset = queryset.filter(
                Q(numero_referencia__iexact=term)
                | Q(folio_bancario__iexact=term)
                | Q(concepto_bancario__iexact=term)
            )
        elif modo_busqueda == "empieza":
            queryset = queryset.filter(
                Q(numero_referencia__istartswith=term)
                | Q(folio_bancario__istartswith=term)
                | Q(concepto_bancario__istartswith=term)
            )
        else:
            queryset = queryset.filter(
                Q(numero_referencia__icontains=term)
                | Q(folio_bancario__icontains=term)
                | Q(concepto_bancario__icontains=term)
            )
    auto_summary = queryset.aggregate(
        sin_analisis=Count(
            "id",
            filter=(
                Q(metadata__isnull=True)
                | Q(metadata={})
                | Q(metadata__auto_match__isnull=True)
            ),
        ),
        revision_manual=Count(
            "id",
            filter=(
                Q(metadata__auto_match__match_status__in=["LOW_CONFIDENCE", "AMBIGUOUS", "PENDING"])
                | Q(metadata__auto_match__confidence_band__in=["BAJA", "MEDIA"])
            ),
        ),
        matched=Count("id", filter=Q(metadata__auto_match__match_status="MATCHED")),
        low_confidence=Count("id", filter=Q(metadata__auto_match__match_status="LOW_CONFIDENCE")),
        ambiguous=Count("id", filter=Q(metadata__auto_match__match_status="AMBIGUOUS")),
        none=Count("id", filter=Q(metadata__auto_match__match_status="NONE")),
        error=Count("id", filter=Q(metadata__auto_match__match_status="ERROR")),
    )
    auto_match_summary = {key: int(value or 0) for key, value in auto_summary.items()}
    auto_match_filter = auto_match.strip().upper()
    if auto_match_filter:
        if auto_match_filter == "SIN_ANALISIS":
            queryset = queryset.filter(
                Q(metadata__isnull=True)
                | Q(metadata={})
                | Q(metadata__auto_match__isnull=True)
            )
        elif auto_match_filter == "REVISION_MANUAL":
            queryset = queryset.filter(
                Q(metadata__auto_match__match_status__in=["LOW_CONFIDENCE", "AMBIGUOUS", "PENDING"])
                | Q(metadata__auto_match__confidence_band__in=["BAJA", "MEDIA"])
            )
        elif auto_match_filter in {"MATCHED", "LOW_CONFIDENCE", "AMBIGUOUS", "NONE", "ERROR"}:
            queryset = queryset.filter(metadata__auto_match__match_status=auto_match_filter)
    ordering_map = {
        "fecha_asc": ("fecha_pago", "id"),
        "fecha_desc": ("-fecha_pago", "-id"),
        "monto_asc": ("monto", "fecha_pago", "id"),
        "monto_desc": ("-monto", "-fecha_pago", "-id"),
    }
    queryset = queryset.order_by(*ordering_map.get(orden, ordering_map["fecha_desc"]))

    summary_totals = queryset.aggregate(
        total_ingresos=Coalesce(
            Sum("monto", filter=Q(tipo_movimiento="INGRESO")),
            Decimal("0"),
        ),
        total_egresos=Coalesce(
            Sum("monto", filter=Q(tipo_movimiento="EGRESO")),
            Decimal("0"),
        ),
        ingresos_count=Count("id", filter=Q(tipo_movimiento="INGRESO")),
        egresos_count=Count("id", filter=Q(tipo_movimiento="EGRESO")),
        total=Count("id"),
    )
    first_transaction = queryset.order_by("fecha_pago", "id").first()
    last_transaction = queryset.order_by("-fecha_pago", "-id").first()
    saldo_inicial_periodo = None
    saldo_final_periodo = None
    if first_transaction and first_transaction.saldo_resultante is not None:
        if first_transaction.tipo_movimiento == "INGRESO":
            saldo_inicial_periodo = first_transaction.saldo_resultante - first_transaction.monto
        else:
            saldo_inicial_periodo = first_transaction.saldo_resultante + first_transaction.monto
    if last_transaction and last_transaction.saldo_resultante is not None:
        saldo_final_periodo = last_transaction.saldo_resultante
    total = int(summary_totals["total"] or 0)
    start = (page - 1) * page_size
    items = queryset[start : start + page_size]
    return {
        "items": [serialize_transaction(item) for item in items],
        "summary": {
            "saldo_inicial": decimal_to_float(saldo_inicial_periodo),
            "saldo_final": decimal_to_float(saldo_final_periodo),
            "ingresos": decimal_to_float(summary_totals["total_ingresos"]),
            "egresos": decimal_to_float(summary_totals["total_egresos"]),
            "ingresos_count": int(summary_totals["ingresos_count"] or 0),
            "egresos_count": int(summary_totals["egresos_count"] or 0),
            "total_movimientos": total,
            "cuenta_bancaria_id": cuenta_bancaria_id,
            "auto_match": auto_match_summary,
        },
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_items": total,
            "total_pages": max(1, (total + page_size - 1) // page_size),
        },
    }


@router.get("/conciliacion/sugerencias/")
def listar_sugerencias_conciliacion(
    request,
    cuenta_bancaria_id: Optional[int] = None,
    limit: int = 80,
):
    require_finance_module(request, "conciliacion")
    return build_bank_reconciliation_suggestions(
        allowed_entity_ids=get_allowed_entity_ids(request),
        capa_id=get_active_business_layer(request).id,
        cuenta_bancaria_id=cuenta_bancaria_id,
        limit=limit,
    )


@router.get("/entidades/{entidad_id}/resumen/")
def obtener_resumen_financiero(request, entidad_id: int):
    require_finance_summary_access(request)
    entidad = get_object_or_404(EntidadNegocio, id=entidad_id, id__in=get_allowed_entity_ids(request))
    return build_finance_summary(
        entidad,
        include_cxc=plan_allows_finance_module(request, "cxc"),
        include_cxp=plan_allows_finance_module(request, "cxp"),
    )




@router.post("/entidades/{entidad_id}/cxc/pagos/")
def registrar_pago_cxc(request, entidad_id: int, payload: PagoCxCIn):
    require_cxc_access(request)
    require_write_access(request)
    entidad = get_object_or_404(EntidadNegocio, id=entidad_id, id__in=get_allowed_entity_ids(request))
    try:
        resultado = register_payment_for_entity(
            entidad,
            monto=payload.monto,
            fecha_pago=payload.fecha_pago,
            metodo=payload.metodo,
            canal_origen=payload.canal_origen,
            referencia=payload.referencia,
            notas=payload.notas,
            cliente_id=payload.cliente_id,
            cuenta_id=payload.cuenta_id,
            evidencia_id=payload.evidencia_id,
        )
        audit(
            actor=request.auth.user,
            capa=entidad.capa_negocio,
            accion="CXC_PAGO_REGISTRADO",
            recurso_tipo="PagoCuentaPorCobrar",
            metadata={
                "entidad_id": entidad.id,
                "cliente_id": payload.cliente_id,
                "cuenta_id": payload.cuenta_id,
                "monto": str(payload.monto),
                "aplicaciones": resultado["aplicaciones"],
            },
        )
        return {
            "mensaje": "Pago registrado correctamente.",
            "aplicaciones": resultado["aplicaciones"],
            "saldo_a_favor": resultado["saldo_a_favor"],
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo registrar el pago: {exc}")


@router.post("/cxc/{cuenta_id}/estatus/")
def actualizar_estatus_cxc(request, cuenta_id: int, payload: EstadoCxCIn):
    require_cxc_access(request)
    require_write_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    cuenta = get_object_or_404(
        CuentaPorCobrar.objects.filter(
            Q(entidad_relacionada_id__in=allowed_entity_ids)
            | Q(
                entidad_relacionada__isnull=True,
                cliente_relacionado__entidad_relacionada_id__in=allowed_entity_ids,
            )
        ),
        id=cuenta_id,
    )
    estatus_validos = {
        "PENDIENTE",
        "PARCIAL",
        "VENCIDO",
        "POR_CONCILIAR",
        "CONCILIADO",
        "CANCELADO",
        "INCOBRABLE",
    }
    if payload.estatus not in estatus_validos:
        raise HttpError(400, "El estatus solicitado no es valido.")

    cuenta.estatus_adeudo = payload.estatus
    cuenta.save(update_fields=["estatus_adeudo"])
    audit(
        actor=request.auth.user,
        capa=resolve_receivable_capa(cuenta),
        accion="CXC_ESTATUS_ACTUALIZADO",
        recurso_tipo="CuentaPorCobrar",
        recurso_id=cuenta.id,
        metadata={"estatus": cuenta.estatus_adeudo, "cliente_id": cuenta.cliente_relacionado_id},
    )
    return {"success": True, "mensaje": "Estatus actualizado correctamente."}


@router.post("/entidades/{entidad_id}/cxp/programaciones/")
def crear_programacion_cxp(request, entidad_id: int, payload: ProgramacionCxPIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    entidad = get_object_or_404(EntidadNegocio, id=entidad_id, id__in=get_allowed_entity_ids(request))
    validate_day(payload.dia_vencimiento)
    if payload.fecha_fin and payload.fecha_fin < payload.fecha_inicio:
        raise HttpError(400, "La fecha fin no puede ser anterior a la fecha inicio.")
    if payload.monto_base <= 0:
        raise HttpError(400, "El monto base debe ser mayor a 0.")

    try:
        programacion = ProgramacionCuentaPorPagar.objects.create(
            entidad_relacionada=entidad,
            nombre=payload.nombre.strip(),
            categoria=payload.categoria,
            naturaleza=payload.naturaleza,
            proveedor_nombre=payload.proveedor_nombre.strip(),
            banco_pago=payload.banco_pago.strip() or None,
            cuenta_pago=payload.cuenta_pago.strip() or None,
            clabe_pago=payload.clabe_pago.strip() or None,
            prioridad=payload.prioridad
            if payload.prioridad in {"BAJA", "MEDIA", "ALTA", "CRITICA"}
            else "MEDIA",
            periodicidad=payload.periodicidad,
            fecha_inicio=payload.fecha_inicio,
            fecha_fin=payload.fecha_fin,
            dia_vencimiento=payload.dia_vencimiento,
            monto_base=payload.monto_base,
            dias_gracia=max(payload.dias_gracia or 0, 0),
            genera_recargo=payload.genera_recargo,
            prorrateable=payload.prorrateable,
            activo=payload.activo,
            observaciones=(payload.observaciones or "").strip() or None,
        )
    except IntegrityError:
        raise HttpError(
            400,
            "Ya existe una programacion de gasto con ese nombre para la entidad.",
        )
    except Exception as exc:
        raise HttpError(500, f"No se pudo crear la programacion: {exc}")

    audit(
        actor=request.auth.user,
        capa=entidad.capa_negocio,
        accion="CXP_PROGRAMACION_CREADA",
        recurso_tipo="ProgramacionCuentaPorPagar",
        recurso_id=programacion.id,
        metadata={"entidad_id": entidad.id, "nombre": programacion.nombre},
    )
    return {"id": programacion.id, "mensaje": "Programacion creada correctamente."}


@router.put("/cxp/programaciones/{programacion_id}/")
def actualizar_programacion_cxp(request, programacion_id: int, payload: ProgramacionCxPIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    programacion = get_object_or_404(
        ProgramacionCuentaPorPagar.objects.filter(
            entidad_relacionada_id__in=get_allowed_entity_ids(request)
        ),
        id=programacion_id,
    )
    validate_day(payload.dia_vencimiento)
    if payload.fecha_fin and payload.fecha_fin < payload.fecha_inicio:
        raise HttpError(400, "La fecha fin no puede ser anterior a la fecha inicio.")
    if payload.monto_base <= 0:
        raise HttpError(400, "El monto base debe ser mayor a 0.")

    programacion.nombre = payload.nombre.strip()
    programacion.categoria = payload.categoria
    programacion.naturaleza = payload.naturaleza
    programacion.proveedor_nombre = payload.proveedor_nombre.strip()
    programacion.banco_pago = payload.banco_pago.strip() or None
    programacion.cuenta_pago = payload.cuenta_pago.strip() or None
    programacion.clabe_pago = payload.clabe_pago.strip() or None
    programacion.prioridad = (
        payload.prioridad
        if payload.prioridad in {"BAJA", "MEDIA", "ALTA", "CRITICA"}
        else "MEDIA"
    )
    programacion.periodicidad = payload.periodicidad
    programacion.fecha_inicio = payload.fecha_inicio
    programacion.fecha_fin = payload.fecha_fin
    programacion.dia_vencimiento = payload.dia_vencimiento
    programacion.monto_base = payload.monto_base
    programacion.dias_gracia = max(payload.dias_gracia or 0, 0)
    programacion.genera_recargo = payload.genera_recargo
    programacion.prorrateable = payload.prorrateable
    programacion.activo = payload.activo
    programacion.observaciones = (payload.observaciones or "").strip() or None
    try:
        programacion.save()
    except IntegrityError:
        raise HttpError(
            400,
            "Ya existe otra programacion de gasto con ese nombre para la entidad.",
        )
    except Exception as exc:
        raise HttpError(500, f"No se pudo actualizar la programacion: {exc}")

    audit(
        actor=request.auth.user,
        capa=programacion.entidad_relacionada.capa_negocio,
        accion="CXP_PROGRAMACION_ACTUALIZADA",
        recurso_tipo="ProgramacionCuentaPorPagar",
        recurso_id=programacion.id,
        metadata={
            "entidad_id": programacion.entidad_relacionada_id,
            "nombre": programacion.nombre,
            "activo": programacion.activo,
        },
    )
    return {"success": True, "mensaje": "Programacion actualizada correctamente."}


@router.delete("/cxp/programaciones/{programacion_id}/")
def borrar_programacion_cxp(request, programacion_id: int):
    require_finance_module(request, "cxp")
    require_write_access(request)
    programacion = get_object_or_404(
        ProgramacionCuentaPorPagar.objects.filter(
            entidad_relacionada_id__in=get_allowed_entity_ids(request)
        ),
        id=programacion_id,
    )
    audit(
        actor=request.auth.user,
        capa=programacion.entidad_relacionada.capa_negocio,
        accion="CXP_PROGRAMACION_ELIMINADA",
        recurso_tipo="ProgramacionCuentaPorPagar",
        recurso_id=programacion.id,
        metadata={"entidad_id": programacion.entidad_relacionada_id, "nombre": programacion.nombre},
    )
    programacion.delete()
    return {"success": True, "mensaje": "Programacion eliminada correctamente."}


@router.post("/entidades/{entidad_id}/cxp/generar/")
def generar_cuentas_por_pagar(request, entidad_id: int, payload: GenerarCargosIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    entidad = get_object_or_404(EntidadNegocio, id=entidad_id, id__in=get_allowed_entity_ids(request))
    try:
        resultado = generate_payables_for_entity(
            entidad,
            through_date=payload.fecha_corte,
        )
        return {
            "mensaje": "Cuentas por pagar generadas correctamente.",
            "creados": resultado["creados"],
            "existentes": resultado["existentes"],
        }
    except Exception as exc:
        raise HttpError(500, f"No se pudieron generar las cuentas por pagar: {exc}")


@router.post("/entidades/{entidad_id}/cxp/manual/")
def crear_gasto_manual(request, entidad_id: int, payload: GastoManualIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    entidad = get_object_or_404(EntidadNegocio, id=entidad_id, id__in=get_allowed_entity_ids(request))
    if payload.fecha_vencimiento < payload.fecha_emision:
        raise HttpError(400, "La fecha de vencimiento no puede ser anterior a la emision.")
    if payload.monto_proyectado <= 0:
        raise HttpError(400, "El monto proyectado debe ser mayor a 0.")
    if payload.monto_real is not None and payload.monto_real <= 0:
        raise HttpError(400, "El monto real debe ser mayor a 0.")

    try:
        cuenta = CuentaPorPagar.objects.create(
            entidad_relacionada=entidad,
            proveedor_nombre=payload.proveedor_nombre.strip(),
            banco_pago=payload.banco_pago.strip() or None,
            cuenta_pago=payload.cuenta_pago.strip() or None,
            clabe_pago=payload.clabe_pago.strip() or None,
            prioridad=payload.prioridad
            if payload.prioridad in {"BAJA", "MEDIA", "ALTA", "CRITICA"}
            else "MEDIA",
            categoria=payload.categoria,
            naturaleza=payload.naturaleza,
            concepto=payload.concepto.strip(),
            periodicidad=payload.periodicidad,
            tipo_registro="MANUAL",
            fecha_periodo_inicio=payload.fecha_periodo_inicio,
            fecha_periodo_fin=payload.fecha_periodo_fin,
            fecha_programada=payload.fecha_emision,
            fecha_emision=payload.fecha_emision,
            fecha_vencimiento=payload.fecha_vencimiento,
            monto_proyectado=payload.monto_proyectado,
            monto_real=payload.monto_real,
            dias_gracia=max(payload.dias_gracia or 0, 0),
            genera_recargo=payload.genera_recargo,
            estatus="PENDIENTE",
            observaciones=(payload.observaciones or "").strip() or None,
        )
    except Exception as exc:
        raise HttpError(500, f"No se pudo crear el gasto manual: {exc}")

    audit(
        actor=request.auth.user,
        capa=entidad.capa_negocio,
        accion="CXP_CREADA",
        recurso_tipo="CuentaPorPagar",
        recurso_id=cuenta.id,
        metadata={
            "entidad_id": entidad.id,
            "concepto": cuenta.concepto,
            "proveedor_nombre": cuenta.proveedor_nombre,
        },
    )
    return {"id": cuenta.id, "mensaje": "Gasto manual creado correctamente."}


@router.post("/cxp/{cuenta_id}/pagos/")
def registrar_pago_cxp(request, cuenta_id: int, payload: PagoCxPIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    cuenta = get_object_or_404(
        CuentaPorPagar.objects.filter(entidad_relacionada_id__in=get_allowed_entity_ids(request)),
        id=cuenta_id,
    )
    if payload.monto <= 0:
        raise HttpError(400, "El monto del pago debe ser mayor a 0.")
    try:
        resultado = register_payment_for_payable(
            cuenta,
            monto=payload.monto,
            fecha_pago=payload.fecha_pago,
            metodo=payload.metodo,
            canal_origen=payload.canal_origen,
            referencia=payload.referencia,
            notas=payload.notas,
            evidencia_id=payload.evidencia_id,
        )
        audit(
            actor=request.auth.user,
            capa=cuenta.entidad_relacionada.capa_negocio,
            accion="CXP_PAGO_REGISTRADO",
            recurso_tipo="PagoCuentaPorPagar",
            metadata={
                "cuenta_id": cuenta.id,
                "monto": str(payload.monto),
                "monto_aplicado": resultado["monto_aplicado"],
            },
        )
        return {
            "mensaje": "Pago del gasto registrado correctamente.",
            "monto_aplicado": resultado["monto_aplicado"],
            "saldo_restante": resultado["saldo_restante"],
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo registrar el pago del gasto: {exc}")


@router.post("/cxc/pagos/{pago_id}/validar/")
def validar_pago_cxc(request, pago_id: int, payload: ValidarPagoIn):
    require_cxc_access(request)
    require_write_access(request)
    payment = get_object_or_404(scoped_payment_cxc_queryset(request), id=pago_id)
    transaction_item = (
        get_scoped_transaction_or_404(request, payload.transaccion_id)
        if payload.transaccion_id
        else None
    )
    validate_receivable_payment(
        payment,
        transaction_item=transaction_item,
        notes=payload.notas_validacion,
    )
    audit(
        actor=request.auth.user,
        capa=resolve_receivable_capa(payment.cuenta_por_cobrar),
        accion="CXC_PAGO_VALIDADO",
        recurso_tipo="PagoCuentaPorCobrar",
        recurso_id=payment.id,
        metadata={"cuenta_id": payment.cuenta_por_cobrar_id, "monto": str(payment.monto)},
    )
    return {"success": True, "mensaje": "Pago CxC validado correctamente."}


@router.post("/cxc/pagos/{pago_id}/rechazar/")
def rechazar_pago_cxc(request, pago_id: int, payload: RechazarPagoIn):
    require_cxc_access(request)
    require_write_access(request)
    payment = get_object_or_404(scoped_payment_cxc_queryset(request), id=pago_id)
    reject_receivable_payment(payment, reason=payload.motivo)
    return {"success": True, "mensaje": "Pago CxC rechazado y revertido correctamente."}


@router.post("/cxp/pagos/{pago_id}/validar/")
def validar_pago_cxp(request, pago_id: int, payload: ValidarPagoIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    payment = get_object_or_404(scoped_payment_cxp_queryset(request), id=pago_id)
    transaction_item = (
        get_scoped_transaction_or_404(request, payload.transaccion_id)
        if payload.transaccion_id
        else None
    )
    validate_payable_payment(
        payment,
        transaction_item=transaction_item,
        notes=payload.notas_validacion,
    )
    audit(
        actor=request.auth.user,
        capa=payment.cuenta_por_pagar.entidad_relacionada.capa_negocio,
        accion="CXP_PAGO_VALIDADO",
        recurso_tipo="PagoCuentaPorPagar",
        recurso_id=payment.id,
        metadata={"cuenta_id": payment.cuenta_por_pagar_id, "monto": str(payment.monto)},
    )
    return {"success": True, "mensaje": "Pago CxP validado correctamente."}


@router.post("/cxp/pagos/{pago_id}/rechazar/")
def rechazar_pago_cxp(request, pago_id: int, payload: RechazarPagoIn):
    require_finance_module(request, "cxp")
    require_write_access(request)
    payment = get_object_or_404(scoped_payment_cxp_queryset(request), id=pago_id)
    reject_payable_payment(payment, reason=payload.motivo)
    return {"success": True, "mensaje": "Pago CxP rechazado y revertido correctamente."}


@router.post("/conciliacion/cargar/", response={200: dict, 400: dict, 409: dict, 500: dict})
def cargar_estado_de_cuenta(request, file: UploadedFile = File(...), async_job: bool = False):
    require_finance_module(request, "conciliacion")
    try:
        return _cargar_estado_de_cuenta_impl(request, file, async_job=async_job)
    except HttpError:
        raise
    except Exception as exc:
        raise HttpError(500, f"No se pudo procesar el estado de cuenta: {exc}") from exc


def _cargar_estado_de_cuenta_impl(request, file: UploadedFile, *, async_job: bool = False):
    require_write_access(request)
    current_capa = get_active_business_layer(request)
    post_data = {key: str(request.POST.get(key, "") or "") for key in request.POST.keys()}
    if async_job:
        upload_info = store_background_upload(file, prefix="finanzas-conciliacion")
        job = enqueue_background_job(
            kind="finanzas.bank_statement_import",
            payload={
                "file_path": upload_info["path"],
                "file_name": upload_info["original_name"],
                "capa_id": current_capa.id,
                "post_data": post_data,
            },
            created_by=getattr(getattr(request, "auth", None), "user", None),
            max_attempts=1,
        )
        return {
            "accepted": True,
            "job_id": job.id,
            "status": job.status,
            "mensaje": "Estado de cuenta enviado a procesamiento en background.",
        }

    parsed_statement = parse_uploaded_bank_file(file)
    movimientos = parsed_statement.movimientos
    cuenta_bancaria = None
    raw_account_id = str(request.POST.get("cuenta_bancaria_id", "") or "").strip()
    raw_account_name = str(request.POST.get("nombre_cuenta", "") or "").strip()
    should_create_account = str(
        request.POST.get("confirmar_creacion_cuenta", "") or ""
    ).strip().lower() in {"1", "true", "si", "sí", "yes"}
    if raw_account_id:
        cuenta_bancaria = get_object_or_404(
            CuentaBancaria,
            id=int(raw_account_id),
            capa_negocio_id=current_capa.id,
            activa=True,
        )
    else:
        cuenta_bancaria = find_matching_bank_account(parsed_statement, current_capa.id)

    if (
        cuenta_bancaria is None
        and not should_create_account
        and not raw_account_name
        and (
            parsed_statement.banco
            or parsed_statement.numero_cuenta
            or parsed_statement.clabe
            or parsed_statement.ultima_4
        )
    ):
        return 409, {
            "detail": (
                "Detectamos una cuenta bancaria nueva para tu capa activa. "
                "Confirma si deseas crearla antes de importar el estado de cuenta."
            ),
            "requires_account_creation_confirmation": True,
            "sugerencia_creacion_cuenta": build_account_creation_suggestion(
                parsed_statement=parsed_statement,
                capa=current_capa,
            ),
        }

    if cuenta_bancaria is None and (
        raw_account_name
        or parsed_statement.banco
        or parsed_statement.numero_cuenta
        or parsed_statement.clabe
        or parsed_statement.ultima_4
    ):
        banco = normalize_bank_identifier(
            str(request.POST.get("banco", "") or "").strip() or parsed_statement.banco,
            digits_only=False,
        )
        numero_cuenta = normalize_bank_identifier(
            str(request.POST.get("numero_cuenta", "") or "").strip()
            or parsed_statement.numero_cuenta
            or ""
        )
        clabe = normalize_bank_identifier(
            str(request.POST.get("clabe", "") or "").strip() or parsed_statement.clabe or ""
        )
        ultima_4 = normalize_bank_identifier(
            str(request.POST.get("ultima_4", "") or "").strip()
            or parsed_statement.ultima_4
            or ""
        ) or (
            numero_cuenta[-4:] if numero_cuenta else (clabe[-4:] if clabe else "")
        )

        existing_account = find_duplicate_bank_account(
            capa_id=current_capa.id,
            numero_cuenta=numero_cuenta,
            clabe=clabe,
            ultima_4=ultima_4,
            banco=banco,
        )
        if existing_account is not None:
            if not existing_account.activa:
                existing_account.activa = True
                existing_account.save(update_fields=["activa", "fecha_actualizacion"])
            cuenta_bancaria = existing_account
        else:
            cuenta_bancaria = CuentaBancaria.objects.create(
                capa_negocio=current_capa,
                nombre=raw_account_name or build_default_bank_account_name(parsed_statement),
                alias=str(request.POST.get("alias_cuenta", "") or "").strip() or None,
                banco=banco or None,
                numero_cuenta=numero_cuenta or None,
                clabe=clabe or None,
                ultima_4=ultima_4 or None,
            )

    if cuenta_bancaria is None:
        raise HttpError(
            400,
            "Selecciona una cuenta bancaria de la capa activa o captura una nueva cuenta vinculada a tu negocio.",
        )

    extension = os.path.splitext(file.name or "")[1].lower().replace(".", "")
    resultado = import_bank_transactions(
        movimientos=movimientos,
        nombre_archivo=file.name or "",
        cuenta_bancaria=cuenta_bancaria,
        formato_archivo=extension or None,
        fecha_desde=parse_optional_request_date(request.POST.get("fecha_desde"))
        or parsed_statement.fecha_desde,
        fecha_hasta=parse_optional_request_date(request.POST.get("fecha_hasta"))
        or parsed_statement.fecha_hasta,
        saldo_inicial=parse_optional_request_decimal(request.POST.get("saldo_inicial"))
        or parsed_statement.saldo_inicial,
        saldo_final=parse_optional_request_decimal(request.POST.get("saldo_final"))
        or parsed_statement.saldo_final,
        observaciones=str(request.POST.get("observaciones", "") or "").strip(),
        attempt_auto_match=False,
    )
    return {
        "mensaje": "Estado de cuenta procesado correctamente.",
        **resultado,
    }


def import_bank_statement_from_storage(payload: dict[str, object]) -> dict[str, object]:
    file_path = str(payload.get("file_path") or "")
    file_name = str(payload.get("file_name") or file_path)
    capa_id = int(payload["capa_id"])
    post_data = {
        str(key): str(value or "")
        for key, value in (payload.get("post_data") or {}).items()
    }
    if not file_path:
        raise ValueError("No se encontro el archivo de estado de cuenta.")
    current_capa = get_object_or_404(CapaNegocio, id=capa_id)
    try:
        with default_storage.open(file_path, "rb") as stored_file:
            parsed_statement = parse_bank_file_object(stored_file, file_name)
            return import_bank_statement_payload(
                parsed_statement=parsed_statement,
                file_name=file_name,
                post_data=post_data,
                current_capa=current_capa,
            )
    finally:
        if default_storage.exists(file_path):
            default_storage.delete(file_path)


def import_bank_statement_payload(
    *,
    parsed_statement: ParsedBankStatement,
    file_name: str,
    post_data: dict[str, str],
    current_capa: CapaNegocio,
) -> dict[str, object]:
    movimientos = parsed_statement.movimientos
    cuenta_bancaria = None
    raw_account_id = str(post_data.get("cuenta_bancaria_id", "") or "").strip()
    raw_account_name = str(post_data.get("nombre_cuenta", "") or "").strip()
    should_create_account = str(
        post_data.get("confirmar_creacion_cuenta", "") or ""
    ).strip().lower() in {"1", "true", "si", "sÃ­", "yes"}
    if raw_account_id:
        cuenta_bancaria = get_object_or_404(
            CuentaBancaria,
            id=int(raw_account_id),
            capa_negocio_id=current_capa.id,
            activa=True,
        )
    else:
        cuenta_bancaria = find_matching_bank_account(parsed_statement, current_capa.id)

    if (
        cuenta_bancaria is None
        and not should_create_account
        and not raw_account_name
        and (
            parsed_statement.banco
            or parsed_statement.numero_cuenta
            or parsed_statement.clabe
            or parsed_statement.ultima_4
        )
    ):
        return {
            "detail": (
                "Detectamos una cuenta bancaria nueva para tu capa activa. "
                "Confirma si deseas crearla antes de importar el estado de cuenta."
            ),
            "requires_account_creation_confirmation": True,
            "sugerencia_creacion_cuenta": build_account_creation_suggestion(
                parsed_statement=parsed_statement,
                capa=current_capa,
            ),
        }

    if cuenta_bancaria is None and (
        raw_account_name
        or parsed_statement.banco
        or parsed_statement.numero_cuenta
        or parsed_statement.clabe
        or parsed_statement.ultima_4
    ):
        banco = normalize_bank_identifier(
            str(post_data.get("banco", "") or "").strip() or parsed_statement.banco,
            digits_only=False,
        )
        numero_cuenta = normalize_bank_identifier(
            str(post_data.get("numero_cuenta", "") or "").strip()
            or parsed_statement.numero_cuenta
            or ""
        )
        clabe = normalize_bank_identifier(
            str(post_data.get("clabe", "") or "").strip() or parsed_statement.clabe or ""
        )
        ultima_4 = normalize_bank_identifier(
            str(post_data.get("ultima_4", "") or "").strip()
            or parsed_statement.ultima_4
            or ""
        ) or (
            numero_cuenta[-4:] if numero_cuenta else (clabe[-4:] if clabe else "")
        )
        existing_account = find_duplicate_bank_account(
            capa_id=current_capa.id,
            numero_cuenta=numero_cuenta,
            clabe=clabe,
            ultima_4=ultima_4,
            banco=banco,
        )
        if existing_account is not None:
            if not existing_account.activa:
                existing_account.activa = True
                existing_account.save(update_fields=["activa", "fecha_actualizacion"])
            cuenta_bancaria = existing_account
        else:
            cuenta_bancaria = CuentaBancaria.objects.create(
                capa_negocio=current_capa,
                nombre=raw_account_name or build_default_bank_account_name(parsed_statement),
                alias=str(post_data.get("alias_cuenta", "") or "").strip() or None,
                banco=banco or None,
                numero_cuenta=numero_cuenta or None,
                clabe=clabe or None,
                ultima_4=ultima_4 or None,
            )

    if cuenta_bancaria is None:
        raise HttpError(
            400,
            "Selecciona una cuenta bancaria de la capa activa o captura una nueva cuenta vinculada a tu negocio.",
        )

    extension = os.path.splitext(file_name or "")[1].lower().replace(".", "")
    resultado = import_bank_transactions(
        movimientos=movimientos,
        nombre_archivo=file_name or "",
        cuenta_bancaria=cuenta_bancaria,
        formato_archivo=extension or None,
        fecha_desde=parse_optional_request_date(post_data.get("fecha_desde"))
        or parsed_statement.fecha_desde,
        fecha_hasta=parse_optional_request_date(post_data.get("fecha_hasta"))
        or parsed_statement.fecha_hasta,
        saldo_inicial=parse_optional_request_decimal(post_data.get("saldo_inicial"))
        or parsed_statement.saldo_inicial,
        saldo_final=parse_optional_request_decimal(post_data.get("saldo_final"))
        or parsed_statement.saldo_final,
        observaciones=str(post_data.get("observaciones", "") or "").strip(),
        attempt_auto_match=False,
    )
    return {
        "mensaje": "Estado de cuenta procesado correctamente.",
        **resultado,
    }


@router.get("/conciliacion/eventos/")
def listar_eventos_financieros(
    request,
    estatus: Optional[str] = None,
    tipo_movimiento: Optional[str] = None,
    entidad_id: Optional[int] = None,
    cliente_id: Optional[int] = None,
    busqueda: str = "",
):
    require_finance_module(request, "conciliacion")
    allowed_entity_ids = get_allowed_entity_ids(request)
    if entidad_id and entidad_id not in allowed_entity_ids:
        raise HttpError(404, "La entidad solicitada no pertenece a tu capa activa.")
    return build_events_dashboard(
        estatus=estatus,
        tipo_movimiento=tipo_movimiento,
        entidad_id=entidad_id,
        cliente_id=cliente_id,
        busqueda=busqueda or None,
        allowed_entity_ids=allowed_entity_ids,
    )


@router.get("/conciliacion/eventos/{event_id}/")
def obtener_evento_financiero(request, event_id: int):
    require_finance_module(request, "conciliacion")
    event = get_object_or_404(
        event_detail_queryset(request).select_related("duplicado_de"),
        id=event_id,
    )
    return serialize_event_detail(event)


@router.post("/conciliacion/eventos/")
def crear_evento_financiero(
    request,
    payload: EventoFinancieroIn,
    include_event_detail: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    ensure_financial_payload_scope(request, payload)
    try:
        event = upsert_financial_event(
            caso_id=payload.caso_id,
            evidencia_id=payload.evidencia_id,
            entidad_id=payload.entidad_id,
            cliente_id=payload.cliente_id,
            transaccion_id=payload.transaccion_id,
            origen_evento=payload.origen_evento,
            tipo_movimiento=payload.tipo_movimiento,
            unidad_detectada=payload.unidad_detectada,
            beneficiario_principal=payload.beneficiario_principal,
            referencia_principal=payload.referencia_principal,
            fecha_evento=payload.fecha_evento,
            monto_total_reportado=payload.monto_total_reportado,
            confianza_global=payload.confianza_global,
            requiere_revision_manual=payload.requiere_revision_manual,
            texto_consolidado=payload.texto_consolidado,
            raw_data=payload.raw_data,
            propuesta_ia=payload.propuesta_ia,
            observaciones=payload.observaciones,
            estatus=payload.estatus,
            partidas=[item.dict() for item in payload.partidas],
        )
        return {
            "id": event.id,
            "mensaje": "Evento financiero registrado correctamente.",
            "evento": serialize_event_mutation_result(
                request,
                event,
                include_event_detail,
            ),
        }
    except Exception as exc:
        raise HttpError(500, f"No se pudo registrar el evento financiero: {exc}")


@router.put("/conciliacion/eventos/{event_id}/")
def actualizar_evento_financiero(
    request,
    event_id: int,
    payload: EventoFinancieroIn,
    include_event_detail: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    get_scoped_event_or_404(request, event_id)
    ensure_financial_payload_scope(request, payload)
    try:
        event = upsert_financial_event(
            event_id=event_id,
            caso_id=payload.caso_id,
            evidencia_id=payload.evidencia_id,
            entidad_id=payload.entidad_id,
            cliente_id=payload.cliente_id,
            transaccion_id=payload.transaccion_id,
            origen_evento=payload.origen_evento,
            tipo_movimiento=payload.tipo_movimiento,
            unidad_detectada=payload.unidad_detectada,
            beneficiario_principal=payload.beneficiario_principal,
            referencia_principal=payload.referencia_principal,
            fecha_evento=payload.fecha_evento,
            monto_total_reportado=payload.monto_total_reportado,
            confianza_global=payload.confianza_global,
            requiere_revision_manual=payload.requiere_revision_manual,
            texto_consolidado=payload.texto_consolidado,
            raw_data=payload.raw_data,
            propuesta_ia=payload.propuesta_ia,
            observaciones=payload.observaciones,
            estatus=payload.estatus,
            partidas=[item.dict() for item in payload.partidas],
        )
        return {
            "success": True,
            "mensaje": "Evento financiero actualizado correctamente.",
            "evento": serialize_event_mutation_result(
                request,
                event,
                include_event_detail,
            ),
        }
    except Exception as exc:
        raise HttpError(500, f"No se pudo actualizar el evento financiero: {exc}")


@router.post("/conciliacion/eventos/{event_id}/aplicar/")
def aplicar_evento_financiero(
    request,
    event_id: int,
    include_event_detail: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    event = get_object_or_404(
        event_mutation_queryset(request),
        id=event_id,
    )
    try:
        event = apply_financial_event(event)
        audit(
            actor=request.auth.user,
            capa=resolve_event_capa(request, event),
            accion="EVENTO_FINANCIERO_APLICADO",
            recurso_tipo="EventoFinanciero",
            recurso_id=event.id,
            metadata={"estatus": event.estatus, "tipo_movimiento": event.tipo_movimiento},
        )
        return {
            "success": True,
            "mensaje": "Evento financiero aplicado correctamente.",
            "evento": serialize_event_mutation_result(
                request,
                event,
                include_event_detail,
            ),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo aplicar el evento financiero: {exc}")


@router.post("/conciliacion/eventos/{event_id}/conciliar/")
def conciliar_evento_financiero(
    request,
    event_id: int,
    payload: ConciliarEventoIn,
    include_event_detail: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    event = get_object_or_404(
        event_reconciliation_queryset(request),
        id=event_id,
    )
    transaction_item = get_scoped_transaction_or_404(request, payload.transaccion_id)

    if event.transaccion_relacionada_id and event.transaccion_relacionada_id != transaction_item.id:
        raise HttpError(
            400,
            "Este evento ya esta vinculado a otra transaccion bancaria.",
        )

    try:
        if event.estatus in {"NUEVO", "PROPUESTO", "PENDIENTE_APLICACION"}:
            event = apply_financial_event(
                event,
                refresh_evidence=False,
                sync_case_status=False,
            )
            for attr_name in (
                "_prefetched_reconciliation_cxc_payments",
                "_prefetched_reconciliation_cxp_payments",
            ):
                if hasattr(event, attr_name):
                    delattr(event, attr_name)
        event = reconcile_event_with_transaction(
            event,
            transaction_item,
            notes=payload.notas or "Conciliado manualmente desde la bandeja.",
        )
        audit(
            actor=request.auth.user,
            capa=resolve_event_capa(request, event),
            accion="EVENTO_FINANCIERO_CONCILIADO",
            recurso_tipo="EventoFinanciero",
            recurso_id=event.id,
            metadata={"transaccion_id": transaction_item.id, "estatus": event.estatus},
        )
        return {
            "success": True,
            "mensaje": "Evento financiero conciliado correctamente.",
            "evento": serialize_event_mutation_result(
                request,
                event,
                include_event_detail,
            ),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo conciliar el evento financiero: {exc}")


@router.post("/conciliacion/eventos/{event_id}/estatus/")
def actualizar_estatus_evento_financiero(
    request,
    event_id: int,
    payload: EventoEstatusIn,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    event = get_scoped_event_or_404(request, event_id)
    valid_statuses = {choice[0] for choice in EventoFinanciero.ESTATUS_CHOICES}
    if payload.estatus not in valid_statuses:
        raise HttpError(400, "El estatus solicitado no es valido.")

    event.estatus = payload.estatus
    if payload.observaciones.strip():
        event.observaciones = payload.observaciones.strip()
    event.save(update_fields=["estatus", "observaciones", "fecha_actualizacion"])

    if payload.estatus in {"DESCARTADO", "DUPLICADO"}:
        event.partidas.exclude(estatus="CONCILIADA").update(estatus="OMITIDA")

    update_event_case_status(event)
    audit(
        actor=request.auth.user,
        capa=resolve_event_capa(request, event),
        accion="EVENTO_FINANCIERO_ESTATUS_ACTUALIZADO",
        recurso_tipo="EventoFinanciero",
        recurso_id=event.id,
        metadata={"estatus": event.estatus},
    )
    return {
        "success": True,
        "mensaje": "Estatus del evento actualizado correctamente.",
        "evento": serialize_event_detail(event),
    }


@router.get("/conciliacion/transacciones/{transaccion_id}/")
def obtener_detalle_transaccion_conciliacion(request, transaccion_id: int):
    require_finance_module(request, "conciliacion")
    transaction_item = get_object_or_404(
        scoped_transaction_queryset(request).select_related(
            "carga_relacionada",
            "cuenta_bancaria_relacionada",
            "cuenta_bancaria_relacionada__entidad_relacionada",
            "cuenta_por_cobrar_relacionada",
            "cuenta_por_cobrar_relacionada__entidad_relacionada",
            "cuenta_por_cobrar_relacionada__cliente_relacionado__entidad_relacionada",
            "cuenta_por_pagar_relacionada",
            "cuenta_por_pagar_relacionada__entidad_relacionada",
            "evento_relacionado",
            "evento_relacionado__entidad_relacionada",
            "evento_relacionado__cliente_relacionado__entidad_relacionada",
        ),
        id=transaccion_id,
    )
    return build_bank_transaction_detail(
        transaction_item,
        allowed_entity_ids=get_allowed_entity_ids(request),
    )


@router.post("/conciliacion/transacciones/{transaccion_id}/aplicar-cxc/")
def aplicar_transaccion_bancaria_a_cxc(
    request,
    transaccion_id: int,
    payload: AplicarTransaccionCxCIn,
    include_transaction: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    transaction_item = get_object_or_404(
        scoped_transaction_queryset(request).select_related(
            "cuenta_bancaria_relacionada",
            "cuenta_bancaria_relacionada__entidad_relacionada",
        ),
        id=transaccion_id,
    )
    if payload.entidad_id:
        get_object_or_404(
            EntidadNegocio,
            id=payload.entidad_id,
            id__in=get_allowed_entity_ids(request),
        )
    ensure_client_scope(request, payload.cliente_id)
    cuenta = None
    if payload.cuenta_id:
        cuenta = get_object_or_404(
            scoped_receivable_queryset(request).select_related(
                "cliente_relacionado",
                "entidad_relacionada",
                "cliente_relacionado__entidad_relacionada",
            ),
            id=payload.cuenta_id,
        )
    try:
        previous_status = transaction_item.estatus_conciliacion
        resultado = apply_bank_transaction_to_receivables(
            transaction_item,
            entidad_id=payload.entidad_id,
            cliente_id=payload.cliente_id,
            cuenta_id=payload.cuenta_id,
            cuenta_obj=cuenta,
            monto=payload.monto,
            notas=payload.notas,
        )
        transaction_item.refresh_from_db()
        audit_metadata = {
            **resultado,
            "transaccion": build_transaction_audit_context(
                transaction_item,
                previous_status=previous_status,
                notes=payload.notas,
            ),
        }
        suggestion_context = build_suggestion_audit_context(payload)
        if suggestion_context:
            audit_metadata["sugerencia"] = suggestion_context
        audit(
            actor=request.auth.user,
            capa=get_current_capa(request),
            accion="TRANSACCION_APLICADA_CXC",
            recurso_tipo="Transaccion",
            recurso_id=transaction_item.id,
            metadata=audit_metadata,
        )
        return {
            "success": True,
            "mensaje": "Transaccion aplicada a CxC correctamente.",
            "resultado": resultado,
            "transaccion": (
                serialize_transaction(transaction_item) if include_transaction else None
            ),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo aplicar la transaccion a CxC: {exc}")


@router.post("/conciliacion/transacciones/{transaccion_id}/aplicar-cxp/")
def aplicar_transaccion_bancaria_a_cxp(
    request,
    transaccion_id: int,
    payload: AplicarTransaccionCxPIn,
    include_transaction: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_finance_module(request, "cxp")
    require_write_access(request)
    transaction_item = get_object_or_404(
        scoped_transaction_queryset(request).select_related(
            "cuenta_bancaria_relacionada",
            "cuenta_bancaria_relacionada__entidad_relacionada",
        ),
        id=transaccion_id,
    )
    cuenta = get_object_or_404(
        scoped_payable_queryset(request).select_related("entidad_relacionada"),
        id=payload.cuenta_id,
    )
    try:
        previous_status = transaction_item.estatus_conciliacion
        resultado = apply_bank_transaction_to_payable(
            transaction_item,
            cuenta_id=payload.cuenta_id,
            cuenta_obj=cuenta,
            monto=payload.monto,
            notas=payload.notas,
        )
        transaction_item.refresh_from_db()
        audit_metadata = {
            **resultado,
            "transaccion": build_transaction_audit_context(
                transaction_item,
                previous_status=previous_status,
                notes=payload.notas,
            ),
        }
        suggestion_context = build_suggestion_audit_context(payload)
        if suggestion_context:
            audit_metadata["sugerencia"] = suggestion_context
        audit(
            actor=request.auth.user,
            capa=get_current_capa(request),
            accion="TRANSACCION_APLICADA_CXP",
            recurso_tipo="Transaccion",
            recurso_id=transaction_item.id,
            metadata=audit_metadata,
        )
        return {
            "success": True,
            "mensaje": "Transaccion aplicada a CxP correctamente.",
            "resultado": resultado,
            "transaccion": (
                serialize_transaction(transaction_item) if include_transaction else None
            ),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo aplicar la transaccion a CxP: {exc}")


@router.post("/conciliacion/transacciones/{transaccion_id}/estatus/")
def actualizar_estatus_transaccion_conciliacion(
    request,
    transaccion_id: int,
    payload: EstadoTransaccionIn,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    transaction_item = get_scoped_transaction_or_404(request, transaccion_id)
    try:
        previous_status = transaction_item.estatus_conciliacion
        transaction_item = set_bank_transaction_status(
            transaction_item,
            status=payload.estatus,
            notes=payload.notas,
        )
        audit(
            actor=request.auth.user,
            capa=get_current_capa(request),
            accion="TRANSACCION_ESTATUS_ACTUALIZADO",
            recurso_tipo="Transaccion",
            recurso_id=transaction_item.id,
            metadata={
                "transaccion": build_transaction_audit_context(
                    transaction_item,
                    previous_status=previous_status,
                    notes=payload.notas,
                ),
                "estatus_anterior": previous_status,
                "estatus_nuevo": transaction_item.estatus_conciliacion,
            },
        )
        return {
            "success": True,
            "mensaje": "Estatus de transaccion actualizado correctamente.",
            "transaccion": serialize_transaction(transaction_item),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo actualizar el estatus de la transaccion: {exc}")


@router.post("/conciliacion/transacciones/{transaccion_id}/auto-vincular/")
def auto_vincular_transaccion_con_evento(
    request,
    transaccion_id: int,
    include_detail: bool = True,
    include_event_detail: bool = True,
    include_transaction: bool = True,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    transaction_item = get_object_or_404(
        transaction_auto_match_queryset(request),
        id=transaccion_id,
    )
    try:
        result = auto_match_bank_transaction(transaction_item)
        event = None
        event_payload = None
        if result.get("event_id"):
            if include_event_detail:
                event = event_detail_queryset(request).filter(id=result["event_id"]).first()
                event_payload = serialize_event_detail(event) if event else None
            else:
                event_payload = {"id": result["event_id"]}
        return {
            "success": True,
            "resultado": result,
            "transaccion": (
                serialize_transaction(transaction_item) if include_transaction else None
            ),
            "detalle": (
                build_bank_transaction_detail(
                    transaction_item,
                    allowed_entity_ids=get_allowed_entity_ids(request),
                )
                if include_detail
                else None
            ),
            "evento": event_payload,
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo auto-vincular la transaccion: {exc}")


@router.post("/conciliacion/auto-conciliar-pendientes/")
def auto_conciliar_transacciones_pendientes(
    request,
    payload: AutoConciliarPendientesIn,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    current_capa = get_active_business_layer(request)
    if payload.cuenta_bancaria_id:
        get_object_or_404(
            CuentaBancaria,
            id=payload.cuenta_bancaria_id,
            capa_negocio_id=current_capa.id,
        )
    if payload.carga_id:
        get_object_or_404(
            CargaConciliacion.objects.select_related("cuenta_bancaria_relacionada"),
            id=payload.carga_id,
            cuenta_bancaria_relacionada__capa_negocio_id=current_capa.id,
        )
    try:
        resultado = auto_reconcile_pending_bank_transactions(
            allowed_entity_ids=allowed_entity_ids,
            cuenta_bancaria_id=payload.cuenta_bancaria_id,
            carga_id=payload.carga_id,
            limit=payload.limite,
        )
        return {
            "success": True,
            "mensaje": "Conciliacion automatica ejecutada.",
            "resultado": resultado,
        }
    except Exception as exc:
        raise HttpError(500, f"No se pudo ejecutar la conciliacion automatica: {exc}")


@router.post("/conciliacion/transacciones/{transaccion_id}/vincular-evento/")
def vincular_transaccion_a_evento(
    request,
    transaccion_id: int,
    payload: VincularEventoIn,
):
    require_finance_module(request, "conciliacion")
    require_write_access(request)
    transaction_item = get_scoped_transaction_or_404(request, transaccion_id)
    event = get_object_or_404(
        event_mutation_queryset(request),
        id=payload.evento_id,
    )

    if event.transaccion_relacionada_id and event.transaccion_relacionada_id != transaction_item.id:
        raise HttpError(400, "El evento ya esta vinculado a otra transaccion.")

    try:
        if event.estatus in {"NUEVO", "PROPUESTO", "PENDIENTE_APLICACION"}:
            event = apply_financial_event(event)
        event = reconcile_event_with_transaction(
            event,
            transaction_item,
            notes=payload.notas or "Transaccion vinculada manualmente.",
        )
        audit(
            actor=request.auth.user,
            capa=get_current_capa(request),
            accion="TRANSACCION_VINCULADA_EVENTO",
            recurso_tipo="Transaccion",
            recurso_id=transaction_item.id,
            metadata={"evento_id": event.id, "estatus_evento": event.estatus},
        )
        return {
            "success": True,
            "mensaje": "Transaccion vinculada al evento correctamente.",
            "transaccion": serialize_transaction(transaction_item),
            "evento": serialize_event_mutation_result(request, event, True),
        }
    except ValueError as exc:
        raise HttpError(400, str(exc))
    except Exception as exc:
        raise HttpError(500, f"No se pudo vincular la transaccion: {exc}")


@router.get("/cxc/")
def listar_cartera_cxc(
    request,
    entidad_id: Optional[int] = None,
    cliente_id: Optional[int] = None,
    status: Optional[str] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    busqueda: str = "",
    agrupacion: str = "",
    page: int = 1,
    page_size: int = 20,
    include_items: bool = True,
):
    require_cxc_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    if entidad_id and entidad_id not in allowed_entity_ids:
        raise HttpError(404, "La entidad solicitada no pertenece a tu capa activa.")

    return build_global_cxc_dashboard(
        entidad_id=entidad_id,
        cliente_id=cliente_id,
        status=status,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        busqueda=busqueda or "",
        page=page,
        page_size=page_size,
        agrupacion=agrupacion or "",
        allowed_entity_ids=allowed_entity_ids,
        include_items=include_items,
    )


@router.get("/cxc/clientes/{cliente_id}/detalle/")
def obtener_detalle_cxc_cliente(
    request,
    cliente_id: int,
    entidad_id: Optional[int] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
):
    require_cxc_access(request)
    allowed_entity_ids = get_allowed_entity_ids(request)
    if entidad_id and entidad_id not in allowed_entity_ids:
        raise HttpError(404, "La entidad solicitada no pertenece a tu capa activa.")

    try:
        return build_cxc_customer_detail(
            cliente_id=cliente_id,
            entidad_id=entidad_id,
            fecha_desde=fecha_desde,
            fecha_hasta=fecha_hasta,
            allowed_entity_ids=allowed_entity_ids,
        )
    except ValueError as exc:
        raise HttpError(404, str(exc))
