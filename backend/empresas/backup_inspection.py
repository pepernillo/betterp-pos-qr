import hashlib
import io

from django.core.exceptions import ValidationError
from openpyxl import load_workbook


REQUIRED_SHEETS = {
    "Resumen",
    "Entidades",
    "Clientes",
    "CxC",
    "Pagos CxC",
    "CxP",
    "Pagos CxP",
    "Comprobantes",
    "Cuentas bancarias",
    "Cargas conciliacion",
    "Transacciones",
    "Eventos financieros",
}


def inspect_backup_workbook(
    content: bytes,
    *,
    expected_checksum: str = "",
) -> dict:
    checksum = hashlib.sha256(content).hexdigest()
    expected = (expected_checksum or "").strip().lower()
    if expected and checksum != expected:
        raise ValidationError("El checksum del backup no coincide con el historial registrado.")

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet_names = set(workbook.sheetnames)
    missing = sorted(REQUIRED_SHEETS - sheet_names)
    if missing:
        raise ValidationError(f"Backup incompleto. Faltan hojas: {', '.join(missing)}")

    sheets = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        sheets.append(
            {
                "nombre": sheet_name,
                "filas": max((sheet.max_row or 1) - 1, 0),
            }
        )

    return {
        "checksum_sha256": checksum,
        "hojas": sheets,
    }
