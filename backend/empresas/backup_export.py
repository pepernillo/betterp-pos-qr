import io
import json
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal

from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from comunicaciones.models import EvidenciaPago
from crm.models import Cliente
from finanzas.models import (
    CargaConciliacion,
    CuentaBancaria,
    CuentaPorCobrar,
    CuentaPorPagar,
    EventoFinanciero,
    PagoCuentaPorCobrar,
    PagoCuentaPorPagar,
    Transaccion,
)

from .models import CapaNegocio, EntidadNegocio

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
MIN_COLUMN_WIDTH = 12
MAX_COLUMN_WIDTH = 42


def format_export_value(value) -> object:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (list, dict)):
        return json_safe_text(value)
    return value


def json_safe_text(value) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def append_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: Iterable[Sequence[object]],
) -> None:
    sheet = workbook.create_sheet(title[:31])
    header_cells = []
    for index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(index)
        sheet.column_dimensions[column_letter].width = min(
            max(len(str(header)) + 2, MIN_COLUMN_WIDTH),
            MAX_COLUMN_WIDTH,
        )
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        header_cells.append(cell)
    sheet.append(header_cells)
    for row in rows:
        formatted_row = [format_export_value(value) for value in row]
        sheet.append(formatted_row)


def build_capa_backup_workbook(capa: CapaNegocio) -> bytes:
    entidad_ids = list(
        EntidadNegocio.objects.filter(capa_negocio=capa).values_list("id", flat=True)
    )
    clientes = list(
        Cliente.objects.select_related("entidad_relacionada")
        .filter(entidad_relacionada_id__in=entidad_ids)
        .order_by("entidad_relacionada__nombre_comercial", "razon_social", "id")
    )
    cliente_ids = [cliente.id for cliente in clientes]

    workbook = Workbook(write_only=True)
    generated_at = timezone.now()

    counts = {
        "entidades": len(entidad_ids),
        "clientes": len(clientes),
        "cxc": CuentaPorCobrar.objects.filter(
            Q(entidad_relacionada_id__in=entidad_ids)
            | Q(cliente_relacionado_id__in=cliente_ids)
        ).count(),
        "cxp": CuentaPorPagar.objects.filter(
            entidad_relacionada_id__in=entidad_ids
        ).count(),
        "comprobantes": EvidenciaPago.objects.filter(
            entidad_relacionada_id__in=entidad_ids
        ).count(),
        "transacciones": Transaccion.objects.filter(
            cuenta_bancaria_relacionada__capa_negocio=capa
        ).count(),
    }
    append_sheet(
        workbook,
        "Resumen",
        ["Campo", "Valor"],
        [
            ["Capa ID", capa.id],
            ["Capa", capa.nombre],
            ["Tipo", capa.tipo_capa],
            ["Generado", generated_at],
            [
                "Alcance",
                "Export operativo MVP por capa; no incluye secretos ni restore avanzado.",
            ],
            ["Entidades", counts["entidades"]],
            ["Clientes", counts["clientes"]],
            ["CxC", counts["cxc"]],
            ["CxP", counts["cxp"]],
            ["Comprobantes", counts["comprobantes"]],
            ["Transacciones", counts["transacciones"]],
        ],
    )

    entidades = list(
        EntidadNegocio.objects.filter(id__in=entidad_ids).order_by(
            "nombre_comercial",
            "id",
        )
    )
    append_sheet(
        workbook,
        "Entidades",
        [
            "ID",
            "Nombre",
            "Ciudad",
            "RFC",
            "Regimen fiscal",
            "Tipo corte",
            "Dia corte global",
            "Activa",
        ],
        (
            [
                entidad.id,
                entidad.nombre_comercial,
                entidad.ciudad,
                entidad.rfc,
                entidad.regimen_fiscal,
                entidad.tipo_fecha_corte,
                entidad.dia_corte_general,
                entidad.activo,
            ]
            for entidad in entidades
        ),
    )
    append_sheet(
        workbook,
        "Clientes",
        [
            "ID",
            "Entidad ID",
            "Entidad",
            "Identificador",
            "Razon social",
            "Nombre comercial",
            "RFC",
            "Correo",
            "Telefono",
            "Dia corte",
            "Dias gracia",
            "Activo",
        ],
        (
            [
                cliente.id,
                cliente.entidad_relacionada_id,
                cliente.entidad_relacionada.nombre_comercial,
                cliente.identificador,
                cliente.razon_social,
                cliente.nombre_comercial,
                cliente.rfc,
                cliente.correo_principal,
                f"{cliente.codigo_pais or ''} {cliente.telefono or ''}".strip(),
                cliente.dia_corte_individual,
                cliente.dias_gracia,
                cliente.activo,
            ]
            for cliente in clientes
        ),
    )
    cxc = list(
        CuentaPorCobrar.objects.select_related(
            "entidad_relacionada",
            "cliente_relacionado",
        )
        .filter(
            Q(entidad_relacionada_id__in=entidad_ids)
            | Q(cliente_relacionado_id__in=cliente_ids)
        )
        .order_by("fecha_vencimiento", "id")
    )
    cxc_ids = [cuenta.id for cuenta in cxc]
    append_sheet(
        workbook,
        "CxC",
        [
            "ID",
            "Entidad",
            "Cliente",
            "Concepto",
            "Origen",
            "Periodo inicio",
            "Periodo fin",
            "Vencimiento",
            "Monto total",
            "Monto pagado",
            "Saldo pendiente",
            "Estatus",
            "Referencia",
        ],
        (
            [
                cuenta.id,
                (
                    cuenta.entidad_relacionada.nombre_comercial
                    if cuenta.entidad_relacionada_id
                    else ""
                ),
                cuenta.cliente_relacionado.razon_social,
                cuenta.concepto,
                cuenta.origen,
                cuenta.fecha_periodo_inicio,
                cuenta.fecha_periodo_fin,
                cuenta.fecha_vencimiento,
                cuenta.monto_total,
                cuenta.monto_pagado,
                cuenta.saldo_pendiente,
                cuenta.estatus_adeudo,
                cuenta.referencia_unica,
            ]
            for cuenta in cxc
        ),
    )
    pagos_cxc = list(
        PagoCuentaPorCobrar.objects.select_related(
            "cuenta_por_cobrar",
            "evidencia_pago_relacionada",
            "transaccion_relacionada",
        )
        .filter(cuenta_por_cobrar_id__in=cxc_ids)
        .order_by("-fecha_pago", "-id")
    )
    append_sheet(
        workbook,
        "Pagos CxC",
        [
            "ID",
            "CxC ID",
            "Fecha",
            "Monto",
            "Metodo",
            "Canal",
            "Estatus",
            "Referencia",
            "Evidencia ID",
            "Transaccion ID",
        ],
        (
            [
                pago.id,
                pago.cuenta_por_cobrar_id,
                pago.fecha_pago,
                pago.monto,
                pago.metodo,
                pago.canal_origen,
                pago.estatus_validacion,
                pago.referencia,
                pago.evidencia_pago_relacionada_id,
                pago.transaccion_relacionada_id,
            ]
            for pago in pagos_cxc
        ),
    )

    cxp = list(
        CuentaPorPagar.objects.select_related(
            "entidad_relacionada",
            "programacion_relacionada",
        )
        .filter(entidad_relacionada_id__in=entidad_ids)
        .order_by("fecha_vencimiento", "id")
    )
    cxp_ids = [cuenta.id for cuenta in cxp]
    append_sheet(
        workbook,
        "CxP",
        [
            "ID",
            "Entidad",
            "Proveedor",
            "Concepto",
            "Categoria",
            "Naturaleza",
            "Periodo inicio",
            "Periodo fin",
            "Vencimiento",
            "Monto proyectado",
            "Monto real",
            "Monto pagado",
            "Saldo pendiente",
            "Estatus",
            "Referencia",
        ],
        (
            [
                cuenta.id,
                cuenta.entidad_relacionada.nombre_comercial,
                cuenta.proveedor_nombre,
                cuenta.concepto,
                cuenta.categoria,
                cuenta.naturaleza,
                cuenta.fecha_periodo_inicio,
                cuenta.fecha_periodo_fin,
                cuenta.fecha_vencimiento,
                cuenta.monto_proyectado,
                cuenta.monto_real,
                cuenta.monto_pagado,
                cuenta.saldo_pendiente,
                cuenta.estatus,
                cuenta.referencia_unica,
            ]
            for cuenta in cxp
        ),
    )
    pagos_cxp = list(
        PagoCuentaPorPagar.objects.select_related(
            "cuenta_por_pagar",
            "evidencia_pago_relacionada",
            "transaccion_relacionada",
        )
        .filter(cuenta_por_pagar_id__in=cxp_ids)
        .order_by("-fecha_pago", "-id")
    )
    append_sheet(
        workbook,
        "Pagos CxP",
        [
            "ID",
            "CxP ID",
            "Fecha",
            "Monto",
            "Metodo",
            "Canal",
            "Estatus",
            "Referencia",
            "Evidencia ID",
            "Transaccion ID",
        ],
        (
            [
                pago.id,
                pago.cuenta_por_pagar_id,
                pago.fecha_pago,
                pago.monto,
                pago.metodo,
                pago.canal_origen,
                pago.estatus_validacion,
                pago.referencia,
                pago.evidencia_pago_relacionada_id,
                pago.transaccion_relacionada_id,
            ]
            for pago in pagos_cxp
        ),
    )

    evidencias = list(
        EvidenciaPago.objects.select_related(
            "entidad_relacionada",
            "cliente_relacionado",
        )
        .filter(
            Q(entidad_relacionada_id__in=entidad_ids)
            | Q(cliente_relacionado_id__in=cliente_ids)
        )
        .order_by("-fecha_registro", "-id")
    )
    append_sheet(
        workbook,
        "Comprobantes",
        [
            "ID",
            "Entidad",
            "Cliente",
            "Canal",
            "Tipo",
            "Origen",
            "Monto reportado",
            "Fecha pago",
            "Referencia",
            "Estatus",
            "Archivo URL",
            "Categoria",
            "Registro",
        ],
        (
            [
                evidencia.id,
                (
                    evidencia.entidad_relacionada.nombre_comercial
                    if evidencia.entidad_relacionada_id
                    else ""
                ),
                (
                    evidencia.cliente_relacionado.razon_social
                    if evidencia.cliente_relacionado_id
                    else ""
                ),
                evidencia.canal,
                evidencia.tipo_movimiento,
                evidencia.origen_deteccion,
                evidencia.monto_reportado,
                evidencia.fecha_pago_reportada,
                evidencia.referencia_reportada,
                evidencia.estatus,
                evidencia.url_archivo,
                evidencia.categoria_sugerida,
                evidencia.fecha_registro,
            ]
            for evidencia in evidencias
        ),
    )

    cuentas_bancarias = list(
        CuentaBancaria.objects.filter(
            Q(capa_negocio=capa) | Q(entidad_relacionada_id__in=entidad_ids)
        )
        .select_related("entidad_relacionada")
        .order_by("nombre", "id")
    )
    cuenta_ids = [cuenta.id for cuenta in cuentas_bancarias]
    append_sheet(
        workbook,
        "Cuentas bancarias",
        ["ID", "Entidad", "Nombre", "Alias", "Banco", "Ultima 4", "Moneda", "Activa"],
        (
            [
                cuenta.id,
                (
                    cuenta.entidad_relacionada.nombre_comercial
                    if cuenta.entidad_relacionada_id
                    else ""
                ),
                cuenta.nombre,
                cuenta.alias,
                cuenta.banco,
                cuenta.ultima_4 or ((cuenta.numero_cuenta or cuenta.clabe or "")[-4:]),
                cuenta.moneda,
                cuenta.activa,
            ]
            for cuenta in cuentas_bancarias
        ),
    )
    cargas = list(
        CargaConciliacion.objects.select_related("cuenta_bancaria_relacionada")
        .filter(cuenta_bancaria_relacionada_id__in=cuenta_ids)
        .order_by("-fecha_carga", "-id")
    )
    append_sheet(
        workbook,
        "Cargas conciliacion",
        [
            "ID",
            "Cuenta bancaria ID",
            "Archivo",
            "Formato",
            "Fecha carga",
            "Desde",
            "Hasta",
            "Estatus",
            "Cuadre",
            "Registros",
            "Conciliados",
            "Pendientes",
        ],
        (
            [
                carga.id,
                carga.cuenta_bancaria_relacionada_id,
                carga.nombre_archivo,
                carga.formato_archivo,
                carga.fecha_carga,
                carga.fecha_desde,
                carga.fecha_hasta,
                carga.estatus_procesamiento,
                carga.estatus_cuadre,
                carga.registros_detectados,
                carga.registros_conciliados,
                carga.registros_pendientes,
            ]
            for carga in cargas
        ),
    )
    transacciones = list(
        Transaccion.objects.select_related(
            "cuenta_bancaria_relacionada",
            "cuenta_por_cobrar_relacionada",
            "cuenta_por_pagar_relacionada",
            "evento_relacionado",
        )
        .filter(cuenta_bancaria_relacionada_id__in=cuenta_ids)
        .order_by("-fecha_pago", "-id")
    )
    append_sheet(
        workbook,
        "Transacciones",
        [
            "ID",
            "Cuenta bancaria ID",
            "Origen",
            "Tipo",
            "Fecha",
            "Monto",
            "Concepto banco",
            "Referencia",
            "Folio",
            "Estatus",
            "CxC ID",
            "CxP ID",
            "Evento ID",
        ],
        (
            [
                tx.id,
                tx.cuenta_bancaria_relacionada_id,
                tx.origen,
                tx.tipo_movimiento,
                tx.fecha_pago,
                tx.monto,
                tx.concepto_bancario,
                tx.numero_referencia,
                tx.folio_bancario,
                tx.estatus_conciliacion,
                tx.cuenta_por_cobrar_relacionada_id,
                tx.cuenta_por_pagar_relacionada_id,
                tx.evento_relacionado_id,
            ]
            for tx in transacciones
        ),
    )
    eventos = list(
        EventoFinanciero.objects.select_related(
            "entidad_relacionada",
            "cliente_relacionado",
            "evidencia_pago_relacionada",
            "transaccion_relacionada",
        )
        .filter(
            Q(entidad_relacionada_id__in=entidad_ids)
            | Q(cliente_relacionado_id__in=cliente_ids)
        )
        .order_by("-fecha_evento", "-id")
    )
    append_sheet(
        workbook,
        "Eventos financieros",
        [
            "ID",
            "Entidad",
            "Cliente",
            "Origen",
            "Tipo",
            "Fecha",
            "Monto",
            "Referencia",
            "Estatus",
            "Evidencia ID",
            "Transaccion ID",
            "Revision manual",
        ],
        (
            [
                evento.id,
                (
                    evento.entidad_relacionada.nombre_comercial
                    if evento.entidad_relacionada_id
                    else ""
                ),
                (
                    evento.cliente_relacionado.razon_social
                    if evento.cliente_relacionado_id
                    else ""
                ),
                evento.origen_evento,
                evento.tipo_movimiento,
                evento.fecha_evento,
                evento.monto_total_reportado,
                evento.referencia_principal,
                evento.estatus,
                evento.evidencia_pago_relacionada_id,
                evento.transaccion_relacionada_id,
                evento.requiere_revision_manual,
            ]
            for evento in eventos
        ),
    )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
