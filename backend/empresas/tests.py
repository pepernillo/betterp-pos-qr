import base64
import io
import json
import tempfile
from calendar import monthrange
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.test import override_settings
from openpyxl import Workbook, load_workbook

from accounts.models import EventoAuditoria
from accounts.test_utils import create_api_auth_context
from billing.models import PlanSaaS, SuscripcionCapa
from crm.models import Cliente
from empresas.api import build_entity_metrics_map
from empresas.backup_export import build_capa_backup_workbook
from empresas.models import BackupCapaExport, CapaNegocio, EntidadNegocio, ReglaMarcoNegocio
from finanzas.models import CuentaPorCobrar, CuentaPorPagar


class EmpresasBatchImportTests(TestCase):
    def setUp(self):
        self.user, self.capa, self.auth_headers = create_api_auth_context(
            email="batch-owner@test.local",
            capa_name="Capa Batch",
        )

    def build_workbook_file(
        self,
        *,
        headers: list[str] | None = None,
        data_rows: list[list[object]] | None = None,
    ) -> SimpleUploadedFile:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Operacion"
        sheet.append(["Mes", "Ano"])
        sheet.append(["Abril", 2026])
        sheet.append(
            headers
            or [
                "Ciudad",
                "Unidad Negocio",
                "Habitacion",
                "Precio Lista",
                "Tipo Espacio",
                "Disponiblidad",
                "Cliente",
                "whats app",
                "correo electronico",
                "Fecha",
                "contrato",
            ]
        )
        for row in (
            data_rows
            or [
                [
                    "CDMX",
                    "Palmas Doral 202 Torre",
                    "Hab 01",
                    15000,
                    "Penthouse",
                    "Ocupada",
                    "Daniel Muller",
                    "5512345678",
                    "daniel@example.com",
                    28,
                    "PALMAS-HAB01",
                ],
                [
                    "CDMX",
                    "Palmas Doral 202 Torre",
                    "Hab 02",
                    9000,
                    "2 recamaras",
                    "Reservado",
                    "",
                    "",
                    "",
                    "",
                    "",
                ],
            ]
        ):
            sheet.append(row)

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return SimpleUploadedFile(
            "batch-espacios.xlsx",
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_batch_template_downloads_workbook(self):
        response = self.client.get(
            "/api/empresas/batch/plantilla/",
            **self.auth_headers,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])

        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn("Carga Batch", workbook.sheetnames)
        self.assertIn("Instrucciones", workbook.sheetnames)
        self.assertIn("Catalogos", workbook.sheetnames)

        sheet = workbook["Carga Batch"]
        self.assertEqual(sheet["A1"].value, "Ciudad o plaza donde opera la unidad.")
        self.assertEqual(sheet["A2"].value, "Ciudad")
        self.assertEqual(sheet["C2"].value, "Tipo Corte")
        self.assertEqual(sheet["H2"].value, "Estatus")
        self.assertEqual(sheet["B3"].value, "Edificio Ejemplo Norte")
        self.assertEqual(sheet["C3"].value, "GENERAL")
        self.assertEqual(sheet["C5"].value, "INDIVIDUAL")
        self.assertEqual(sheet["I3"].value, "Cliente Ejemplo 01")
        self.assertEqual(sheet["K3"].value, "cliente01@example.com")
        validations = list(sheet.data_validations.dataValidation)
        self.assertGreaterEqual(len(validations), 6)
        validation_prompts = " ".join(
            str(validation.prompt or "") for validation in validations
        )
        self.assertIn("GENERAL usa un dia", validation_prompts)
        self.assertIn("monto mayor a 0", validation_prompts)
        catalog = workbook["Catalogos"]
        self.assertEqual(catalog["C2"].value, "Departamento")
        self.assertEqual(catalog["C11"].value, "Otro")

        visible_values = " ".join(
            str(value)
            for row in sheet.iter_rows(min_row=1, max_row=7, values_only=True)
            for value in row
            if value is not None
        )
        self.assertNotIn("Maya", visible_values)
        self.assertNotIn("Daniel Muller", visible_values)
        self.assertNotIn("Daniel", visible_values)








    def test_batch_import_requires_plan_feature(self):
        PlanSaaS.objects.create(
            clave="starter-no-batch",
            nombre="Starter No Batch",
            max_usuarios=3,
            max_entidades=3,
            max_productos=40,
            modulos_habilitados=["dashboard", "entidades", "clientes"],
            funciones_habilitadas=[],
        )
        SuscripcionCapa.objects.create(
            capa_negocio=self.capa,
            plan=PlanSaaS.objects.get(clave="starter-no-batch"),
            estatus="ACTIVA",
        )

        response = self.client.get(
            "/api/empresas/batch/plantilla/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 403)
        self.assertIn("funcion", response.json()["detail"])

    def test_batch_import_requires_entities_module(self):
        plan = PlanSaaS.objects.create(
            clave="starter-batch-no-entities",
            nombre="Starter Batch Sin Entidades",
            max_usuarios=3,
            max_entidades=3,
            max_productos=40,
            modulos_habilitados=["dashboard", "clientes"],
            funciones_habilitadas=["batch_import"],
        )
        SuscripcionCapa.objects.create(capa_negocio=self.capa, plan=plan, estatus="ACTIVA")

        response = self.client.get(
            "/api/empresas/batch/plantilla/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 403)
        self.assertIn("modulo", response.json()["detail"])




class EmpresasCapaNegocioApiTests(TestCase):
    def setUp(self):
        self.user, self.capa, self.auth_headers = create_api_auth_context(
            email="capa-owner@test.local",
            capa_name="Operadora Maya",
        )
        self.capa.periodicidad_cobro_default = "MENSUAL"
        self.capa.dia_vencimiento_default = 5
        self.capa.dias_gracia_default = 3
        self.capa.plazo_meses_default = 12
        self.capa.auto_renueva_default = True
        self.capa.aplicacion_pagos = "ADEUDO_MAS_ANTIGUO"
        self.capa.save()

    def test_crear_entidad_vincula_capa_y_expone_defaults(self):
        response = self.client.post(
            "/api/empresas/crear/",
            data=json.dumps(
                {
                    "nombre_comercial": "Casa Coral",
                    "ciudad": "Merida",
                    "tipo_fecha_corte": "INDIVIDUAL",
                    "activo": True,
                    "capa_negocio_nombre": "Operadora Maya",
                    "capa_negocio_tipo": "OPERADORA",
                    "capa_nombre_administrador": "Equipo Central",
                    "capa_periodicidad_cobro_default": "MENSUAL",
                    "capa_dia_vencimiento_default": 5,
                    "capa_dias_gracia_default": 3,
                    "capa_plazo_meses_default": 12,
                    "capa_auto_renueva_default": True,
                    "capa_aplicacion_pagos": "ADEUDO_MAS_ANTIGUO",
                }
            ),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        entidad = EntidadNegocio.objects.select_related("capa_negocio").get(
            nombre_comercial="Casa Coral"
        )
        self.assertIsNotNone(entidad.capa_negocio)
        self.assertEqual(entidad.capa_negocio.nombre, "Operadora Maya")
        self.assertEqual(entidad.capa_negocio.dia_vencimiento_default, 5)

        detail_response = self.client.get(
            f"/api/empresas/{entidad.id}/",
            **self.auth_headers,
        )
        self.assertEqual(detail_response.status_code, 200)
        body = detail_response.json()
        self.assertEqual(body["capa_negocio_nombre"], "Operadora Maya")
        self.assertEqual(body["capa_periodicidad_cobro_default"], "MENSUAL")
        self.assertEqual(body["capa_dia_vencimiento_default"], 5)
        self.assertEqual(body["capa_dias_gracia_default"], 3)

    def test_plan_entity_limit_blocks_create_entity(self):
        plan = PlanSaaS.objects.create(
            clave="starter-entity-limit",
            nombre="Starter Entity Limit",
            max_usuarios=3,
            max_entidades=1,
            max_productos=40,
            modulos_habilitados=["dashboard", "entidades", "clientes"],
        )
        SuscripcionCapa.objects.create(capa_negocio=self.capa, plan=plan, estatus="ACTIVA")
        EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Casa Existente",
        )

        response = self.client.post(
            "/api/empresas/crear/",
            data=json.dumps(
                {
                    "nombre_comercial": "Casa Nueva",
                    "ciudad": "CDMX",
                    "tipo_fecha_corte": "INDIVIDUAL",
                    "activo": True,
                }
            ),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 403)
        self.assertIn("Tu plan permite hasta 1 entidades", response.json()["detail"])
        self.assertFalse(
            EntidadNegocio.objects.filter(
                capa_negocio=self.capa,
                nombre_comercial="Casa Nueva",
            ).exists()
        )

    def test_plan_entity_limit_blocks_reactivating_entity(self):
        plan = PlanSaaS.objects.create(
            clave="starter-entity-reactivation-limit",
            nombre="Starter Entity Reactivation Limit",
            max_usuarios=3,
            max_entidades=1,
            max_productos=40,
            modulos_habilitados=["dashboard", "entidades", "clientes"],
        )
        SuscripcionCapa.objects.create(capa_negocio=self.capa, plan=plan, estatus="ACTIVA")
        EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Casa Activa",
            activo=True,
        )
        inactive_entity = EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Casa Pausada",
            ciudad="Merida",
            tipo_fecha_corte="INDIVIDUAL",
            activo=False,
        )

        response = self.client.put(
            f"/api/empresas/{inactive_entity.id}/",
            data=json.dumps(
                {
                    "nombre_comercial": "Casa Pausada",
                    "ciudad": "Merida",
                    "tipo_fecha_corte": "INDIVIDUAL",
                    "activo": True,
                }
            ),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 403)
        self.assertIn("Tu plan permite hasta 1 entidades", response.json()["detail"])
        inactive_entity.refresh_from_db()
        self.assertFalse(inactive_entity.activo)

    def test_entidades_routes_require_entities_module(self):
        plan = PlanSaaS.objects.create(
            clave="starter-sin-entidades",
            nombre="Starter Sin Entidades",
            max_usuarios=3,
            max_entidades=3,
            max_productos=40,
            modulos_habilitados=["dashboard", "clientes", "cxc"],
        )
        SuscripcionCapa.objects.create(capa_negocio=self.capa, plan=plan, estatus="ACTIVA")
        entidad = EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Casa Sin Modulo Existente",
            ciudad="CDMX",
            tipo_fecha_corte="INDIVIDUAL",
        )

        list_response = self.client.get("/api/empresas/lista/", **self.auth_headers)
        detail_response = self.client.get(
            f"/api/empresas/{entidad.id}/",
            **self.auth_headers,
        )
        create_response = self.client.post(
            "/api/empresas/crear/",
            data=json.dumps(
                {
                    "nombre_comercial": "Casa Sin Modulo",
                    "ciudad": "CDMX",
                    "tipo_fecha_corte": "INDIVIDUAL",
                    "activo": True,
                }
            ),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(list_response.status_code, 403)
        self.assertEqual(detail_response.status_code, 403)
        self.assertEqual(create_response.status_code, 403)
        self.assertIn("modulo", list_response.json()["detail"])
        self.assertFalse(
            EntidadNegocio.objects.filter(
                capa_negocio=self.capa,
                nombre_comercial="Casa Sin Modulo",
            ).exists()
        )

    def test_entidad_ve_regla_marco_y_luego_override_personalizado(self):
        capa = CapaNegocio.objects.create(
            nombre="Grupo Palmas",
            tipo_capa="OPERADORA",
            periodicidad_cobro_default="MENSUAL",
        )
        _, _, auth_headers = create_api_auth_context(
            email="grupo-palmas-owner@test.local",
            capa=capa,
        )
        entidad = EntidadNegocio.objects.create(
            nombre_comercial="Palmas Centro",
            tipo_fecha_corte="INDIVIDUAL",
            capa_negocio=capa,
        )
        regla_marco = ReglaMarcoNegocio.objects.create(
            capa_negocio=capa,
            nombre="Interes moratorio",
            tipo_calculo="PORCENTAJE_RECARGO",
            valor=10,
            periodicidad="MENSUAL",
            aplica_a_todos=True,
            dias_condicion=3,
            activo=True,
        )

        response = self.client.get(
            f"/api/empresas/{entidad.id}/reglas/",
            **auth_headers,
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body), 1)
        self.assertEqual(body[0]["origen"], "GLOBAL")
        self.assertEqual(body[0]["regla_marco_id"], regla_marco.id)

        personalize_response = self.client.post(
            f"/api/empresas/{entidad.id}/reglas/personalizar/",
            data=json.dumps(
                {
                    "regla_marco_id": regla_marco.id,
                    "nombre": "Interes moratorio local",
                    "tipo_calculo": "PORCENTAJE_RECARGO",
                    "valor": 12,
                    "periodicidad": "MENSUAL",
                    "aplica_a_todos": True,
                    "dias_condicion": 4,
                    "activo": True,
                }
            ),
            content_type="application/json",
            **auth_headers,
        )

        self.assertEqual(personalize_response.status_code, 200)

        response = self.client.get(
            f"/api/empresas/{entidad.id}/reglas/",
            **auth_headers,
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body), 1)
        self.assertEqual(body[0]["origen"], "OVERRIDE")
        self.assertEqual(body[0]["regla_marco_id"], regla_marco.id)
        self.assertEqual(body[0]["nombre"], "Interes moratorio local")

    def test_actualizar_y_bloquear_borrado_de_capa_con_entidades_vinculadas(self):
        capa = CapaNegocio.objects.create(
            nombre="Operadora Centro",
            tipo_capa="OPERADORA",
            periodicidad_cobro_default="MENSUAL",
        )
        _, _, auth_headers = create_api_auth_context(
            email="operadora-centro-owner@test.local",
            capa=capa,
        )
        EntidadNegocio.objects.create(
            nombre_comercial="Casa Centro",
            tipo_fecha_corte="INDIVIDUAL",
            capa_negocio=capa,
        )

        update_response = self.client.put(
            f"/api/empresas/capas/{capa.id}/",
            data=json.dumps(
                {
                    "nombre": "Operadora Centro Actualizada",
                    "tipo_capa": "OPERADORA",
                    "nombre_administrador": "Backoffice",
                    "razon_social": "",
                    "rfc": "",
                    "regimen_fiscal": "",
                    "correo_contacto": "",
                    "telefono_contacto": "",
                    "logo_url": "",
                    "plazo_meses_default": 18,
                    "periodicidad_cobro_default": "MENSUAL",
                    "dia_vencimiento_default": 7,
                    "dias_gracia_default": 2,
                    "auto_renueva_default": False,
                    "aplicacion_pagos": "ADEUDO_MAS_ANTIGUO",
                    "activo": True,
                }
            ),
            content_type="application/json",
            **auth_headers,
        )
        self.assertEqual(update_response.status_code, 200)
        capa.refresh_from_db()
        self.assertEqual(capa.nombre, "Operadora Centro Actualizada")
        self.assertEqual(capa.plazo_meses_default, 18)
        self.assertEqual(capa.dia_vencimiento_default, 7)

        delete_response = self.client.delete(
            f"/api/empresas/capas/{capa.id}/",
            **auth_headers,
        )
        self.assertEqual(delete_response.status_code, 400)

    def test_activar_facturacion_cfdi_requiere_modulo_del_plan(self):
        capa = CapaNegocio.objects.create(
            nombre="Operadora CFDI",
            tipo_capa="OPERADORA",
        )
        _, _, auth_headers = create_api_auth_context(
            email="operadora-cfdi-owner@test.local",
            capa=capa,
        )
        plan = PlanSaaS.objects.create(
            clave="starter-sin-cfdi",
            nombre="Starter Sin CFDI",
            modulos_habilitados=["dashboard", "entidades", "clientes", "cxc"],
        )
        SuscripcionCapa.objects.create(capa_negocio=capa, plan=plan, estatus="ACTIVA")

        response = self.client.put(
            f"/api/empresas/capas/{capa.id}/",
            data=json.dumps(
                {
                    "nombre": "Operadora CFDI",
                    "tipo_capa": "OPERADORA",
                    "facturacion_activa": True,
                }
            ),
            content_type="application/json",
            **auth_headers,
        )

        self.assertEqual(response.status_code, 403)
        self.assertIn("modulo", response.json()["detail"])
