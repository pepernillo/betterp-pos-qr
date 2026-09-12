import json
import os
import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook

from accounts.models import MembresiaCapaNegocio
from accounts.security import create_access_session, issue_token_pair
from accounts.test_utils import create_api_auth_context
from billing.background_jobs import run_due_jobs
from billing.models import PlanSaaS, SuscripcionCapa
from comunicaciones.models import ConfiguracionComunicacion
from empresas.models import CapaNegocio
from empresas.models import EntidadNegocio
from finanzas.models import CuentaPorCobrar

from .api import extract_csf_data_from_text
from .models import Cliente

TEST_STORAGES = {
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
}


class CsfExtractionTests(TestCase):
    def test_extracts_persona_moral_name_and_address(self):
        texto = """
        Constancia de Situacion Fiscal
        RFC: ABC120101AB1
        Denominacion/Razon Social:
        ACME SERVICIOS INTEGRALES SA DE CV
        Regimen Fiscal:
        Regimen General de Ley Personas Morales
        Tipo de Vialidad:
        Avenida
        Nombre de Vialidad:
        Paseo de la Reforma
        Numero Exterior:
        123
        Numero Interior:
        Piso 4
        Nombre de la Colonia:
        Juarez
        Nombre de la Localidad:
        Ciudad de Mexico
        Nombre del Municipio o Demarcacion Territorial:
        Cuauhtemoc
        Nombre de la Entidad Federativa:
        Ciudad de Mexico
        Codigo Postal:
        06600
        """

        datos = extract_csf_data_from_text(texto)

        self.assertEqual(datos["rfc"], "ABC120101AB1")
        self.assertEqual(datos["razon_social"], "ACME SERVICIOS INTEGRALES SA DE CV")
        self.assertTrue(datos["es_persona_moral"])
        self.assertEqual(datos["calle"], "Avenida Paseo de la Reforma")
        self.assertEqual(datos["numero_exterior"], "123")
        self.assertEqual(datos["numero_interior"], "Piso 4")
        self.assertEqual(datos["colonia"], "Juarez")
        self.assertEqual(datos["ciudad"], "Ciudad de Mexico")
        self.assertEqual(datos["estado"], "Ciudad de Mexico")
        self.assertEqual(datos["codigo_postal"], "06600")

    def test_extracts_persona_fisica_full_name(self):
        texto = """
        Cedula de Identificacion Fiscal
        RFC: PEGJ9001011A2
        Nombre (s):
        JUAN CARLOS
        Primer Apellido:
        PEREZ
        Segundo Apellido:
        GOMEZ
        Regimen Fiscal:
        Regimen Simplificado de Confianza
        Nombre de Vialidad:
        Rio Lerma
        Numero Exterior:
        45
        Nombre de la Colonia:
        Cuauhtemoc
        Nombre de la Entidad Federativa:
        Ciudad de Mexico
        Codigo Postal:
        06500
        """

        datos = extract_csf_data_from_text(texto)

        self.assertEqual(datos["rfc"], "PEGJ9001011A2")
        self.assertEqual(datos["razon_social"], "JUAN CARLOS PEREZ GOMEZ")
        self.assertFalse(datos["es_persona_moral"])
        self.assertEqual(datos["regimen_fiscal"], "Regimen Simplificado de Confianza")
        self.assertEqual(datos["calle"], "Rio Lerma")
        self.assertEqual(datos["colonia"], "Cuauhtemoc")
        self.assertEqual(datos["codigo_postal"], "06500")

    def test_extracts_flattened_persona_fisica_fields(self):
        texto = """
        Cedula de Identificacion Fiscal
        RFC
        MUVD780315U21
        """
        texto_lineal = """
        Cedula de Identificacion Fiscal RFC MUVD780315U21 Nombre (s): DANIEL ANTONIO
        Primer Apellido: MUNOZ Segundo Apellido: VALDEZ Regimen Fiscal:
        Regimen Simplificado de Confianza Nombre de Vialidad: INSURGENTES SUR
        Numero Exterior: 742 Numero Interior: 3B Nombre de la Colonia: DEL VALLE
        Nombre de la Localidad: BENITO JUAREZ Nombre de la Entidad Federativa:
        CIUDAD DE MEXICO Codigo Postal: 03100
        """

        datos = extract_csf_data_from_text(texto, texto_lineal)

        self.assertEqual(datos["rfc"], "MUVD780315U21")
        self.assertEqual(datos["razon_social"], "DANIEL ANTONIO MUNOZ VALDEZ")
        self.assertFalse(datos["es_persona_moral"])
        self.assertEqual(datos["regimen_fiscal"], "Regimen Simplificado de Confianza")
        self.assertEqual(datos["calle"], "INSURGENTES SUR")
        self.assertEqual(datos["numero_exterior"], "742")
        self.assertEqual(datos["numero_interior"], "3B")
        self.assertEqual(datos["colonia"], "DEL VALLE")
        self.assertEqual(datos["ciudad"], "BENITO JUAREZ")
        self.assertEqual(datos["estado"], "CIUDAD DE MEXICO")
        self.assertEqual(datos["codigo_postal"], "03100")

    def test_extracts_user_reported_sat_format_without_overcapturing(self):
        texto = """
        Constancia de Situacion Fiscal
        RFC
        MUVD780315U21
        """
        texto_lineal = """
        DANIELANTONIO PrimerApellido: MULLER SegundoApellido: VEGA
        Fechainiciodeoperaciones: 24DEABRILDE2006 Estatusenelpadron: ACTIVO
        Fechadeultimocambiodeestado: 24DEABRILDE2006 NombreComercial:
        Datos del domicilio registrado CodigoPostal:01729 TipodeVialidad:CALZADA
        NombredeVialidad:DESIERTODELOSLEONES NumeroExterior:5547 NumeroInterior:B1001
        NombredelaColonia:ALCANTARILLA NombredelaLocalidad:ALVAROOBREGON
        NombredelMunicipiooDemarcacionTerritorial:ALVAROOBREGON
        NombredelaEntidadFederativa:CIUDADDEMEXICO EntreCalle: AVTOLUCA
        YCalle:TLAHUICOLE Actividades Economicas: Orden
        Regimenes: Regimen Fecha Inicio Fecha Fin Regimen Simplificado de Confianza 30/10/2023
        Obligaciones: Descripcion de la Obligacion Descripcion Vencimiento
        CURP: MUVD780315HSLLGN07
        """

        datos = extract_csf_data_from_text(texto, texto_lineal)

        self.assertEqual(datos["rfc"], "MUVD780315U21")
        self.assertEqual(datos["razon_social"], "DANIEL ANTONIO MULLER VEGA")
        self.assertEqual(datos["regimen_fiscal"], "Regimen Simplificado de Confianza")
        self.assertEqual(datos["codigo_postal"], "01729")
        self.assertEqual(datos["calle"], "CALZADA DESIERTO DE LOS LEONES")
        self.assertEqual(datos["numero_exterior"], "5547")
        self.assertEqual(datos["numero_interior"], "B1001")
        self.assertEqual(datos["colonia"], "ALCANTARILLA")
        self.assertEqual(datos["ciudad"], "ALVARO OBREGON")
        self.assertEqual(datos["estado"], "CIUDAD DE MEXICO")

    def test_detects_multiple_regimes_and_requires_selection(self):
        texto = """
        Constancia de Situacion Fiscal
        RFC GOVA9009195Y0
        Nombre (s): ALFREDO
        Primer Apellido: GODINEZ
        Segundo Apellido: VARGAS
        Codigo Postal: 01729
        """
        texto_lineal = """
        Regimenes: Regimen Fecha Inicio Fecha Fin
        Regimen de las Actividades Empresariales y Profesionales 01/01/2024
        Regimen de las Actividades Empresariales con ingresos a traves de Plataformas Tecnologicas 01/02/2024
        """

        datos = extract_csf_data_from_text(texto, texto_lineal)

        self.assertEqual(datos["rfc"], "GOVA9009195Y0")
        self.assertEqual(datos["regimen_fiscal"], "")
        self.assertTrue(datos["regimen_fiscal_pendiente_seleccion"])
        self.assertEqual(
            [item["codigo"] for item in datos["regimenes_fiscales_detectados"]],
            ["612", "625"],
        )

    def test_ignores_sat_idcif_header_when_extracting_person_name(self):
        texto = """
        Constancia de Situacion Fiscal
        RFC
        LERL700825JI6
        """
        texto_lineal = """
        MAYO DE 2026 idCIF: 16110296416 VALIDA TU INFORMACION FISCAL
        LERL700825JI6 Datos de Identificacion del Contribuyente
        RFC: LERL700825JI6 CURP: LERL700825HDFXXX00 Nombre (s): LUIS ENRIQUE
        Primer Apellido: RODRIGUEZ Segundo Apellido: LOPEZ
        Regimen Fiscal: Sueldos y Salarios e Ingresos Asimilados a Salarios
        Codigo Postal: 52780
        """

        datos = extract_csf_data_from_text(texto, texto_lineal)

        self.assertEqual(datos["rfc"], "LERL700825JI6")
        self.assertEqual(datos["razon_social"], "LUIS ENRIQUE RODRIGUEZ LOPEZ")
        self.assertNotIn("idCIF", datos["razon_social"])
        self.assertFalse(datos["es_persona_moral"])
        self.assertEqual(
            datos["regimen_fiscal"],
            "Sueldos y Salarios e Ingresos Asimilados a Salarios",
        )


class ClienteCreateTests(TestCase):
    def setUp(self):
        self.user, self.capa, self.auth_headers = create_api_auth_context(
            email="crm-owner@test.local",
            capa_name="Capa CRM",
        )
        self.entidad = EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Maya Coliving",
            tipo_fecha_corte="INDIVIDUAL",
        )

    def activate_plan(self, *, modules: list[str], features: list[str] | None = None):
        plan = PlanSaaS.objects.create(
            clave=f"crm-plan-{PlanSaaS.objects.count() + 1}",
            nombre="CRM Plan",
            max_usuarios=5,
            max_entidades=5,
            max_productos=50,
            modulos_habilitados=modules,
            funciones_habilitadas=features or [],
        )
        return SuscripcionCapa.objects.create(
            capa_negocio=self.capa,
            plan=plan,
            estatus="ACTIVA",
        )

    def test_plan_without_clientes_module_blocks_crm_endpoints(self):
        self.activate_plan(modules=["dashboard", "cxc"])

        list_response = self.client.get("/api/crm/lista/", **self.auth_headers)
        create_response = self.client.post(
            "/api/crm/crear/",
            data=json.dumps(
                {
                    "entidad_relacionada_id": self.entidad.id,
                    "razon_social": "Cliente bloqueado",
                    "rfc": "BLO900101AB1",
                    "telefono": "5512345678",
                }
            ),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(list_response.status_code, 403)
        self.assertEqual(create_response.status_code, 403)
        self.assertIn("modulo", list_response.json()["detail"])

    def test_plan_without_batch_import_blocks_client_batch_routes(self):
        self.activate_plan(modules=["clientes"], features=[])

        template_response = self.client.get(
            "/api/crm/batch/plantilla/",
            **self.auth_headers,
        )
        upload = SimpleUploadedFile(
            "clientes.xlsx",
            b"archivo-no-procesado",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        import_response = self.client.post(
            "/api/crm/batch/importar/",
            {"file": upload},
            **self.auth_headers,
        )

        self.assertEqual(template_response.status_code, 403)
        self.assertEqual(import_response.status_code, 403)
        self.assertIn("funcion", import_response.json()["detail"])

    @override_settings(
        STORAGES=TEST_STORAGES,
        MEDIA_ROOT=os.path.join(tempfile.gettempdir(), "betterp-test-media"),
    )
    def test_async_client_batch_import_creates_background_job(self):
        self.activate_plan(modules=["clientes"], features=["batch_import"])
        content = (
            "Entidad,Razon social,RFC,Telefono,Correo principal\n"
            f"{self.entidad.nombre_comercial},Cliente Async,AAA010101AAA,5511112222,async@example.com\n"
        ).encode("utf-8")
        upload = SimpleUploadedFile(
            "clientes.csv",
            content,
            content_type="text/csv",
        )

        response = self.client.post(
            "/api/crm/batch/importar/?async_job=true",
            {"file": upload},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["accepted"])
        self.assertEqual(body["status"], "PENDING")

        summary = run_due_jobs(limit=1, worker_id="crm-test-worker")

        self.assertEqual(summary["success"], 1)
        cliente = Cliente.objects.get(rfc="AAA010101AAA")
        self.assertEqual(cliente.razon_social, "Cliente Async")

        job_response = self.client.get(
            f"/api/billing/jobs/{body['job_id']}/",
            **self.auth_headers,
        )
        self.assertEqual(job_response.status_code, 200)
        self.assertEqual(job_response.json()["status"], "SUCCESS")
        self.assertEqual(job_response.json()["result"]["clientes_creados"], 1)

    def test_create_endpoint_persists_fiscal_and_address_fields(self):
        payload = {
            "entidad_relacionada_id": self.entidad.id,
            "es_persona_moral": False,
            "nombre_comercial": "Depto 3",
            "razon_social": "Juan Perez Gomez",
            "rfc": "pegj9001011a2",
            "regimen_fiscal": "Regimen Simplificado de Confianza",
            "identificador": "Depto 3",
            "condiciones_pago": "Contado",
            "archivo_csf_url": "https://example.com/csf.pdf",
            "correo_principal": "juan@example.com",
            "codigo_pais": "+52",
            "telefono": "5512345678",
            "pais": "Mexico",
            "estado": "Ciudad de Mexico",
            "ciudad": "Ciudad de Mexico",
            "colonia": "Cuauhtemoc",
            "calle": "Rio Lerma",
            "numero_exterior": "45",
            "numero_interior": "2B",
            "codigo_postal": "06500",
            "agente_cobranza": "Equipo Cobranza",
        }

        response = self.client.post(
            "/api/crm/crear/",
            data=json.dumps(payload),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)

        cliente = Cliente.objects.get()
        self.assertEqual(cliente.razon_social, "Juan Perez Gomez")
        self.assertEqual(cliente.rfc, "PEGJ9001011A2")
        self.assertEqual(cliente.regimen_fiscal, "Regimen Simplificado de Confianza")
        self.assertFalse(cliente.es_persona_moral)
        self.assertEqual(cliente.calle, "Rio Lerma")
        self.assertEqual(cliente.numero_exterior, "45")
        self.assertEqual(cliente.numero_interior, "2B")
        self.assertEqual(cliente.colonia, "Cuauhtemoc")
        self.assertEqual(cliente.estado, "Ciudad de Mexico")
        self.assertEqual(cliente.codigo_postal, "06500")

    def test_create_endpoint_allows_operational_client_without_fiscal_data(self):
        payload = {
            "entidad_relacionada_id": self.entidad.id,
            "razon_social": "Cliente Prueba WhatsApp",
            "telefono": "5512345678",
            "codigo_pais": "+52",
            "condiciones_pago": "Contado",
        }

        response = self.client.post(
            "/api/crm/crear/",
            data=json.dumps(payload),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        cliente = Cliente.objects.get(razon_social="Cliente Prueba WhatsApp")
        self.assertEqual(cliente.rfc, "")
        self.assertIsNone(cliente.regimen_fiscal)
        self.assertEqual(cliente.codigo_postal, "")

    def test_list_endpoint_returns_paginated_payload(self):
        for index in range(24):
            Cliente.objects.create(
                entidad_relacionada=self.entidad,
                razon_social=f"Cliente {index}",
                rfc=f"AAA900101A{index:02d}"[:13],
                telefono=f"5512345{index:03d}",
            )

        response = self.client.get(
            "/api/crm/lista/",
            {"page": 2, "page_size": 20},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["total"], 24)
        self.assertEqual(body["page"], 2)
        self.assertEqual(body["page_size"], 20)
        self.assertEqual(body["total_pages"], 2)
        self.assertEqual(len(body["items"]), 4)
        self.assertEqual(len(body["entidades"]), 1)



    def test_list_endpoint_hides_inactive_clients_by_default(self):
        active = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente activo",
            rfc="ACT900101AB1",
            telefono="5511111111",
            activo=True,
        )
        inactive = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente inactivo",
            rfc="INA900101AB1",
            telefono="5522222222",
            activo=False,
        )

        response = self.client.get("/api/crm/lista/", **self.auth_headers)

        self.assertEqual(response.status_code, 200)
        names = [item["nombre"] for item in response.json()["items"]]
        self.assertIn(active.razon_social, names)
        self.assertNotIn(inactive.razon_social, names)

        inactive_response = self.client.get(
            "/api/crm/lista/",
            {"estatus": "Inactivo"},
            **self.auth_headers,
        )

        self.assertEqual(inactive_response.status_code, 200)
        inactive_names = [item["nombre"] for item in inactive_response.json()["items"]]
        self.assertIn(inactive.razon_social, inactive_names)
        self.assertNotIn(active.razon_social, inactive_names)

    def test_list_endpoint_filters_clients_with_open_balance(self):
        cliente_con_adeudo = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente con adeudo",
            rfc="ADE900101AB1",
            telefono="5511111111",
        )
        cliente_pagado = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente pagado",
            rfc="PAG900101AB1",
            telefono="5522222222",
        )
        Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente sin cuentas",
            rfc="SIN900101AB1",
            telefono="5533333333",
        )
        CuentaPorCobrar.objects.create(
            entidad_relacionada=self.entidad,
            cliente_relacionado=cliente_con_adeudo,
            concepto="Renta pendiente",
            fecha_emision=date(2026, 5, 1),
            fecha_vencimiento=date(2026, 5, 5),
            monto_total=Decimal("5000"),
            monto_pagado=Decimal("1500"),
            estatus_adeudo="PARCIAL",
            periodicidad="MENSUAL",
        )
        CuentaPorCobrar.objects.create(
            entidad_relacionada=self.entidad,
            cliente_relacionado=cliente_pagado,
            concepto="Renta pagada",
            fecha_emision=date(2026, 5, 1),
            fecha_vencimiento=date(2026, 5, 5),
            monto_total=Decimal("5000"),
            monto_pagado=Decimal("5000"),
            estatus_adeudo="PAGADO",
            periodicidad="MENSUAL",
        )

        response = self.client.get(
            "/api/crm/lista/",
            {"con_adeudo": "true"},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["total"], 1)
        self.assertEqual(body["items"][0]["id"], cliente_con_adeudo.id)
        self.assertEqual(body["items"][0]["saldo"], 3500)

    def test_list_endpoint_filters_clients_with_incomplete_data(self):
        complete = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente completo",
            rfc="COM900101AB1",
            regimen_fiscal="Regimen Simplificado de Confianza",
            correo_principal="completo@example.com",
            telefono="5511111111",
            codigo_postal="06500",
            contrato_digital_url="https://r2.test/contratos/completo.pdf",
            contrato_digital_nombre="completo.pdf",
        )
        missing_contract = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente sin contrato",
            rfc="SCO900101AB1",
            regimen_fiscal="Regimen Simplificado de Confianza",
            correo_principal="sincontrato@example.com",
            telefono="5522222222",
            codigo_postal="06500",
        )
        missing_contact = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente sin contacto",
            rfc="SCO900101AB2",
            regimen_fiscal="Regimen Simplificado de Confianza",
            telefono="",
            codigo_postal="06500",
            contrato_digital_url="https://r2.test/contratos/contacto.pdf",
            contrato_digital_nombre="contacto.pdf",
        )

        response = self.client.get(
            "/api/crm/lista/",
            {"incompletos": "true"},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.json()["items"]}
        self.assertNotIn(complete.id, ids)
        self.assertIn(missing_contract.id, ids)
        self.assertIn(missing_contact.id, ids)

    def test_list_endpoint_avoids_csf_header_noise_as_display_name(self):
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social=(
                "MAYO DE 2026 idCIF: 16110296416 VALIDA TU INFORMACION FISCAL "
                "LERL700825JI6 Datos de Identificacion del Contribuyente"
            ),
            rfc="LERL700825JI6",
            telefono="5511111111",
        )

        response = self.client.get("/api/crm/lista/", **self.auth_headers)

        self.assertEqual(response.status_code, 200)
        item = next(item for item in response.json()["items"] if item["id"] == cliente.id)
        self.assertEqual(item["nombre"], "LERL700825JI6")

    def test_update_endpoint_modifies_existing_client(self):
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Original",
            rfc="ABC900101AB1",
            telefono="5511111111",
        )

        payload = {
            "entidad_relacionada_id": self.entidad.id,
            "es_persona_moral": True,
            "nombre_comercial": "Suite 7",
            "razon_social": "Cliente Editado",
            "rfc": "ABC900101AB1",
            "regimen_fiscal": "Regimen General de Ley Personas Morales",
            "identificador": "CLI-007",
            "condiciones_pago": "Credito 30 dias",
            "archivo_csf_url": "",
            "correo_principal": "cliente@example.com",
            "codigo_pais": "+52",
            "telefono": "5599999999",
            "pais": "Mexico",
            "estado": "Yucatan",
            "ciudad": "Merida",
            "colonia": "Centro",
            "calle": "60",
            "numero_exterior": "100",
            "numero_interior": "2A",
            "codigo_postal": "97000",
            "agente_cobranza": "Equipo 1",
            "dia_corte_individual": 12,
            "dias_gracia": 3,
            "activo": True,
        }

        response = self.client.put(
            f"/api/crm/cliente/{cliente.id}/",
            data=json.dumps(payload),
            content_type="application/json",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertEqual(cliente.razon_social, "Cliente Editado")
        self.assertEqual(cliente.nombre_comercial, "Suite 7")
        self.assertTrue(cliente.es_persona_moral)
        self.assertEqual(cliente.telefono, "5599999999")
        self.assertEqual(cliente.dia_corte_individual, 12)
        self.assertEqual(cliente.dias_gracia, 3)

    def test_portal_preview_endpoint_returns_signed_customer_view(self):
        self.capa.portal_clientes_url_base = "https://app.test"
        self.capa.save(update_fields=["portal_clientes_url_base"])
        ConfiguracionComunicacion.objects.create(
            capa_negocio=self.capa,
            portal_token_horas=48,
        )
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Portal",
            rfc="POR900101AB1",
            regimen_fiscal="Regimen de Arrendamiento",
            codigo_postal="06500",
            correo_principal="portal@example.com",
            telefono="5511111111",
        )
        CuentaPorCobrar.objects.create(
            entidad_relacionada=self.entidad,
            cliente_relacionado=cliente,
            concepto="Renta mayo",
            fecha_emision=date(2026, 5, 1),
            fecha_vencimiento=date(2026, 5, 5),
            monto_total=Decimal("7500"),
            periodicidad="MENSUAL",
        )

        response = self.client.get(
            f"/api/crm/cliente/{cliente.id}/portal-preview/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["cliente"]["nombre"], "Cliente Portal")
        self.assertEqual(body["datos_fiscales"]["rfc"], "POR900101AB1")
        self.assertEqual(body["portal"]["token_horas"], 48)
        self.assertTrue(body["portal"]["url"].startswith("https://app.test/portal-cliente/"))
        self.assertEqual(len(body["cuentas"]), 1)

    def test_portal_link_endpoint_returns_token_without_building_preview(self):
        self.capa.portal_clientes_url_base = "https://app.test"
        self.capa.save(update_fields=["portal_clientes_url_base"])
        ConfiguracionComunicacion.objects.create(
            capa_negocio=self.capa,
            portal_token_horas=48,
        )
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Portal",
            rfc="POR900101AB1",
            regimen_fiscal="Regimen de Arrendamiento",
            codigo_postal="06500",
            correo_principal="portal@example.com",
            telefono="5511111111",
        )

        with patch(
            "crm.api.build_portal_payload",
            side_effect=AssertionError("El link ligero no debe cargar el portal completo."),
        ):
            response = self.client.get(
                f"/api/crm/cliente/{cliente.id}/portal-link/",
                **self.auth_headers,
            )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["cliente"]["nombre"], "Cliente Portal")
        self.assertEqual(body["portal"]["token_horas"], 48)
        self.assertTrue(body["portal"]["token"])
        self.assertTrue(body["portal"]["url"].startswith("https://app.test/portal-cliente/"))
        self.assertNotIn("cuentas", body)

    def test_cliente_csf_upload_updates_fiscal_data_and_portal_preview(self):
        self.activate_plan(modules=["clientes"])
        self.capa.portal_clientes_url_base = "https://app.test"
        self.capa.save(update_fields=["portal_clientes_url_base"])
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Pendiente",
            nombre_comercial="Ocupante Pendiente",
            rfc="PEN900101AB1",
            telefono="5511111111",
            activo=True,
        )
        upload = SimpleUploadedFile(
            "csf.pdf",
            b"%PDF-1.4 contenido simulado",
            content_type="application/pdf",
        )

        with patch(
            "crm.api.extract_csf_data_from_upload",
            return_value={
                "razon_social": "CLIENTE FISCAL SA DE CV",
                "rfc": "CFS900101AB1",
                "regimen_fiscal": "Regimen General de Ley Personas Morales",
                "codigo_postal": "06600",
                "es_persona_moral": True,
            },
        ), patch(
            "crm.api.upload_csf_to_r2",
            return_value="https://r2.test/csf/csf.pdf",
        ):
            response = self.client.post(
                f"/api/crm/cliente/{cliente.id}/extraer-csf/",
                data={"file": upload},
                **self.auth_headers,
            )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertEqual(cliente.razon_social, "CLIENTE FISCAL SA DE CV")
        self.assertEqual(cliente.nombre_comercial, "Ocupante Pendiente")
        self.assertEqual(cliente.rfc, "CFS900101AB1")
        self.assertEqual(cliente.regimen_fiscal, "Regimen General de Ley Personas Morales")
        self.assertEqual(cliente.codigo_postal, "06600")
        self.assertEqual(cliente.archivo_csf_url, "https://r2.test/csf/csf.pdf")
        body = response.json()
        self.assertEqual(body["portal"]["datos_fiscales"]["rfc"], "CFS900101AB1")
        self.assertTrue(body["portal"]["portal"]["url"].startswith("https://app.test/portal-cliente/"))

    def test_cliente_csf_upload_preserves_visible_name_when_fiscal_receiver_differs(self):
        self.activate_plan(modules=["clientes"])
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Daniel Bolanos",
            nombre_comercial="",
            rfc="BOD900101AB1",
            telefono="5511111111",
            activo=True,
        )
        upload = SimpleUploadedFile(
            "csf.pdf",
            b"%PDF-1.4 contenido simulado",
            content_type="application/pdf",
        )

        with patch(
            "crm.api.extract_csf_data_from_upload",
            return_value={
                "razon_social": "PEDRO BOLANOS",
                "rfc": "BOP900101AB1",
                "regimen_fiscal": "Regimen Simplificado de Confianza",
                "codigo_postal": "01729",
                "es_persona_moral": False,
            },
        ), patch(
            "crm.api.upload_csf_to_r2",
            return_value="https://r2.test/csf/pedro.pdf",
        ):
            response = self.client.post(
                f"/api/crm/cliente/{cliente.id}/extraer-csf/",
                data={"file": upload},
                **self.auth_headers,
            )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertEqual(cliente.nombre_comercial, "Daniel Bolanos")
        self.assertEqual(cliente.razon_social, "PEDRO BOLANOS")
        self.assertEqual(cliente.rfc, "BOP900101AB1")

    def test_cliente_contract_upload_updates_document_fields(self):
        self.activate_plan(modules=["clientes"])
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Contrato",
            rfc="CON900101AB1",
            telefono="5511111111",
            activo=True,
        )
        upload = SimpleUploadedFile(
            "contrato.pdf",
            b"%PDF-1.4 contrato simulado",
            content_type="application/pdf",
        )

        with patch(
            "crm.api.upload_contract_to_r2",
            return_value=(
                "https://r2.test/contratos/clientes/1/contrato.pdf",
                "contrato.pdf",
            ),
        ):
            response = self.client.post(
                f"/api/crm/cliente/{cliente.id}/contrato/",
                data={"file": upload},
                **self.auth_headers,
            )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertEqual(
            cliente.contrato_digital_url,
            "https://r2.test/contratos/clientes/1/contrato.pdf",
        )
        self.assertEqual(cliente.contrato_digital_nombre, "contrato.pdf")

        list_response = self.client.get("/api/crm/lista/", **self.auth_headers)
        item = next(
            item for item in list_response.json()["items"] if item["id"] == cliente.id
        )
        self.assertTrue(item["contrato_digital_cargado"])
        self.assertEqual(item["contrato_digital_nombre"], "contrato.pdf")

    def test_cliente_contract_upload_rejects_invalid_file_type(self):
        self.activate_plan(modules=["clientes"])
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Contrato",
            rfc="CON900101AB1",
            telefono="5511111111",
            activo=True,
        )
        upload = SimpleUploadedFile(
            "contrato.txt",
            b"texto no valido",
            content_type="text/plain",
        )

        response = self.client.post(
            f"/api/crm/cliente/{cliente.id}/contrato/",
            data={"file": upload},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("contrato debe ser", response.json()["detail"].lower())

    def test_delete_endpoint_soft_disables_client_with_history(self):
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente con historial",
            rfc="XYZ900101AB1",
            telefono="5512222222",
            activo=True,
        )
        CuentaPorCobrar.objects.create(
            entidad_relacionada=self.entidad,
            cliente_relacionado=cliente,
            concepto="Renta abril",
            fecha_emision=date(2026, 4, 1),
            fecha_vencimiento=date(2026, 4, 5),
            monto_total=Decimal("5000"),
            periodicidad="MENSUAL",
        )

        response = self.client.delete(
            f"/api/crm/cliente/{cliente.id}/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertFalse(cliente.activo)
        self.assertEqual(response.json()["accion"], "desactivado")

    def test_reactivate_endpoint_enables_inactive_client(self):
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente por reactivar",
            rfc="REA900101AB1",
            telefono="5514444444",
            activo=False,
        )

        response = self.client.post(
            f"/api/crm/cliente/{cliente.id}/reactivar/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertTrue(cliente.activo)
        self.assertEqual(response.json()["accion"], "reactivado")

    def test_batch_import_updates_existing_client_by_id(self):
        cliente = Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente Batch",
            rfc="BAT900101AB1",
            telefono="5513333333",
            correo_principal="antes@example.com",
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Cliente ID",
                "Entidad ID",
                "Entidad",
                "Activo",
                "Es persona moral",
                "Nombre comercial",
                "Razon social",
                "RFC",
                "Regimen fiscal",
                "Identificador",
                "Condiciones pago",
                "Archivo CSF URL",
                "Correo principal",
                "Codigo pais",
                "Telefono",
                "Pais",
                "Estado",
                "Ciudad",
                "Colonia",
                "Calle",
                "Numero exterior",
                "Numero interior",
                "Codigo postal",
                "Agente cobranza",
                "Dia corte individual",
                "Dias gracia",
            ]
        )
        sheet.append(
            [
                cliente.id,
                self.entidad.id,
                self.entidad.nombre_comercial,
                "Si",
                "No",
                "Nuevo comercial",
                "Cliente Batch Actualizado",
                "BAT900101AB1",
                "Regimen Simplificado de Confianza",
                "L-100",
                "Contado",
                "",
                "despues@example.com",
                "+52",
                "5514444444",
                "Mexico",
                "Yucatan",
                "Merida",
                "Centro",
                "60",
                "200",
                "3B",
                "97000",
                "Equipo Batch",
                15,
                5,
            ]
        )

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        upload = SimpleUploadedFile(
            "clientes.xlsx",
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/crm/batch/importar/",
            {"file": upload},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        cliente.refresh_from_db()
        self.assertEqual(cliente.razon_social, "Cliente Batch Actualizado")
        self.assertEqual(cliente.correo_principal, "despues@example.com")
        self.assertEqual(cliente.telefono, "5514444444")
        self.assertEqual(cliente.codigo_postal, "97000")
        self.assertEqual(response.json()["clientes_actualizados"], 1)


class ClienteBatchTenantIsolationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="owner-a@test.local",
            email="owner-a@test.local",
            password="secret123",
        )
        self.capa = CapaNegocio.objects.create(nombre="Capa A", tipo_capa="OPERADORA")
        self.other_capa = CapaNegocio.objects.create(nombre="Capa B", tipo_capa="OPERADORA")
        MembresiaCapaNegocio.objects.create(
            user=self.user,
            capa_negocio=self.capa,
            rol="OWNER_ADMIN",
            activo=True,
        )
        session = create_access_session(self.user, provider="PASSWORD")
        tokens = issue_token_pair(session)
        self.auth_headers = {
            "HTTP_AUTHORIZATION": f"Bearer {tokens['access_token']}",
            "HTTP_X_BETTERP_CAPA_ID": str(self.capa.id),
        }
        self.entidad = EntidadNegocio.objects.create(
            capa_negocio=self.capa,
            nombre_comercial="Entidad A",
        )
        self.other_entidad = EntidadNegocio.objects.create(
            capa_negocio=self.other_capa,
            nombre_comercial="Entidad B",
        )

    def build_batch_file(self, rows: list[list[object]]) -> SimpleUploadedFile:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Cliente ID",
                "Entidad ID",
                "Entidad",
                "Activo",
                "Es persona moral",
                "Nombre comercial",
                "Razon social",
                "RFC",
                "Regimen fiscal",
                "Identificador",
                "Condiciones pago",
                "Archivo CSF URL",
                "Correo principal",
                "Codigo pais",
                "Telefono",
                "Pais",
                "Estado",
                "Ciudad",
                "Colonia",
                "Calle",
                "Numero exterior",
                "Numero interior",
                "Codigo postal",
                "Agente cobranza",
                "Dia corte individual",
                "Dias gracia",
            ]
        )
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return SimpleUploadedFile(
            "clientes.xlsx",
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_batch_template_only_exports_current_capa_clients(self):
        Cliente.objects.create(
            entidad_relacionada=self.entidad,
            razon_social="Cliente visible",
            rfc="AAA900101AA1",
            telefono="5511111111",
        )
        Cliente.objects.create(
            entidad_relacionada=self.other_entidad,
            razon_social="Cliente oculto",
            rfc="BBB900101BB1",
            telefono="5522222222",
        )

        response = self.client.get(
            "/api/crm/batch/plantilla/",
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        exported_names = [row[6] for row in rows[1:]]
        self.assertIn("Cliente visible", exported_names)
        self.assertNotIn("Cliente oculto", exported_names)

    def test_batch_import_does_not_update_client_from_another_capa(self):
        other_client = Cliente.objects.create(
            entidad_relacionada=self.other_entidad,
            razon_social="Cliente externo",
            rfc="EXT900101AA1",
            telefono="5599999999",
        )
        upload = self.build_batch_file(
            [
                [
                    other_client.id,
                    self.entidad.id,
                    self.entidad.nombre_comercial,
                    "Si",
                    "No",
                    "Nuevo",
                    "Cliente nuevo interno",
                    "INT900101AA1",
                    "",
                    "INT-001",
                    "",
                    "",
                    "interno@example.com",
                    "+52",
                    "5512345678",
                    "Mexico",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    1,
                    0,
                ]
            ]
        )

        response = self.client.post(
            "/api/crm/batch/importar/",
            {"file": upload},
            **self.auth_headers,
        )

        self.assertEqual(response.status_code, 200)
        other_client.refresh_from_db()
        self.assertEqual(other_client.razon_social, "Cliente externo")
        self.assertTrue(
            Cliente.objects.filter(
                entidad_relacionada=self.entidad,
                razon_social="Cliente nuevo interno",
            ).exists()
        )
